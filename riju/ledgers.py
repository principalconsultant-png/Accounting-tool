"""Read-only ledger books and customer handover exports from posted vouchers."""
import csv
import html
import io
import os
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from . import core
from .foundation import branches

COLUMNS=('Date','Voucher','Reference','Narration','Branch','Department','Debit (INR)','Credit (INR)','Balance (INR)','Dr / Cr')


def number(paise,grouped=False):
    whole,fraction=divmod(abs(paise),100)
    return ('-' if paise<0 else '')+(f'{whole:,}' if grouped else str(whole))+f'.{fraction:02}'


def balance(paise):
    return number(abs(paise),True)+(' Dr' if paise>0 else ' Cr' if paise<0 else '')


def ledger_pack(state,codes,start,end,branch='ALL'):
    """One snapshot shared by the screen and all exports. No drafts or side effects.

    Initial opening vouchers dated on the first selected day belong to opening
    balance. An opening later in the selected range is a dated movement. Closing
    journals remain movements so closing agrees with the actual trial balance.
    """
    core.verify_audit(state)
    if date.fromisoformat(start).isoformat()!=start or date.fromisoformat(end).isoformat()!=end or start>end:
        raise ValueError('Use a valid From/Through date range in YYYY-MM-DD format.')
    names=branches(state)
    if branch!='ALL' and branch not in names:raise ValueError('Select an existing branch or all operating units.')
    by_code={a['code']:a for a in state['accounts']}
    codes=list(dict.fromkeys(codes))
    if not codes:raise ValueError('Select at least one ledger account.')
    if any(c not in by_code for c in codes):raise ValueError('A selected ledger no longer exists.')
    books={c:dict(code=c,name=by_code[c]['name'],kind=by_code[c]['kind'],opening=0,debit=0,credit=0,closing=0,rows=[],opening_refs=[]) for c in codes}
    for _,voucher in sorted(enumerate(state['vouchers']),key=lambda x:(x[1]['date'],x[0])):
        if voucher['date']>end:continue
        for line in voucher['lines']:
            c=line['account']
            if c not in books or (branch!='ALL' and line['branch']!=branch):continue
            book=books[c];delta=line['debit']-line['credit']
            if voucher['date']<start or (voucher.get('opening') and voucher['date']==start):
                book['opening']+=delta
                if voucher.get('opening') and voucher['date']==start and voucher['id'] not in book['opening_refs']:book['opening_refs'].append(voucher['id'])
                continue
            metadata=voucher.get('metadata',{})
            book['rows'].append(dict(date=voucher['date'],voucher=voucher['id'],reference=metadata.get('reference',''),
              narration=voucher['narration'],branch=names[line['branch']],department=line.get('cost_center',''),
              debit=line['debit'],credit=line['credit'],closing_entry=bool(voucher.get('closing')),opening_entry=bool(voucher.get('opening'))))
    for book in books.values():
        running=book['opening']
        for row in book['rows']:
            running+=row['debit']-row['credit'];row['balance']=running
            book['debit']+=row['debit'];book['credit']+=row['credit']
        book['closing']=running
    # Independently compare with the accounting engine, including year-end closes.
    closing=core.balances(state,end,branch)
    if any(b['closing']!=closing[c] for c,b in books.items()):raise ValueError('Ledger closing does not reconcile to the trial balance. No export was created.')
    settings=state['settings']
    return dict(entity=settings['entity'],registration=dict(settings.get('registration',{})),scope='All operating units — entity total' if branch=='ALL' else names[branch],
      branch=branch,start=start,end=end,prepared=datetime.now().astimezone().isoformat(timespec='seconds'),
      books=list(books.values()),currency='INR')


def book_rows(book,start,end):
    """Opening, movements, period totals and closing, consistently for all formats."""
    def summary(day,label,debit,credit,value,refs=''):
        return (day,refs,'',label,'','',debit,credit,abs(value),'Dr' if value>0 else 'Cr' if value<0 else '')
    yield summary(start,'Opening balance',0,0,book['opening'],', '.join(book['opening_refs']))
    for row in book['rows']:
        n=row['balance']
        yield (row['date'],row['voucher'],row['reference'],row['narration'],row['branch'],row['department'],row['debit'],row['credit'],abs(n),'Dr' if n>0 else 'Cr' if n<0 else '')
    yield summary(end,'Period totals',book['debit'],book['credit'],book['closing'])
    yield summary(end,'Closing balance',0,0,book['closing'])


def safe_csv(value):
    text=str(value)
    # Excel must never execute a narration/ledger name as a formula.
    return "'"+text if text.lstrip().startswith(('=','+','-','@')) or text.startswith(('\t','\r','\n')) else text


def export_csv(pack,path):
    out=io.StringIO(newline='');writer=csv.writer(out)
    writer.writerow(['Entity',safe_csv(pack['entity'])]);writer.writerow(['Scope',safe_csv(pack['scope'])])
    writer.writerow(['From',pack['start'],'Through',pack['end'],'Currency','INR']);writer.writerow(['Prepared',pack['prepared']])
    for key,value in pack['registration'].items():
        if value:writer.writerow([key.upper(),safe_csv(value)])
    for book in pack['books']:
        writer.writerow([]);writer.writerow(['Ledger',book['code'],safe_csv(book['name']),book['kind']]);writer.writerow(COLUMNS)
        for row in book_rows(book,pack['start'],pack['end']):
            writer.writerow([number(x) if i in (6,7,8) else safe_csv(x) for i,x in enumerate(row)])
    Path(path).write_text(out.getvalue(),encoding='utf-8-sig',newline='')


def export_xlsx(pack,path):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    wb=Workbook();wb.remove(wb.active)
    for book in pack['books']:
        ws=wb.create_sheet(book['code'])
        # Force all text to literal strings, including strings beginning '='.
        def append(values):
            ws.append(values)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str):cell.data_type='s'
        append([pack['entity']]);append([book['code']+' · '+book['name']]);append([pack['scope']])
        append(['From',pack['start'],'Through',pack['end'],'Currency','INR']);append(['Prepared',pack['prepared']])
        for key,value in pack['registration'].items():
            if value:append([key.upper(),value])
        append(list(COLUMNS));header=ws.max_row
        for cell in ws[header]:cell.font=Font(bold=True);cell.fill=PatternFill('solid',fgColor='E3EEDF')
        for row in book_rows(book,pack['start'],pack['end']):
            append([Decimal(x)/100 if i in (6,7,8) else x for i,x in enumerate(row)])
            for col in (7,8,9):ws.cell(ws.max_row,col).number_format='#,##0.00'
            ws.cell(ws.max_row,4).alignment=Alignment(wrap_text=True,vertical='top')
            if row[3] in ('Opening balance','Period totals','Closing balance'):
                for c in ws[ws.max_row]:c.font=Font(bold=True)
        for i,width in enumerate((13,23,22,65,24,24,17,17,19,9),1):ws.column_dimensions[get_column_letter(i)].width=width
        ws.freeze_panes=f'A{header+1}';ws.print_title_rows=f'1:{header}';ws.print_area=f'A1:J{ws.max_row}'
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.oddFooter.center.text='Page &P of &N';ws.oddFooter.right.text='Riju Books — INR'
    wb.save(path);wb.close()


def html_document(pack):
    esc=lambda value:html.escape(str(value),quote=True)
    sections=[]
    for book in pack['books']:
        reg=' · '.join(esc(k.upper()+': '+v) for k,v in pack['registration'].items() if v)
        rows=[]
        for row in book_rows(book,pack['start'],pack['end']):
            cells=''.join('<td'+(' class="amount"' if i in (6,7,8) else '')+'>'+esc(number(v,True) if i in (6,7,8) else v)+'</td>' for i,v in enumerate(row))
            rows.append('<tr'+(' class="total"' if row[3] in ('Opening balance','Period totals','Closing balance') else '')+'>'+cells+'</tr>')
        sections.append('<section><header><h1>'+esc(pack['entity'])+'</h1><p>'+reg+'</p><h2>'+esc(book['code']+' · '+book['name'])+'</h2><p>'+esc(pack['scope'])+' | '+esc(pack['start'])+' to '+esc(pack['end'])+' | INR</p><p>Prepared: '+esc(pack['prepared'])+'</p></header><table><thead><tr><th colspan="10">'+esc(book['code']+' · '+book['name']+' | '+pack['scope']+' | '+pack['start']+' to '+pack['end'])+'</th></tr><tr>'+''.join('<th>'+esc(c)+'</th>' for c in COLUMNS)+'</tr></thead><tbody>'+''.join(rows)+'</tbody></table><footer>Prepared by: ____________________ &nbsp; Reviewed by / CEO: ____________________</footer></section>')
    return '''<!doctype html><html><head><meta charset="utf-8"><meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'; script-src 'unsafe-inline'"><title>Ledger books</title><style>
body{font:12px Arial,sans-serif;color:#111;margin:20px}h1{font-size:20px}h2{font-size:17px}p{margin:5px 0}table{border-collapse:collapse;width:100%;margin:15px 0;table-layout:fixed}th,td{border:1px solid #999;padding:6px;vertical-align:top;overflow-wrap:anywhere;white-space:pre-wrap}th{background:#eee}th:nth-child(4){width:22%}.amount{text-align:right;font-variant-numeric:tabular-nums}.total{font-weight:bold}footer{margin:25px 0}section+section{break-before:page;margin-top:35px}.no-print{padding:12px;background:#f9e5d1}button{padding:10px}thead{display:table-header-group}@page{size:A4 landscape;margin:12mm}@media print{body{margin:0;font-size:9px}.no-print{display:none}th{background:#eee}section{margin:0}tr{break-inside:avoid}}
</style></head><body><div class="no-print"><button onclick="window.print()">Print / Save as PDF</button> Use the print dialog to choose your printer or save a PDF. This document works offline.</div>'''+''.join(sections)+'</body></html>'


def export_html(pack,path):
    Path(path).write_text(html_document(pack),encoding='utf-8')


def export_pdf(pack,path):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,LongTable,TableStyle,PageBreak
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4,landscape
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    font='Helvetica'
    # Native Windows and common Linux fonts provide Unicode without online fonts.
    for candidate in (Path(os.environ.get('WINDIR','C:/Windows'))/'Fonts'/'segoeui.ttf',Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')):
        if candidate.is_file():
            font='RijuLedger'
            if font not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont(font,str(candidate)))
            break
    style=ParagraphStyle('Ledger',fontName=font,fontSize=7,leading=10,wordWrap='CJK')
    numeric=ParagraphStyle('LedgerNumber',parent=style,alignment=2)
    heading=ParagraphStyle('LedgerHeading',parent=style,fontSize=12,leading=16,spaceAfter=7)
    esc=lambda s:html.escape(str(s)).replace('\n','<br/>')
    p=lambda s:Paragraph(esc(s),style)
    doc=SimpleDocTemplate(str(path),pagesize=landscape(A4),leftMargin=26,rightMargin=26,topMargin=28,bottomMargin=37,
                         title='Ledger books',author=pack['entity'])
    story=[]
    for index,book in enumerate(pack['books']):
        if index:story.append(PageBreak())
        story.append(Paragraph(esc(pack['entity']),heading))
        reg=' | '.join(k.upper()+': '+v for k,v in pack['registration'].items() if v)
        if reg:story.extend([p(reg),Spacer(1,6)])
        caption=pack['entity']+' | '+book['code']+' · '+book['name']+' | '+pack['scope']+' | '+pack['start']+' to '+pack['end']+' | INR'
        data=[[p(caption)]+['']*9,[p(c) for c in COLUMNS]]
        for row in book_rows(book,pack['start'],pack['end']):data.append([Paragraph(esc(number(v,True)),numeric) if i in (6,7,8) else p(v) for i,v in enumerate(row)])
        table=LongTable(data,colWidths=[49,80,62,170,70,70,65,65,65,44],repeatRows=2,splitInRow=1)
        rules=[('SPAN',(0,0),(-1,0)),('BACKGROUND',(0,0),(-1,1),colors.HexColor('#E3EEDF')),('GRID',(0,1),(-1,-1),0.3,colors.grey),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),5)]
        for r in (2,len(data)-2,len(data)-1):rules.append(('BACKGROUND',(0,r),(-1,r),colors.HexColor('#F4F4F4')))
        table.setStyle(TableStyle(rules));story.extend([table,Spacer(1,18),p('Prepared by: ____________________     Reviewed by / CEO: ____________________')])
    def footer(canvas,document):
        canvas.saveState();canvas.setFont(font,7)
        canvas.drawString(26,17,'Prepared: '+pack['prepared'])
        canvas.drawRightString(landscape(A4)[0]-26,17,'Riju Books | Page '+str(document.page));canvas.restoreState()
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
