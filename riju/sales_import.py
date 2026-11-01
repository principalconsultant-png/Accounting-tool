"""Read the hospital's daily-sales XLSX into a reviewable, balanced queue."""
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
import copy
import hashlib
import io
import json
from pathlib import Path
import re
from . import core
from .revenue import MODES, receipt_lines

def normalized(value):return re.sub(r'[^a-z0-9]','',str(value or '').lower())

def amount(value, cell):
    if value is None or value=='':return 0
    if isinstance(value,bool):raise ValueError(f'{cell}: amount cannot be TRUE/FALSE.')
    if isinstance(value,str) and value.startswith('='):raise ValueError(f'{cell}: amount columns must contain numeric values, not formulas.')
    try:
        raw=Decimal(str(value).replace(',','').strip())
        if not raw.is_finite() or raw<0 or raw*100!=(raw*100).to_integral_value():raise ValueError()
    except (InvalidOperation,ValueError):raise ValueError(f'{cell}: enter a non-negative amount with at most two decimals.') from None
    return int(raw*100)

def row_key(row):
    reference=[row['branch'],row['date']]
    if row.get('receipt'):reference += ['receipt',row['receipt']]
    elif row.get('invoice'):reference += ['invoice',row['invoice']]
    else:reference += [row['revenues'],row['payments']]
    return hashlib.sha256(json.dumps(reference,sort_keys=True).encode()).hexdigest()

def read_sales_xlsx(raw,branch,valid_branches=None):
    from openpyxl import load_workbook
    if branch not in (valid_branches if valid_branches is not None else core.BRANCHES):raise ValueError('Select the operating unit for this report.')
    if len(raw)>25*1024*1024:raise ValueError('Sales report exceeds 25 MB.')
    from zipfile import ZipFile
    with ZipFile(io.BytesIO(raw)) as archive:
        if sum(i.file_size for i in archive.infolist())>100*1024*1024:raise ValueError('Uncompressed workbook exceeds 100 MB.')
    source=load_workbook(io.BytesIO(raw),data_only=False,read_only=True,keep_links=False)
    cached=load_workbook(io.BytesIO(raw),data_only=True,read_only=True,keep_links=False)
    rows=[];warnings=[];known=set();found_sheets=0
    try:
        for sheet in source:
            if sheet.max_row>100000 or sheet.max_column>100:raise ValueError('Unexpectedly large sales worksheet.')
            header=None
            # Source has a two-row heading; accept shifts but require all modes.
            for heading_row,cells in enumerate(sheet.iter_rows(min_row=1,max_row=min(20,sheet.max_row)),1):
                values={normalized(c.value):col for col,c in enumerate(cells,1) if c.value is not None}
                if all(normalized(m[1]) in values for m in MODES):
                    header=heading_row;columns={m[0]:values[normalized(m[1])] for m in MODES}
                    payments={k:values.get(normalized(label)) for k,label in [('card','Card'),('upi','UPI'),('cash','Cash'),('other','Other source')]}
                    subtotal=values.get('subtotal');total=values.get('total');break
            if header is None:
                if any(c.value is not None for row in sheet.iter_rows() for c in row):
                    raise ValueError(f'{sheet.title}: cannot find all 15 daily-sales headings. No sheets were silently skipped.')
                continue
            found_sheets+=1
            if not all(payments.values()):raise ValueError(f'{sheet.title}: Card, UPI, Cash and Other source columns are required.')
            value_rows=cached[sheet.title].iter_rows(min_row=header+1)
            for r,(cells,cached_cells) in enumerate(zip(sheet.iter_rows(min_row=header+1),value_rows),header+1):
                values={col:c.value for col,c in enumerate(cells,1)};cache={col:c.value for col,c in enumerate(cached_cells,1)}
                relevant=[values.get(i) for i in list(columns.values())+list(payments.values())]
                if not any(v not in (None,'',0) for v in relevant):continue
                raw_date=values.get(2)
                if isinstance(raw_date,datetime):day=raw_date.date().isoformat()
                elif isinstance(raw_date,date):day=raw_date.isoformat()
                else:
                    try:day=date.fromisoformat(str(raw_date).strip()).isoformat()
                    except ValueError:raise ValueError(f'{sheet.title} row {r}: transaction date must be a real date or YYYY-MM-DD.') from None
                revenues={key:amount(values.get(col),f'{sheet.title} row {r} {key}') for key,col in columns.items()}
                collected={key:amount(values.get(col),f'{sheet.title} row {r} {key}') for key,col in payments.items()}
                service_total=sum(revenues.values());payment_total=sum(collected.values())
                if service_total<=0 or service_total!=payment_total:
                    raise ValueError(f'{sheet.title} row {r}: all 15 service columns total {core.rupees(service_total)}, payments total {core.rupees(payment_total)}. Correct the report before importing.')
                for label,col,expected in [('Sub total',subtotal,service_total),('Total',total,payment_total)]:
                    if not col:continue
                    v=cache.get(col)
                    if v is not None and amount(v,f'row {r} {label}')!=expected:
                        raise ValueError(f'{sheet.title} row {r}: {label} disagrees with the independently calculated total; check formulas, including Treatment.')
                    if v is None:warnings.append(f'{sheet.title} row {r}: {label} has no cached result; source amounts were independently summed.')
                row={'date':day,'branch':branch,'revenues':revenues,'payments':collected,'amount':payment_total,
                     'receipt':str(values.get(7) or '').strip(),'invoice':str(values.get(6) or '').strip(),
                     'source_row':r,'sheet':sheet.title}
                row['key']=row_key(row)
                if row['key'] in known:raise ValueError(f'{sheet.title} row {r}: duplicate receipt/invoice in this report.')
                known.add(row['key']);rows.append(row)
    finally:source.close();cached.close()
    if not found_sheets or not rows:raise ValueError('No daily-sales receipts found.')
    return rows, list(dict.fromkeys(warnings))

def stage_sales(state,raw,filename,branch,treatment_policy='earned',other_policy='clearing'):
    digest=hashlib.sha256(raw).hexdigest()
    if any(q['digest']==digest and q.get('branch',branch)==branch for q in state['queue']):raise ValueError('This report is already in the queue for that branch.')
    from .foundation import branches
    rows,warnings=read_sales_xlsx(raw,branch,branches(state,active_only=True))
    prior={r['key'] for q in state['queue'] if q.get('kind')=='daily_sales' for i,r in enumerate(q['rows']) if q.get('row_reviews',{}).get(str(i),{}).get('status')!='REJECTED'}
    if any(r['key'] in prior for r in rows):raise ValueError('A receipt in this report is already queued or posted for this branch. Review the earlier import.')
    if any(r['revenues']['treatment_receipt'] for r in rows) and treatment_policy=='review':
        warnings.append('Trtmt Pymt RECD requires earned-revenue versus advance classification before approval.')
    q={'id':hashlib.sha256((digest+branch).encode()).hexdigest()[:12],'digest':digest,'filename':Path(filename).name,
       'kind':'daily_sales','branch':branch,'rows':rows,'status':'PENDING_REVIEW','warnings':warnings,
       'treatment_policy':treatment_policy,'other_policy':other_policy}
    state['queue'].append(q);core.add_audit(state,'SALES_REPORT_STAGED',{'id':q['id'],'branch':branch,'receipts':len(rows)})
    return q

def approve_sales(state,queue_id):
    candidate=copy.deepcopy(state)
    q=next(q for q in candidate['queue'] if q['id']==queue_id)
    if q['status']!='PENDING_REVIEW':raise ValueError('Only pending reports may be posted.')
    prior={r['key'] for item in candidate['queue'] if item.get('kind')=='daily_sales' and item['status']=='PROCESSED' for i,r in enumerate(item['rows']) if item.get('row_reviews',{}).get(str(i),{}).get('status')!='REJECTED'}
    reviews=q.get('row_reviews')
    if reviews and (len(reviews)!=len(q['rows']) or any(x['status'] not in ('APPROVED','REJECTED') for x in reviews.values())):raise ValueError('Review every sales row before approval.')
    selected=[row for i,row in enumerate(q['rows']) if not reviews or reviews[str(i)]['status']=='APPROVED']
    if not selected:raise ValueError('No approved sales rows.')
    for row in selected:
        if row['key'] in prior:raise ValueError('A source receipt has already been posted.')
        lines=receipt_lines(candidate,row['branch'],row['revenues'],row['payments'],q['treatment_policy'],q['other_policy'])
        narration=f"Daily sales {row['date']} · Receipt {row['receipt'] or '—'} · Invoice {row['invoice'] or '—'} · {q['filename']} / {row['sheet']} row {row['source_row']}"
        core.post(candidate,row['date'],narration,lines,source={'kind':'daily_sales','document_id':q['id'],'record_id':row['key'],'row_count':1})
    q['status']='PROCESSED';core.add_audit(candidate,'SALES_REPORT_APPROVED',{'id':q['id'],'receipts':len(selected),'rejected':len(q['rows'])-len(selected)})
    state.clear();state.update(candidate)


def review_sales_row(state,queue_id,index,approved,reason):
    q=next((q for q in state['queue'] if q['id']==queue_id and q.get('kind')=='daily_sales'),None)
    if not q or q['status']!='PENDING_REVIEW' or type(index)!=int or not 0<=index<len(q['rows']) or not str(reason).strip():raise ValueError('Review a pending sales row with a reason.')
    if approved:receipt_lines(state,q['branch'],q['rows'][index]['revenues'],q['rows'][index]['payments'],q['treatment_policy'],q['other_policy'])
    q.setdefault('row_reviews',{})[str(index)]={'status':'APPROVED' if approved else 'REJECTED','reason':str(reason).strip()}
    core.add_audit(state,'SALES_ROW_REVIEWED',dict(batch=queue_id,row=index,review=q['row_reviews'][str(index)]))
