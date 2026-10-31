"""INR bank statements, reviewed matching and dated reconciliation proofs."""
import copy
import csv
import hashlib
import io
import json
import re
import uuid
from collections import defaultdict,Counter
from datetime import date,datetime,timedelta,timezone
from decimal import Decimal,InvalidOperation,ROUND_HALF_UP
from difflib import SequenceMatcher
from pathlib import Path
from . import core,foundation


STATEMENT_ALIASES={'date':('date','transactiondate','posteddate'), 'value_date':('valuedate',), 'narration':('narration','description','particulars','memo'), 'reference':('reference','utr','chequeno','chequenumber','refno'), 'debit':('debit','withdrawal','withdrawalamt','withdrawalamount'), 'credit':('credit','deposit','depositamt','depositamount'), 'amount':('amount','signedamount'), 'balance':('balance','closingbalance','runningbalance'), 'fitid':('fitid',)}

def _day(value):
    if date.fromisoformat(value).isoformat()!=value:raise ValueError('Use YYYY-MM-DD for statement/report dates.')
    return value


def amount(value):
    if value in (None,''):return 0
    if isinstance(value,bool):raise ValueError('Bank amounts cannot be TRUE/FALSE.')
    s=str(value).strip().replace('₹','').replace('INR','').replace(',','').strip()
    negative=s.startswith('(') and s.endswith(')')
    if negative:s=s[1:-1]
    if s.upper().endswith((' CR',' DR')):
        negative=s.upper().endswith(' DR');s=s[:-3].strip()
    try:
        d=Decimal(s)
        if not d.is_finite() or abs(d)>=Decimal('1e18') or d*100!=(d*100).to_integral_value():raise ValueError()
        return -abs(int(d*100)) if negative else int(d*100)
    except (InvalidOperation,ValueError):raise ValueError('Enter an INR amount with at most two decimal places.') from None


def _date(value,fmt):
    if isinstance(value,datetime):return value.date().isoformat()
    if isinstance(value,date):return value.isoformat()
    text=str(value).strip()
    if fmt=='ISO':return _day(text)
    formats={'DD/MM/YYYY':'%d/%m/%Y','DD-MM-YYYY':'%d-%m-%Y'}
    if fmt not in formats:raise ValueError('Choose a supported statement date format.')
    return datetime.strptime(text,formats[fmt]).date().isoformat()


def _normal(value):return re.sub(r'[^a-z0-9]','',str(value or '').lower())
def _now():return datetime.now(timezone.utc).isoformat()
def active(match,through=None):
    if through and match.get('effective_from') and match['effective_from']>through:return False
    if match.get('status','ACTIVE')=='ACTIVE':return True
    return bool(through and match.get('undo_date') and match['undo_date']>through)


def bank_accounts(state):
    return [a for a in state['accounts'] if a['code']!='1000' and a['kind']=='Asset' and state['account_details'].get(a['code'],{}).get('group')=='Cash and bank' and a['code'] not in state.get('foreign_accounts',{})]


def _account(state,account):
    if account not in {a['code'] for a in bank_accounts(state)}:raise ValueError('Choose an INR bank ledger; foreign-currency accounts use the FX workflow.')


def _scope(state,scope):
    if scope!='ALL' and scope not in foundation.branches(state):raise ValueError('Unknown statement branch scope.')


def _canonical(row):
    r=copy.deepcopy(row)
    r.setdefault('bank_account','1010');r.setdefault('value_date',r['date'])
    r.setdefault('credit',max(r.get('amount',0),0));r.setdefault('debit',max(-r.get('amount',0),0))
    r['amount']=r['credit']-r['debit'];r.setdefault('balance',None);r.setdefault('fitid','')
    return r


def bank_rows(state,account):return [_canonical(b) for b in state['bank_lines'] if b.get('bank_account','1010')==account and not b.get('void')]


def _raw_rows(raw,kind,sheet=None,header_row=1):
    if len(raw)>25*1024*1024:raise ValueError('Statement exceeds 25 MB.')
    if kind=='csv':
        try:
            stream=io.StringIO(raw.decode('utf-8-sig'))
            for _ in range(header_row-1):next(stream,None)
            reader=csv.DictReader(stream)
            headers=reader.fieldnames or []
            if len(set(headers))!=len(headers):raise ValueError('Statement column headings must be unique.')
            return list(reader),{'columns':headers}
        except UnicodeDecodeError:raise ValueError('Save the bank CSV using UTF-8 encoding, or use XLSX.') from None
    if kind=='xlsx':
        from openpyxl import load_workbook
        from zipfile import ZipFile
        with ZipFile(io.BytesIO(raw)) as z:
            if sum(x.file_size for x in z.infolist())>100*1024*1024:raise ValueError('Uncompressed workbook exceeds 100 MB.')
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
        try:
            if len(wb.sheetnames)>1 and sheet is None:raise ValueError('This workbook has multiple sheets. Enter the sheet name explicitly.')
            ws=wb[sheet or wb.sheetnames[0]]
            if ws.max_row>50001 or ws.max_column>100:raise ValueError('Unexpectedly large statement sheet.')
            rows=ws.iter_rows(values_only=True)
            for _ in range(header_row-1):next(rows,None)
            headers=next(rows,None)
            if not headers:raise ValueError('Empty statement workbook.')
            headers=[str(v).strip() if v is not None else '' for v in headers]
            if len(set(h for h in headers if h))!=len([h for h in headers if h]):raise ValueError('Statement column headings must be unique.')
            return [dict(zip(headers,r)) for r in rows if any(v is not None for v in r)],{'columns':headers}
        finally:wb.close()
    if kind=='ofx':
        text=raw.decode('utf-8-sig',errors='strict')
        def tag(block,name):
            m=re.search(r'<'+name+r'>\s*([^<\r\n]+)',block,re.I)
            if not m:return ''
            from html import unescape
            return unescape(m.group(1).strip())
        currency=tag(text,'CURDEF')
        if currency and currency.upper()!='INR':raise ValueError('This reconciliation module accepts INR statements; retain foreign units through FX accounts.')
        rows=[]
        for block in re.findall(r'<STMTTRN>(.*?)</STMTTRN>',text,re.I|re.S):
            posted=tag(block,'DTPOSTED')
            if len(posted)<8:raise ValueError('OFX transaction has no valid posted date.')
            val=tag(block,'DTUSER') or posted
            rows.append(dict(date=posted[:4]+'-'+posted[4:6]+'-'+posted[6:8],value_date=val[:4]+'-'+val[4:6]+'-'+val[6:8],narration=' | '.join(x for x in (tag(block,'NAME'),tag(block,'MEMO')) if x),reference=tag(block,'CHECKNUM') or tag(block,'REFNUM') or tag(block,'FITID'),amount=tag(block,'TRNAMT'),fitid=tag(block,'FITID')))
        if not rows and not re.search(r'<LEDGERBAL>',text,re.I):raise ValueError('No OFX statement transactions or balance found.')
        ledger=re.search(r'<LEDGERBAL>(.*?)</LEDGERBAL>',text,re.I|re.S)
        header={'ofx_account':tag(text,'ACCTID')}
        if ledger:header['ofx_closing']=tag(ledger.group(1),'BALAMT')
        return rows,header
    raise ValueError('Use CSV, XLSX or OFX.')


def read_statement(raw,kind,mapping=None,date_format='ISO',order='Oldest first',sheet=None,header_row=1):
    if type(header_row)!=int or not 1<=header_row<=100:raise ValueError('Header row must be between 1 and 100.')
    rows,header=_raw_rows(raw,kind,sheet,header_row)
    if kind=='ofx' and not rows:return [],header
    if len(rows)>50000:raise ValueError('Maximum 50,000 statement rows per import.')
    if not rows and not header.get('columns'):raise ValueError('No statement headings or rows found.')
    aliases=STATEMENT_ALIASES
    fields={_normal(k):k for k in (rows[0] if rows else header['columns']) if k}
    chosen={key:next((fields[a] for a in names if a in fields),None) for key,names in aliases.items()}
    for key,value in (mapping or {}).items():
        if key not in aliases:raise ValueError('Unknown column mapping: '+key)
        if value and value not in (rows[0] if rows else header['columns']):raise ValueError('Column not found: '+value)
        chosen[key]=value or None
    if not chosen['date'] or not (chosen['amount'] or chosen['debit'] and chosen['credit']):raise ValueError('Map Date plus either signed Amount or both Debit and Credit columns.')
    output=[]
    for n,row in enumerate(rows,2):
        if not any(v not in (None,'') for v in row.values()):continue
        try:
            if None in row:raise ValueError('More cells than column headings; quote amounts/narrations containing commas.')
            def value(key):return row.get(chosen[key]) if chosen[key] else None
            d=_date(value('date'),date_format);vd=_date(value('value_date'),date_format) if value('value_date') not in (None,'') else d
            if chosen['debit'] and chosen['credit']:
                debit,credit=amount(value('debit')),amount(value('credit'))
                if debit<0 or credit<0 or (debit>0)==(credit>0):raise ValueError('Exactly one positive bank debit (withdrawal) or credit (deposit) is required.')
            else:
                signed=amount(value('amount'))
                if not signed:raise ValueError('Transaction amount cannot be zero.')
                debit,credit=max(-signed,0),max(signed,0)
            output.append(dict(date=d,value_date=vd,debit=debit,credit=credit,amount=credit-debit,narration=str(value('narration') or '').strip(),reference=str(value('reference') or '').strip(),fitid=str(value('fitid') or '').strip(),balance=amount(value('balance')) if value('balance') not in (None,'') else None,source_row=n))
        except (ValueError,TypeError) as e:raise ValueError(f'Statement row {n}: {e}') from e
    if order=='Newest first':output.reverse()
    elif order!='Oldest first':raise ValueError('Choose Oldest first or Newest first.')
    if any(a['date']>b['date'] for a,b in zip(output,output[1:])):raise ValueError('Transactions are not chronological. Select the correct row order.')
    return output,header


def _signature(row):return (row['date'],row.get('value_date',row['date']),row['amount'],_normal(row.get('reference')),_normal(row.get('narration')))


def _check_balances(rows,opening,closing):
    if opening is None:return 'Missing opening balance — not verified'
    running=opening;checked=0
    for r in rows:
        running+=r['amount']
        if r.get('balance') is not None:
            if running!=r['balance']:raise ValueError('Running balance disagrees at source row '+str(r.get('source_row','?'))+'. Check order, debit/credit mapping and opening balance.')
            checked+=1
    if closing is None:return 'Missing closing balance — not verified'
    if running!=closing:raise ValueError('Opening balance plus statement movements does not equal the closing balance. No partial import was saved.')
    return 'Opening/closing verified; '+str(checked)+' running balances checked'


def _finals(state,account,scope='ALL'):
    return [s for s in state.get('bank_statements',[]) if s['bank_account']==account and s['status']=='RECONCILED' and (scope=='ALL' or s['scope']=='ALL' or s['scope']==scope)]


def _editable_bank(state,rows):
    for r in rows:
        if any(r['date']<=s['through'] for s in _finals(state,r['bank_account'],r.get('branch','ALL'))):raise ValueError('This bank period is reconciled. Reopen the statement with a reason before changing its matches or transactions.')


def import_statement(state,raw,filename,account,scope,start,end,opening='',closing='',mapping=None,date_format='ISO',order='Oldest first',sheet=None,header_row=1,template_id=None):
    template_snapshot=None
    if template_id:
        if mapping is not None or date_format!='ISO' or order!='Oldest first' or sheet is not None or header_row!=1:raise ValueError('Use either a saved layout or manual parser settings, not both.')
        from . import bank_templates
        opts,template_snapshot=bank_templates.resolve(state,template_id,raw,filename,account,scope)
        mapping=opts['mapping'];date_format=opts['date_format'];order=opts['order'];sheet=opts['sheet'];header_row=opts['header_row']
    _account(state,account);_scope(state,scope);_day(start);_day(end)
    if start>end:raise ValueError('Statement start must precede its end.')
    rows,header=read_statement(raw,Path(filename).suffix.lower().lstrip('.'),mapping,date_format,order,sheet,header_row)
    fitids=[r['fitid'] for r in rows if r['fitid']]
    if len(fitids)!=len(set(fitids)):raise ValueError('Repeated OFX FITID within one statement.')
    if any(not start<=r['date']<=end for r in rows):raise ValueError('A transaction is outside the selected statement period.')
    if scope!='ALL' and any(l['account']==account and l['branch']!=scope for v in state['vouchers'] for l in v['lines']):raise ValueError('This bank ledger is shared across branches. Import its statement using All operating units.')
    prior=[s for s in state.get('bank_statements',[]) if s['bank_account']==account and s['status']!='VOID']
    if any(s['scope']!=scope for s in prior):raise ValueError('Keep a consistent branch/shared scope for each physical bank ledger.')
    if header.get('ofx_account') and any(s.get('ofx_account') and s['ofx_account']!=header['ofx_account'] for s in prior):raise ValueError('This OFX account differs from previous statements assigned to this ledger.')
    op=amount(opening) if str(opening).strip() else None
    cl=amount(closing) if str(closing).strip() else amount(header['ofx_closing']) if header.get('ofx_closing') else None
    if not rows and (op is None or cl is None):raise ValueError('For an empty period, enter both bank-verified opening and closing balances.')
    coverage=_check_balances(rows,op,cl)
    digest=hashlib.sha256(raw).hexdigest()
    if any(s['digest']==digest and s['bank_account']==account and s['start']==start and s['through']==end and s['status']!='VOID' for s in state.get('bank_statements',[])):raise ValueError('This statement file is already imported for this bank ledger.')
    candidate=copy.deepcopy(state);existing=bank_rows(candidate,account);sid=str(uuid.uuid4());ids=[];now=_now();new_count=0;used=set()
    for r in rows:
        possible=[b for b in existing if b['id'] not in used and ((r['fitid'] and b.get('fitid')==r['fitid']) or (not r['fitid'] and _signature(b)==_signature(r) and (r['balance'] is None or b.get('balance') is None or b['balance']==r['balance'])))]
        if len(possible)>1 or possible and not r['fitid'] and not r['reference'] and (r['balance'] is None or possible[0].get('balance') is None):raise ValueError('Overlapping transactions cannot be distinguished safely. Supply FITID, a unique reference or running balances; no rows were guessed or skipped.')
        if possible:
            b=possible[0]
            if r['fitid'] and (b['date']!=r['date'] or b['amount']!=r['amount']):raise ValueError('An existing OFX FITID has different date/amount. Review the bank source.')
            if scope!='ALL' and b.get('branch')!=scope:raise ValueError('An overlapping bank record belongs to another scope.')
            bank_id=b['id']
        else:
            bank_id=str(uuid.uuid4());record=dict(r,id=bank_id,bank_account=account,branch=scope,status='UNMATCHED',statement_id=sid,imported_at=now)
            _editable_bank(candidate,[record]);candidate['bank_lines'].append(record);existing.append(record);new_count+=1
        used.add(bank_id);ids.append(bank_id)
    statement=dict(id=sid,digest=digest,filename=Path(filename).name,bank_account=account,scope=scope,start=start,through=end,opening=op,closing=cl,coverage=coverage,row_ids=ids,imported_at=now,status='OPEN',ofx_account=header.get('ofx_account',''))
    statement['parser_options']=dict(mapping=copy.deepcopy(mapping),date_format=date_format,order=order,sheet=sheet,header_row=header_row)
    statement['mapping_template']=template_snapshot
    candidate.setdefault('bank_statements',[]).append(statement)
    core.add_audit(candidate,'BANK_STATEMENT_IMPORTED',{'statement':copy.deepcopy(statement),'new_rows':new_count,'reused_rows':len(ids)-new_count})
    state.clear();state.update(candidate)
    return statement


def book_items(state,account,scope='ALL',through=None):
    result=[]
    for v in state['vouchers']:
        if through and v['date']>through or v.get('opening'):continue
        for i,l in enumerate(v['lines']):
            if l['account']==account and (scope=='ALL' or l['branch']==scope):
                result.append(dict(key=v['id']+':'+str(i),voucher=v['id'],date=v['date'],amount=l['debit']-l['credit'],branch=l['branch'],reference=v.get('metadata',{}).get('reference',''),narration=v['narration'],payee=v.get('metadata',{}).get('payee','')))
    return result


def _edges(state,match):
    if 'edges' in match:return match['edges']
    if not match.get('posting') or not match.get('bank_id'):return []
    row=next((b for b in state['bank_lines'] if b['id']==match['bank_id']),None)
    if not row:return []
    items=[b for b in book_items(state,row.get('bank_account','1010')) if b['voucher']==match['posting']]
    return [dict(bank_id=match['bank_id'],book_key=b['key'],amount=b['amount']) for b in items]


def used_allocations(state,through=None):
    banks={b['id']:_canonical(b) for b in state['bank_lines']};books={b['key']:b for a in bank_accounts(state) for b in book_items(state,a['code'])}
    bank_used=defaultdict(int);book_used=defaultdict(int)
    for m in state['matches']:
        if not active(m,through):continue
        for e in _edges(state,m):
            b,k=banks.get(e['bank_id']),books.get(e['book_key'])
            if b is None or k is None:raise ValueError('Bank match references a missing statement/book record.')
            if through and (b['date']>through or k['date']>through):continue
            bank_used[e['bank_id']]+=e['amount'];book_used[e['book_key']]+=e['amount']
    return bank_used,book_used


def _cancelled_book_keys(state,account,through=None):
    _,used=used_allocations(state);cancelled=set();keys=set()
    by_id={v['id']:v for v in state['vouchers']}
    for v in state['vouchers']:
        parent=by_id.get(v.get('reverses'))
        if not parent or parent['id'] in cancelled or through and (v['date']>through or parent['date']>through):continue
        involved=[x['id']+':'+str(i) for x in (v,parent) for i,l in enumerate(x['lines']) if l['account']==account]
        if not any(used[k] for k in involved):cancelled.update((v['id'],parent['id']));keys.update(involved)
    return keys


def outstanding(state,account,scope='ALL',through=None):
    bu,ku=used_allocations(state)
    bank=[dict(b,remaining=b['amount']-bu[b['id']]) for b in bank_rows(state,account) if (scope=='ALL' or b['branch']==scope) and (not through or b['date']<=through)]
    cancelled=_cancelled_book_keys(state,account,through)
    books=[dict(b,remaining=b['amount']-ku[b['key']]) for b in book_items(state,account,scope,through) if b['key'] not in cancelled]
    return [b for b in bank if b['remaining']],[b for b in books if b['remaining']]


def _greedy_edges(banks,books):
    # Negative lines in a single net settlement voucher are explicit offsets.
    edges=[];capacity={b['id']:b['remaining'] for b in banks}
    negatives=[b for b in books if b['remaining']<0]
    positive=[b for b in books if b['remaining']>0]
    sign=1 if sum(b['remaining'] for b in banks)>0 else -1
    if sign<0:
        result=_greedy_edges([dict(b,remaining=-b['remaining']) for b in banks],[dict(b,remaining=-b['remaining']) for b in books])
        return [dict(e,amount=-e['amount']) for e in result]
    for k in negatives:
        edges.append(dict(bank_id=banks[0]['id'],book_key=k['key'],amount=k['remaining']));capacity[banks[0]['id']]-=k['remaining']
    for k in positive:
        remaining=k['remaining']
        for b in banks:
            take=min(remaining,capacity[b['id']])
            if take>0:edges.append(dict(bank_id=b['id'],book_key=k['key'],amount=take));capacity[b['id']]-=take;remaining-=take
            if not remaining:break
        if remaining:raise ValueError('Book allocations exceed the selected bank amounts.')
    if any(capacity.values()):raise ValueError('Bank allocations do not equal book allocations.')
    return edges


def _status(state):
    used,_=used_allocations(state)
    for b in state['bank_lines']:
        total=_canonical(b)['amount'];remaining=total-used[b['id']]
        b['status']='MATCHED' if remaining==0 else 'PARTIAL' if remaining!=total else 'UNMATCHED'


def _record_match(state,banks,books,kind='existing',posting=None,clearing=None):
    if not banks or not books:raise ValueError('Select statement and book records.')
    _editable_bank(state,banks)
    if any(b['id'] in state.get('bank_exceptions',{}) for b in banks):raise ValueError('Remove the bank-exception classification before matching this row.')
    edges=_greedy_edges(banks,books)
    prior_undos=[m.get('undo_date','') for m in state['matches'] if not active(m) and any(e['bank_id'] in {b['id'] for b in banks} for e in _edges(state,m))]
    record=dict(id=str(uuid.uuid4()),effective_from=max(prior_undos,default=''),kind=kind,status='ACTIVE',bank_id=banks[0]['id'],bank_ids=[b['id'] for b in banks],bank_account=banks[0]['bank_account'],edges=edges,created_at=_now(),branches=sorted({b['branch'] for b in books}),vouchers=sorted({b['voucher'] for b in books}),posting=posting,clearing_allocations=clearing or [],policy=copy.deepcopy(state['settings'].get('bank_reconciliation',{})))
    state['matches'].append(record);_status(state)
    core.add_audit(state,'BANK_MATCHED',copy.deepcopy(record))
    return record


def match_existing(state,bank_ids,book_keys,partial_minor=None):
    candidate=copy.deepcopy(state);accounts={b.get('bank_account','1010') for b in candidate['bank_lines'] if b['id'] in bank_ids}
    if len(accounts)!=1:raise ValueError('Select bank rows from one bank ledger.')
    account=accounts.pop();all_b,all_k=outstanding(candidate,account)
    banks=[b for b in all_b if b['id'] in set(bank_ids)];books=[b for b in all_k if b['key'] in set(book_keys)]
    if len(banks)!=len(set(bank_ids)) or len(books)!=len(set(book_keys)):raise ValueError('A selected record is already matched or missing.')
    if any((b['remaining']>0)!=(banks[0]['remaining']>0) for b in banks):raise ValueError('Match deposits and withdrawals separately.')
    mixed=any((b['remaining']>0)!=(banks[0]['remaining']>0) for b in books)
    if mixed and (len(banks)!=1 or len({b['voucher'] for b in books})!=1 or partial_minor is not None):raise ValueError('Mixed debit/credit book lines must be one net compound voucher matched to one bank row, without partial matching.')
    if any(b['branch']!='ALL' and any(k['branch']!=b['branch'] for k in books) for b in banks):raise ValueError('Book branches differ from the branch-exclusive statement.')
    if partial_minor is not None:
        if type(partial_minor)!=int or partial_minor<=0:raise ValueError('Partial match must be positive integer paise.')
        if partial_minor>min(sum(abs(b['remaining']) for b in banks),sum(abs(b['remaining']) for b in books)):raise ValueError('Partial match exceeds the unallocated amount.')
        def trim(items):
            remaining=partial_minor;out=[]
            for item in items:
                take=min(abs(item['remaining']),remaining)
                if take:out.append(dict(item,remaining=take if item['remaining']>0 else -take));remaining-=take
            return out
        banks,books=trim(banks),trim(books)
    if sum(b['remaining'] for b in banks)!=sum(b['remaining'] for b in books):raise ValueError('Selected remaining bank and book amounts do not match. Use a reviewed partial amount or a separate adjustment.')
    record=_record_match(candidate,banks,books)
    state.clear();state.update(candidate)
    return record


def business_day(state,d):
    policy=state['settings'].get('bank_reconciliation',{})
    return d.isoformat() in policy.get('working_days',[]) or core.is_business_day(d,state['holidays'])


def next_business_day(state,value):
    d=date.fromisoformat(value)+timedelta(days=1)
    while not business_day(state,d):d+=timedelta(days=1)
    return d.isoformat()


def _date_near(state,book_date,bank_date):
    a,b=sorted((book_date,bank_date));d=date.fromisoformat(a);last=date.fromisoformat(b);count=0
    while d<last:
        d+=timedelta(days=1)
        if business_day(state,d):count+=1
        if count>state['settings'].get('bank_reconciliation',{}).get('tolerance',1):return False
    return True


def suggestions(state,bank_id):
    row=next(_canonical(b) for b in state['bank_lines'] if b['id']==bank_id)
    banks,books=outstanding(state,row['bank_account'],row['branch']);bank=next((b for b in banks if b['id']==bank_id),None)
    if not bank:return []
    policy=state['settings'].get('bank_reconciliation',{});bank_date=bank['value_date'] if policy.get('use_value_date') else bank['date'];out=[]
    for b in books:
        if b['remaining']!=bank['remaining']:continue
        near=_date_near(state,b['date'],bank_date);reference=bool(bank['reference'] and _normal(bank['reference'])==_normal(b['reference']))
        similarity=SequenceMatcher(None,_normal(bank['narration']),_normal(b['payee'] or b['narration'])).ratio()
        if len(_normal(b['payee']))>=4 and _normal(b['payee']) in _normal(bank['narration']):similarity=1.0
        if reference or near:
            reason='Amount, reference and date' if reference and near else 'Amount/reference; outside date tolerance' if reference else 'Amount/date and similar narration' if similarity>=.6 else 'Amount/date only — verify narration'
            out.append(dict(book_key=b['key'],score=100 if reference and near else 80 if reference else 70 if similarity>=.6 else 40,reason=reason))
    return sorted(out,key=lambda x:-x['score'])[:30]


def clearing_items(state,account,scope='ALL',through=None):
    if account not in ('1000','1110','1120'):raise ValueError('Choose cash, UPI or card clearing.')
    used=defaultdict(int);owned={m.get('posting') for m in state['matches'] if m.get('posting')}
    def effective(v):
        reversal=next((r for r in state['vouchers'] if r.get('reverses')==v['id']),None)
        return not effective(reversal) if reversal else True
    for m in state['matches']:
        if not active(m):continue
        if 'clearing_allocations' in m:
            for a in m['clearing_allocations']:used[a['key']]+=a['amount']
        elif m.get('account')==account:
            for v in state['vouchers']:
                if v['id'] in m.get('vouchers',[]):
                    for i,l in enumerate(v['lines']):
                        if l['account']==account:used[v['id']+':'+str(i)]+=l['debit']-l['credit']
    result=[]
    for v in state['vouchers']:
        if v.get('opening') or v.get('closing') or v.get('reverses') or not effective(v) or v['id'] in owned or through and v['date']>through:continue
        if any(l['account'] in {a['code'] for a in bank_accounts(state)} for l in v['lines']):continue
        for i,l in enumerate(v['lines']):
            if l['account']==account and (scope=='ALL' or l['branch']==scope):
                key=v['id']+':'+str(i);remaining=l['debit']-l['credit']-used[key]
                if remaining:result.append(dict(key=key,voucher=v['id'],date=v['date'],branch=l['branch'],remaining=remaining,narration=v['narration'],account=account))
    return sorted(result,key=lambda b:(b['date'],b['key']))


def _split(total,weights):
    weights={k:v for k,v in weights.items() if v>0};den=sum(weights.values())
    if not total:return {k:0 for k in weights}
    result={k:total*v//den for k,v in weights.items()}
    left=total-sum(result.values())
    for key in sorted(weights,key=lambda k:(-(total*weights[k]%den),k))[:left]:result[key]+=1
    return result


def settle(state,bank_id,clearing_account,keys,fees_minor=0,gst_minor=0,gst_mode='None',partial_gross_minor=None):
    candidate=copy.deepcopy(state);bank=next(_canonical(b) for b in candidate['bank_lines'] if b['id']==bank_id)
    remaining,_=outstanding(candidate,bank['bank_account']);bank=next((b for b in remaining if b['id']==bank_id),None)
    if bank is None or bank['remaining']<=0:raise ValueError('Choose an unmatched incoming bank deposit.')
    items=[b for b in clearing_items(candidate,clearing_account,bank['branch'],bank['date']) if b['key'] in set(keys)]
    if len(items)!=len(set(keys)) or not items:raise ValueError('A receipt is missing, already settled or outside this statement scope/date.')
    if partial_gross_minor is not None:
        if any(b['remaining']<0 for b in items):raise ValueError('Partial gross allocation requires positive receipt lines; review refunds separately.')
        if type(partial_gross_minor)!=int or not 0<partial_gross_minor<=sum(b['remaining'] for b in items):raise ValueError('Invalid partial gross amount.')
        left=partial_gross_minor;trim=[]
        for b in items:
            take=min(left,b['remaining'])
            if take:trim.append(dict(b,remaining=take));left-=take
        items=trim
    gross=sum(b['remaining'] for b in items)
    if any(type(x)!=int or x<0 for x in (fees_minor,gst_minor)) or gst_minor>fees_minor:raise ValueError('Enter valid INR fees and eligible GST; GST cannot exceed total fees.')
    if gst_mode not in ('None','Intra-state','Inter-state') or gst_mode=='None' and gst_minor:raise ValueError('Choose the reviewed fee GST treatment.')
    if gross<=0 or gross-fees_minor!=bank['remaining']:raise ValueError('Selected gross receipts minus total fees must equal the unmatched bank deposit.')
    gross_branch=defaultdict(int)
    for b in items:gross_branch[b['branch']]+=b['remaining']
    for branch,total in gross_branch.items():
        available=core.balances(candidate,bank['date'],branch)[clearing_account]
        if total>0 and total>available:raise ValueError('Settlement exceeds the cash/clearing balance in '+branch+'. Review refunds, expenses and prior settlements.')
    fee_branch=_split(fees_minor,gross_branch);tax_branch=_split(gst_minor,fee_branch) if gst_minor else {};cgst_branch=_split(gst_minor//2,tax_branch) if gst_mode=='Intra-state' and gst_minor else {}
    lines=[]
    def add(account,branch,signed):
        if signed:lines.append(dict(account=account,branch=branch,debit=max(signed,0),credit=max(-signed,0)))
    for branch,g in gross_branch.items():
        fee=fee_branch.get(branch,0);tax=tax_branch.get(branch,0)
        add(bank['bank_account'],branch,g-fee);add(clearing_account,branch,-g);add('4560',branch,fee-tax)
        if tax:
            if gst_mode=='Inter-state':add('1220',branch,tax)
            else:add('1200',branch,cgst_branch.get(branch,0));add('1210',branch,tax-cgst_branch.get(branch,0))
    v=core.post(candidate,bank['date'],'Reviewed '+{'1000':'cash deposit','1110':'UPI settlement','1120':'card settlement'}[clearing_account]+' '+bank['reference'],lines,metadata={'type':'Transfer','reference':bank['reference'],'bank_account':bank['bank_account'],'source':'Bank settlement'})
    books=[dict(b,remaining=b['amount']) for b in book_items(candidate,bank['bank_account']) if b['voucher']==v['id']]
    record=_record_match(candidate,[bank],books,'settlement',v['id'],[dict(key=b['key'],amount=b['remaining'],account=clearing_account) for b in items])
    state.clear();state.update(candidate)
    return record


def adjust(state,bank_id,branch,contra,narration):
    candidate=copy.deepcopy(state);row=next(_canonical(b) for b in candidate['bank_lines'] if b['id']==bank_id)
    banks,_=outstanding(candidate,row['bank_account']);row=next((b for b in banks if b['id']==bank_id),None)
    if row is None:raise ValueError('The bank row is already fully matched.')
    if row['branch']!='ALL' and row['branch']!=branch:raise ValueError('Adjustment branch differs from the statement scope.')
    if contra==row['bank_account'] or contra in candidate.get('foreign_accounts',{}):raise ValueError('Choose a different INR contra ledger.')
    n=row['remaining'];lines=[dict(account=row['bank_account'],branch=branch,debit=max(n,0),credit=max(-n,0)),dict(account=contra,branch=branch,debit=max(-n,0),credit=max(n,0))]
    v=core.post(candidate,row['date'],narration,lines,metadata={'type':'Receipt' if n>0 else 'Payment','reference':row['reference'],'bank_account':row['bank_account'],'source':'Reviewed bank adjustment'})
    books=[dict(b,remaining=b['amount']) for b in book_items(candidate,row['bank_account']) if b['voucher']==v['id']]
    record=_record_match(candidate,[row],books,'adjustment',v['id']);state.clear();state.update(candidate)
    return record


def classify_exception(state,bank_id,reason):
    reason=reason.strip()
    if not reason:raise ValueError('Record the bank-error/dispute reason; ordinary charges should be posted as an adjustment.')
    row=next(_canonical(b) for b in state['bank_lines'] if b['id']==bank_id);_editable_bank(state,[row]);used,_=used_allocations(state);remaining=row['amount']-used[bank_id]
    if not remaining:raise ValueError('This row is already matched.')
    record=dict(reason=reason,amount=remaining,created_at=_now())
    old=copy.deepcopy(state.setdefault('bank_exceptions',{}).get(bank_id));state['bank_exceptions'][bank_id]=record
    core.add_audit(state,'BANK_EXCEPTION',{'bank_id':bank_id,'old':old,'new':copy.deepcopy(record)})


def remove_exception(state,bank_id,reason):
    if not reason.strip():raise ValueError('Enter a reason for removing the exception.')
    row=next(_canonical(b) for b in state['bank_lines'] if b['id']==bank_id);_editable_bank(state,[row])
    old=state.setdefault('bank_exceptions',{}).pop(bank_id,None)
    if old is None:raise ValueError('No exception exists for this row.')
    core.add_audit(state,'BANK_EXCEPTION_REMOVED',{'bank_id':bank_id,'old':old,'reason':reason})


def validate_post(state,day,lines,reverses):
    for l in lines:
        if any(day<=s['through'] for s in _finals(state,l['account'],l['branch'])):raise ValueError('This bank ledger is reconciled through that date. Reopen its statement before backdating a bank entry.')
    if reverses:
        for m in state['matches']:
            if not active(m):continue
            refs=[e['book_key'].rsplit(':',1)[0] for e in _edges(state,m)]+[a['key'].rsplit(':',1)[0] for a in m.get('clearing_allocations',[])]
            if reverses in refs or reverses in m.get('vouchers',[]) or reverses==m.get('posting'):raise ValueError('Undo the active bank match before reversing a matched journal/receipt.')


def undo_match(state,match_id,reason,reversal_date=None):
    if not reason.strip():raise ValueError('Enter an undo reason.')
    candidate=copy.deepcopy(state);m=next((m for m in candidate['matches'] if m.get('id',m.get('bank_id'))==match_id),None)
    if not m or not active(m):raise ValueError('Choose an active bank match.')
    rows=[_canonical(b) for b in candidate['bank_lines'] if b['id'] in {e['bank_id'] for e in _edges(candidate,m)}];_editable_bank(candidate,rows)
    undo_day=reversal_date or datetime.now(timezone(timedelta(hours=5,minutes=30))).date().isoformat();_day(undo_day)
    if undo_day<max(b['date'] for b in rows):raise ValueError('Undo date cannot precede the bank transaction.')
    m['status']='UNDONE';m['undo_reason']=reason;m['undone_at']=_now();m['undo_date']=undo_day
    if m.get('posting') and m.get('kind','settlement') in ('settlement','adjustment'):
        original=next(v for v in candidate['vouchers'] if v['id']==m['posting'])
        v=core.reverse(candidate,original['id'],undo_day);m['reversal']=v['id']
    _status(candidate);core.add_audit(candidate,'BANK_MATCH_UNDONE',copy.deepcopy(m));state.clear();state.update(candidate)


def statement(state,sid):return next(s for s in state.get('bank_statements',[]) if s['id']==sid)


def update_balances(state,sid,opening,closing):
    s=statement(state,sid)
    if s['status']!='OPEN':raise ValueError('Only an open statement can have its balances changed.')
    if not str(opening).strip() or not str(closing).strip():raise ValueError('Enter both bank-verified opening and closing balances.')
    rows={r['id']:_canonical(r) for r in state['bank_lines']};op,cl=amount(opening),amount(closing)
    coverage=_check_balances([rows[k] for k in s['row_ids']],op,cl);old=copy.deepcopy(s)
    s.update(opening=op,closing=cl,coverage=coverage);core.add_audit(state,'BANK_BALANCES_UPDATED',{'old':old,'new':copy.deepcopy(s)})


def brs(state,sid):
    core.verify_audit(state);s=statement(state,sid);account,scope,end=s['bank_account'],s['scope'],s['through']
    bu,ku=used_allocations(state,end)
    cancelled=_cancelled_book_keys(state,account,end)
    books=[dict(b,remaining=b['amount']-ku[b['key']]) for b in book_items(state,account,scope,end) if b['amount']!=ku[b['key']] and b['key'] not in cancelled]
    banks=[dict(b,remaining=b['amount']-bu[b['id']]) for b in bank_rows(state,account) if b['date']<=end and (scope=='ALL' or b['branch']==scope) and b['amount']!=bu[b['id']]]
    all_books={b['key']:b for b in book_items(state,account,scope)}
    for b in banks:
        ex=state.get('bank_exceptions',{}).get(b['id'])
        b['exception']=ex['reason'] if ex and ex['amount']==b['remaining'] else ''
        future=[e for m in state['matches'] if active(m,end) for e in _edges(state,m) if e['bank_id']==b['id'] and e['book_key'] in all_books and all_books[e['book_key']]['date']>end]
        if not b['exception'] and future and sum(e['amount'] for e in future)==b['remaining']:
            b['exception']='Timing: recorded in books after statement date ('+', '.join(sorted({all_books[e['book_key']]['voucher'] for e in future}))+')'
    ledger=core.balances(state,end,scope)[account]
    deposits=sum(b['remaining'] for b in books if b['remaining']>0);payments=-sum(b['remaining'] for b in books if b['remaining']<0)
    bank_credits=sum(b['remaining'] for b in banks if b['remaining']>0);bank_debits=-sum(b['remaining'] for b in banks if b['remaining']<0)
    adjusted=None if s['closing'] is None else s['closing']+deposits-payments-bank_credits+bank_debits
    variance=None if adjusted is None else ledger-adjusted
    verified=s['opening'] is not None and s['closing'] is not None
    unclassified=[b for b in banks if not b['exception']]
    ready=verified and variance==0 and not unclassified and s['status']!='VOID'
    return dict(entity=state['settings']['entity'],registration=copy.deepcopy(state['settings'].get('registration',{})),statement=copy.deepcopy(s),account_name=next(a['name'] for a in state['accounts'] if a['code']==account),scope_name='All operating units / shared bank' if scope=='ALL' else foundation.branches(state)[scope],books=books,banks=banks,
        bank_closing=s['closing'],deposits_in_transit=deposits,unpresented_payments=payments,unrecorded_bank_credits=bank_credits,unrecorded_bank_debits=bank_debits,adjusted_bank_balance=adjusted,book_balance=ledger,variance=variance,ready=ready,unclassified_count=len(unclassified),status=s['status'],prepared=datetime.now(timezone(timedelta(hours=5,minutes=30))).isoformat(timespec='seconds'))


def finalize(state,sid):
    report=brs(state,sid)
    if not report['ready']:raise ValueError('Reconciliation is incomplete: verify statement balances, resolve unexplained differences, and post or explicitly classify all outstanding bank rows.')
    s=statement(state,sid);s['status']='RECONCILED';s['reconciled_at']=_now();core.add_audit(state,'BANK_STATEMENT_RECONCILED',{'id':sid,'through':s['through'],'bank_account':s['bank_account'],'book_balance':report['book_balance'],'bank_balance':report['bank_closing']})


def reopen(state,sid,reason):
    if not reason.strip():raise ValueError('Enter a reason for reopening the statement.')
    s=statement(state,sid)
    if s['status']!='RECONCILED':raise ValueError('This statement is already open.')
    s['status']='OPEN';s['reopened_at']=_now();core.add_audit(state,'BANK_STATEMENT_REOPENED',{'id':sid,'reason':reason,'prior_reconciled_at':s.get('reconciled_at')})


def set_policy(state,tolerance,use_value_date,holiday_csv='',working_days=''):
    if str(tolerance) not in ('0','1'):raise ValueError('Date tolerance must be 0 or 1 business day.')
    candidate=copy.deepcopy(state);sources=[]
    if holiday_csv.strip():
        reader=csv.DictReader(io.StringIO(holiday_csv.lstrip('\ufeff')))
        if not {'date','name','source'}.issubset(reader.fieldnames or []):raise ValueError('Holiday CSV needs date,name,source.')
        for r in reader:
            _day(r['date'])
            if not r['source'].strip():raise ValueError('Record the holiday calendar source.')
            candidate['holidays']=sorted(set(candidate['holidays']+[r['date']]))
            sources.append(dict(date=r['date'],name=r['name'],source=r['source']))
    working=[d.strip() for d in working_days.split(',') if d.strip()]
    for d in working:_day(d)
    old=copy.deepcopy(candidate['settings'].get('bank_reconciliation',{}))
    policy=dict(tolerance=int(tolerance),use_value_date=bool(use_value_date),working_days=working,holiday_sources=old.get('holiday_sources',[])+sources)
    candidate['settings']['bank_reconciliation']=policy;core.add_audit(candidate,'BANK_CALENDAR_UPDATED',{'old':old,'new':copy.deepcopy(policy),'holidays':candidate['holidays']});state.clear();state.update(candidate)


def settlement_suggestions(state,bank_id,clearing_account):
    row=next(_canonical(b) for b in state['bank_lines'] if b['id']==bank_id)
    banks,_=outstanding(state,row['bank_account']);row=next((b for b in banks if b['id']==bank_id),None)
    if not row or row['remaining']<=0:return []
    tol=state['settings'].get('bank_reconciliation',{}).get('tolerance',1)
    items=clearing_items(state,clearing_account,row['branch'],row['date']);eligible=[]
    for item in items:
        expected=next_business_day(state,item['date']);latest=expected
        for _ in range(tol):latest=next_business_day(state,latest)
        if clearing_account=='1000' or expected<=row['date']<=latest:eligible.append(item)
    out=[]
    # Bounded contiguous batches; manual selection remains available for other groupings.
    for i in range(len(eligible)):
        total=0
        for j in range(i,min(i+200,len(eligible))):
            total+=eligible[j]['remaining'];fee=total-row['remaining']
            if total>0 and 0<=fee<=total//20 and (clearing_account!='1000' or fee==0):
                out.append(dict(keys=[x['key'] for x in eligible[i:j+1]],gross=total,proposed_fee=fee,from_date=eligible[i]['date'],through=eligible[j]['date']))
                if len(out)>=20:return out
    return out


def void_statement(state,sid,reason):
    if not reason.strip():raise ValueError('Enter a reason for cancelling the import.')
    candidate=copy.deepcopy(state);s=statement(candidate,sid)
    if s['status']!='OPEN':raise ValueError('Only an open statement can be voided; reopen reconciled statements first.')
    related=set(s['row_ids'])
    if any(active(m) and any(e['bank_id'] in related for e in _edges(candidate,m)) for m in candidate['matches']):raise ValueError('Undo the active matches before voiding this statement.')
    other_rows={k for other in candidate['bank_statements'] if other['id']!=sid and other['status']!='VOID' for k in other['row_ids']}
    for b in candidate['bank_lines']:
        if b['id'] in related and b['id'] not in other_rows and b.get('statement_id')==sid:
            _editable_bank(candidate,[_canonical(b)]);b['void']=True;b['status']='VOID';b['void_reason']=reason;b['voided_at']=_now()
    s['status']='VOID';s['void_reason']=reason;s['voided_at']=_now();core.add_audit(candidate,'BANK_STATEMENT_VOIDED',{'id':sid,'reason':reason})
    state.clear();state.update(candidate)
