"""Standard, row-reviewed CSV/XLSX sales receipts; no journal posts at upload."""
import copy
import csv
import hashlib
import io
from pathlib import Path
import re
import uuid
from datetime import date,datetime
from . import core,foundation,revenue,sales_import,bulk_import,financial_reports as f

HEADERS=('receipt_date','branch','receipt_no','invoice_no','party','record_type','tax_mode',*[m[0] for m in revenue.MODES],'cash','bank','upi','card','other','total')
PAYMENTS=('cash','bank','upi','card','other')


def identity(row):return (row['branch'],core.financial_year(row['date']),re.sub('[^a-zA-Z0-9]','',row['receipt']).upper())
def _batch(s,key):
    q=next((q for q in s['queue'] if q['id']==key and q.get('kind')=='standard_sales'),None)
    if not q:raise ValueError('Select a standard sales receipt import.')
    return q

def known(s,exclude=None):
    keys=set()
    for q in s['queue']:
        if q['id']==exclude or q.get('kind') not in ('daily_sales','standard_sales'):continue
        for i,r in enumerate(q['rows']):
            status=r.get('review') if q.get('kind')=='standard_sales' else q.get('row_reviews',{}).get(str(i),{}).get('status')
            if r.get('validation') in ('INVALID','DUPLICATE') or status in ('REJECTED','INVALID','DUPLICATE') or not r.get('receipt') or not r.get('date'):continue
            keys.add(identity(r))
    return keys


def lines(s,r):
    payments={k:v for k,v in r['payments'].items() if k!='bank'};revenues=r['revenues'];bank=r['payments']['bank']
    # Receipt taxonomy remains unchanged; a bank column is an explicit 1010 receipt,
    # never guessed from legacy Other source clearing.
    if bank:
        payments['other']=payments.get('other',0)+bank
        result=revenue.receipt_lines(s,r['branch'],revenues,payments)
        other=s['revenue_mode_accounts']['other_receipts'];entry=next(l for l in result if l['account']==other and l['debit'])
        entry['debit']-=bank
        if not entry['debit']:result.remove(entry)
        result.insert(0,dict(account='1010',branch=r['branch'],debit=bank,credit=0));return result
    return revenue.receipt_lines(s,r['branch'],revenues,payments)


def _parse(s,raw,source_row):
    for value in raw.values():
        if isinstance(value,str) and value.startswith('='):raise ValueError('Use literal values, not spreadsheet formulas. Paste values from the source report.')
    d=raw['receipt_date']
    if isinstance(d,datetime):d=d.date().isoformat()
    elif isinstance(d,date):d=d.isoformat()
    else:d=str(d or '').strip()
    if date.fromisoformat(d).isoformat()!=d:raise ValueError('Receipt date must be YYYY-MM-DD or an actual Excel date.')
    branch=str(raw['branch'] or '').strip();receipt=str(raw['receipt_no'] or '').strip()
    if branch not in foundation.branches(s,True) or not receipt or not re.search('[a-zA-Z0-9]',receipt):raise ValueError('Use an active unit code and a unique receipt number.')
    if str(raw['record_type'] or '').strip().lower()!='earned':raise ValueError('This template supports earned receipts only. Advances, refunds and separate bill collections use Service revenue.')
    if str(raw['tax_mode'] or '').strip().lower()!='no_separate_tax':raise ValueError('Review tax treatment. This collection template does not calculate separate GST/TDS; taxable source splits use GST/Service revenue.')
    revenues={key:sales_import.amount(raw[key],key) for key,*_ in revenue.MODES};payments={k:sales_import.amount(raw[k],k) for k in PAYMENTS};amount=sum(revenues.values())
    if amount<=0 or amount!=sum(payments.values()):raise ValueError('All 15 revenue columns must equal actual payment splits, with a positive total.')
    if raw['total'] not in (None,'') and sales_import.amount(raw['total'],'total')!=amount:raise ValueError('Total disagrees with the independently summed revenue/payment amounts.')
    r=dict(date=d,branch=branch,receipt=receipt,invoice=str(raw['invoice_no'] or '').strip(),party=str(raw['party'] or '').strip(),revenues=revenues,payments=payments,amount=amount,source_row=source_row,sheet='Receipts');r['key']=sales_import.row_key(r);return r


def stage(s,raw,filename):
    core.verify_audit(s);digest=hashlib.sha256(raw).hexdigest()
    if any(q.get('digest')==digest for q in s['queue']):raise ValueError('This exact report is already queued. Review its existing batch.')
    headers,rows,label=bulk_import._read(raw,filename,'Receipts')
    if set(headers)!=set(HEADERS):raise ValueError('Use the standard receipt template headings exactly. Missing: '+', '.join(sorted(set(HEADERS)-set(headers)))+'; unknown: '+', '.join(sorted(set(headers)-set(HEADERS))))
    existing=known(s);seen=set();parsed=[];probe=copy.deepcopy(s)
    for n,row in enumerate(rows,2):
        if not any(v not in (None,'') for v in row):continue
        rawrow=dict(zip(headers,list(row)+[None]*len(headers)));record=dict(source_row=n,sheet=label,raw={k:('' if v is None else v.isoformat() if isinstance(v,(date,datetime)) else str(v)) for k,v in rawrow.items()},review='PENDING',reason='',error='',validation='VALID')
        try:
            if len(row)>len(headers) and any(v not in (None,'') for v in row[len(headers):]):raise ValueError('Source row has extra fields beyond the standard headings.')
            r=_parse(s,rawrow,n);record.update(r);key=identity(r)
            if key in existing or key in seen:record.update(validation='DUPLICATE',error='Receipt number already queued/posted in this unit/fiscal year, or repeated in this file.')
            else:core.post(probe,r['date'],'Validate receipt '+r['receipt'],lines(probe,r));seen.add(key)
        except (ValueError,KeyError,TypeError) as error:record.update(validation='INVALID',error=str(error))
        parsed.append(record)
    if not parsed:raise ValueError('The report has no receipt data rows.')
    q=dict(id=str(uuid.uuid4()),digest=digest,filename=Path(filename).name,kind='standard_sales',status='PENDING_REVIEW',rows=parsed,branch='ALL',format_version=1)
    s['queue'].append(q);core.add_audit(s,'STANDARD_RECEIPTS_STAGED',dict(id=q['id'],filename=q['filename'],rows=len(parsed),valid=sum(r['validation']=='VALID' for r in parsed)));return q


def review(s,key,indices,approved,reason,confirmed=False):
    q=_batch(s,key)
    if q['status']!='PENDING_REVIEW' or not reason.strip() or type(approved)!=bool or not confirmed:raise ValueError('Review pending source rows with a reason and confirm earned-service/tax/payment treatment.')
    if not indices or len(indices)!=len(set(indices)) or any(type(i)!=int or not 0<=i<len(q['rows']) for i in indices):raise ValueError('Select actual source rows.')
    if approved and any(q['rows'][i]['validation']!='VALID' for i in indices):raise ValueError('Invalid/duplicate rows cannot be approved. Reject them with a reason and upload a corrected report.')
    changes=[dict(index=i,status='APPROVED' if approved else 'REJECTED',reason=reason.strip()) for i in indices]
    for r in changes:q['rows'][r['index']].update(review=r['status'],reason=r['reason'])
    core.add_audit(s,'STANDARD_RECEIPTS_REVIEWED',dict(batch=key,rows=changes))


def post(s,key):
    core.verify_audit(s);c=copy.deepcopy(s);q=_batch(c,key)
    if q['status']!='PENDING_REVIEW' or any(r['review'] not in ('APPROVED','REJECTED') for r in q['rows']):raise ValueError('Approve or reject every source row before atomic batch posting.')
    selected=[r for r in q['rows'] if r['review']=='APPROVED']
    if not selected:raise ValueError('There are no approved receipts to post.')
    existing=known(c,key);seen=set()
    for r in selected:
        if r['validation']!='VALID' or identity(r) in existing or identity(r) in seen:raise ValueError('A selected receipt is invalid or now duplicated. Review sources again.')
        seen.add(identity(r));core.post(c,r['date'],'Sales receipt '+r['receipt']+' | '+q['filename']+' row '+str(r['source_row']),lines(c,r),metadata={'type':'Receipt','reference':r['receipt'],'payee':r['party']},source={'kind':'daily_sales','document_id':q['id'],'record_id':r['key'],'row_count':1})
    q['status']='PROCESSED';core.add_audit(c,'STANDARD_RECEIPTS_POSTED',dict(id=key,posted=len(selected),rejected=len(q['rows'])-len(selected)));s.clear();s.update(c);return len(selected)


def snapshot(s,key):
    q=_batch(s,key);valid=[r for r in q['rows'] if r['validation']=='VALID'];start=min([r['date'] for r in valid] or ['2026-01-01']);end=max([r['date'] for r in valid] or [start]);p=f.snapshot(s,'Trial balance',start,end)
    p.update(mode='Sales receipt import review',columns=['Source row','Date','Unit','Receipt','Revenue INR','Payments INR','Validation','Decision','Error / review reason'],money_columns=[4,5],rows=[[r['source_row'],r.get('date',''),r.get('branch',''),r.get('receipt',''),r.get('amount',0),sum(r.get('payments',{}).values()),r['validation'],r['review'],r['error']+' '+r['reason']] for r in q['rows']],controls={},note=q['filename']+' | '+q['status']+' | Source SHA256 '+q['digest']+'. Every approved row becomes one balanced receipt journal; rejected rows never post. All source rows remain in the review history. No GST, TDS, advance or refund is guessed.');return p


def template(path):
    path=Path(path)
    if path.suffix.lower()=='.csv':
        with path.open('w',encoding='utf-8-sig',newline='') as out:csv.writer(out).writerow(HEADERS)
    elif path.suffix.lower()=='.xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
        wb=Workbook();ws=wb.active;ws.title='Receipts';ws.append(list(HEADERS));ws.freeze_panes='D2';ws.auto_filter.ref=f'A1:{get_column_letter(len(HEADERS))}1'
        for cell in ws[1]:cell.font=Font(bold=True);cell.fill=PatternFill('solid',fgColor='E3EEDF');cell.alignment=Alignment(wrap_text=True)
        for i,h in enumerate(HEADERS,1):ws.column_dimensions[get_column_letter(i)].width=23 if i<8 else 18
        ws.row_dimensions[1].height=34
        for col in ('C','D'):ws.column_dimensions[col].number_format='@'
        for col,formula in [('F','"earned"'),('G','"no_separate_tax"')]:
            dv=DataValidation(type='list',formula1=formula);ws.add_data_validation(dv);dv.add(f'{col}2:{col}100001')
        guide=wb.create_sheet('Instructions')
        for text in ['Riju standard sales receipt template v1 — no sample receipt rows to accidentally post.','One earned receipt per row. Worksheet must be named Receipts; preserve all headings. Columns may be reordered.','receipt_date: YYYY-MM-DD or real Excel date. branch: actual unit code, e.g. JP, KM, VK.','receipt_no: required, unique per unit/fiscal year. Preserve leading zeros as text. invoice_no and party optional.','record_type = earned; tax_mode = no_separate_tax, after source/tax review.','15 service columns use INR amounts, at most two decimals. Blank money cells mean zero.','cash/bank/upi/card/other are actual splits. bank debits actual collections to ledger 1010; other remains Other receipts clearing.','Service total must equal payment total; optional total, if filled, must match both. Paste values, not formulas.','Advances/refunds/unpaid bills and taxable receipt splits use Service revenue / GST workflows. Do not import them as earned receipts.','Example (do not post): 2026-11-01, JP, SAMPLE-001; consultation 500; treatment 300; cash 200; upi 600; total 800.','Upload from Sales receipt → Import sales receipt report. Review valid rows in bulk, reject flagged rows, then preview and approve batch.','Duplicates are checked across queued/posted standard and legacy sales reports, including files renamed or edited. Manual/unlinked receipts still require source review.']:
            guide.append([text]);guide.cell(guide.max_row,1).alignment=Alignment(wrap_text=True);guide.row_dimensions[guide.max_row].height=40
        guide.column_dimensions['A'].width=125;wb.save(path);wb.close()
    else:raise ValueError('Save the standard template as CSV or XLSX.')


def totals_snapshot(s,key):
    q=_batch(s,key);p=snapshot(s,key);rows=[]
    for source,keys in [('Revenue',[m[0] for m in revenue.MODES]),('Payment',list(PAYMENTS))]:
        for k in keys:
            field='revenues' if source=='Revenue' else 'payments';valid=sum(r.get(field,{}).get(k,0) for r in q['rows'] if r['validation']=='VALID');approved=sum(r.get(field,{}).get(k,0) for r in q['rows'] if r['review']=='APPROVED');rows.append([source,k,valid,approved])
    p.update(mode='Receipt category payment totals',columns=['Side','Category / method','Valid source INR','Approved to post INR'],money_columns=[2,3],rows=rows,note=p['note']+' Totals exclude invalid/duplicate rows; approved columns exclude all rejected/pending rows. Compare both revenue and payment totals to the source and final posting decisions.');return p
