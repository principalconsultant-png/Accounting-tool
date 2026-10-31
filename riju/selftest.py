"""Packaged executable readiness check; operates only in a disposable directory."""
import io
import json
import os
from pathlib import Path
import tempfile
from datetime import datetime
from . import APP_VERSION,core,foundation
from .revenue import MODES
from .sales_import import stage_sales,approve_sales
from .storage import Vault

def run(output):
    result={'version':APP_VERSION,'passed':False,'platform':os.name}
    try:
        from openpyxl import Workbook
        w=Workbook();s=w.active
        for col,(_,label,*_) in enumerate(MODES,8):s.cell(4,col,label)
        for col,label in enumerate(('Sub total','Card','UPI','Cash','Other source','Total'),23):s.cell(4,col,label)
        s.cell(5,1,1);s.cell(5,2,datetime(2026,6,1));s.cell(5,6,'TEST-INVOICE');s.cell(5,7,'TEST-RECEIPT')
        for col in range(8,23):s.cell(5,col,100)
        s.cell(5,23,1500);s.cell(5,24,500);s.cell(5,25,500);s.cell(5,26,500);s.cell(5,27,0);s.cell(5,28,1500)
        raw=io.BytesIO();w.save(raw);w.close()
        with tempfile.TemporaryDirectory() as directory:
            v=Vault(Path(directory)/'check.riju');state=v.load()
            q=stage_sales(state,raw.getvalue(),'synthetic.xlsx','JP');approve_sales(state,q['id'])
            v.save(state);loaded=v.load();core.verify_audit(loaded)
            if state!=loaded or len(loaded['vouchers'])!=1:raise ValueError('Vault round-trip mismatch')
            if sum(core.balances(loaded,'2026-06-30').values())!=0:raise ValueError('Unbalanced report')
            folder=Path(directory)/'external';folder.mkdir();dest=v.backup(folder)
            if Vault(dest).load()!=loaded:raise ValueError('Backup restore mismatch')
            foundation.save_entity(loaded,'Synthetic hospital',{})
            foundation.save_branch(loaded,'NEW',{'name':'Synthetic branch','beds':0})
            foundation.save_account(loaded,'1011','Synthetic bank','Asset','Cash and bank')
            lines=[dict(account='1011',branch='NEW',debit=10000,credit=0),dict(account='3000',branch='NEW',debit=0,credit=10000)]
            draft=foundation.save_draft(loaded,'2024-04-01','Synthetic opening',lines,{'type':'Journal'},opening=True)
            foundation.post_draft(loaded,draft['id']);v.save(loaded)
            if v.load()!=loaded:raise ValueError('Foundation vault round-trip mismatch')
            # A real previous-format encrypted vault must be preserved before migration.
            from .storage import MAGIC,protect
            legacy=core.empty_state();legacy['version']=1
            for key in ('schema_revision','branches','account_details','financial_years','drafts','patient_control'):legacy.pop(key,None)
            old=Vault(Path(directory)/'legacy.riju')
            original=MAGIC+protect(json.dumps(legacy).encode());old.path.write_bytes(original)
            upgraded=old.load();upgraded['settings']['backup_dir']=str(folder);old.save(upgraded)
            if old.last_upgrade_backup.read_bytes()!=original:raise ValueError('Upgrade backup differs from original vault')
            if old.load()!=upgraded:raise ValueError('Upgrade round-trip mismatch')
            result['foundation_checks_passed']=True
            from . import ledgers
            pack=ledgers.ledger_pack(loaded,['1011','3000'],'2024-04-01','2026-06-30','NEW')
            if pack['books'][0]['opening']!=10000 or pack['books'][0]['closing']!=10000:
                raise ValueError('Ledger opening/closing mismatch')
            root=Path(directory)
            ledgers.export_pdf(pack,root/'ledgers.pdf')
            ledgers.export_xlsx(pack,root/'ledgers.xlsx')
            ledgers.export_csv(pack,root/'ledgers.csv')
            ledgers.export_html(pack,root/'ledgers.html')
            if not (root/'ledgers.pdf').read_bytes().startswith(b'%PDF-'):raise ValueError('Packaged PDF export failed')
            from openpyxl import load_workbook
            workbook=load_workbook(root/'ledgers.xlsx',read_only=True)
            try:
                if workbook.sheetnames!=['1011','3000']:raise ValueError('Packaged Excel ledger pack mismatch')
            finally:workbook.close()
            result['ledger_exports_passed']=True
            from . import fx,activity
            fx.save_rate(loaded,'2026-06-01','USD','83.25','Synthetic bank advice')
            receipt=fx.post_receipt(loaded,payment_date='2026-06-01',code='USD',foreign_amount='20',credit_account='2300',reference='SELFTEST-FX-1')
            if receipt['fx']['inr_minor']!=166500:raise ValueError('Packaged FX conversion mismatch')
            v.save(loaded)
            if v.load()!=loaded:raise ValueError('Packaged FX vault round-trip mismatch')
            metrics=activity.report(loaded)['metrics']
            if metrics['fx_receipts_posted']!=1 or metrics['journal_vouchers_posted']!=3:raise ValueError('Packaged work-volume mismatch')
            result['fx_and_work_volume_passed']=True
            from . import reconciliation as bank,brs_exports
            receipt=core.post(loaded,'2026-06-02','Synthetic bank receipt',[
                dict(account='1011',branch='NEW',debit=5000,credit=0),
                dict(account='4000',branch='NEW',debit=0,credit=5000)])
            statement=bank.import_statement(loaded,b'date,narration,reference,amount\n2026-06-02,Receipt,SELFTEST-BANK,50\n','selftest-bank.csv','1011','NEW','2026-06-01','2026-06-30','100','150')
            bank.match_existing(loaded,statement['row_ids'],[receipt['id']+':0'])
            report=bank.brs(loaded,statement['id'])
            if not report['ready']:raise ValueError('Packaged bank reconciliation proof failed')
            brs_exports.export_pdf(report,root/'brs.pdf')
            if not (root/'brs.pdf').read_bytes().startswith(b'%PDF-'):raise ValueError('Packaged BRS export failed')
            bank.finalize(loaded,statement['id']);v.save(loaded)
            if v.load()!=loaded:raise ValueError('Packaged bank reconciliation vault round-trip mismatch')
            result['bank_reconciliation_passed']=True
            from . import history
            if not history.verify_saved_state(loaded) or not history.integrity_report(loaded)['passed']:
                raise ValueError('Packaged history integrity review failed')
            trace=history.trace_voucher(loaded,receipt['id'])
            if len(trace['matches'])!=1:raise ValueError('Packaged bank history trace failed')
            history.export_rows(history.entries(loaded),root/'history.csv')
            history.export_rows(history.entries(loaded),root/'history.json','json')
            result['history_integrity_passed']=True
            from . import tax
            tax.save_rule(loaded,dict(family='GST',key='SELFTEST-GST',label='Synthetic taxable service',effective='2026-04-01',rate='18',treatment='Taxable',hsn_sac='999311',applicability='Invented test rule only',source='Packaged synthetic test',reviewed=True))
            party=tax.save_party(loaded,dict(name='Synthetic supplier',effective='2026-04-01',category='Unknown'))
            preview=tax.preview(loaded,'SELFTEST-GST','2026-06-01',party['party_id'])
            if preview['rule']['rate']!='18':raise ValueError('Packaged tax configuration preview failed')
            tax.export_config(loaded,root/'tax-config.json');v.save(loaded)
            if v.load()!=loaded:raise ValueError('Packaged tax configuration vault round-trip mismatch')
            result['tax_configuration_passed']=True
            from . import gst
            gst.post(loaded,dict(date='2026-06-03',branch='JP',kind='Sale',registration='29ABCDE1234F1Z5',place_of_supply='29',counterparty_gstin='',reference='SELFTEST-GST-1',rule_key='SELFTEST-GST',amount='100',inclusive=False,account='4000',counter_account='1010',itc='Not applicable',narration='Synthetic GST sale',confirmed=True))
            working=gst.summary(loaded,'2026-06-01','2026-06-30')
            if working['totals']['output_cgst']!=900:raise ValueError('Packaged GST journal/summary failed')
            gst.export_csv(working,root/'gst-working.csv');gst.export_report(working,root/'gst-working.json');v.save(loaded)
            if v.load()!=loaded:raise ValueError('Packaged GST vault round-trip mismatch')
            result['gst_accounting_passed']=True
            original=loaded['gst_documents'][-1]
            gst.post_note(loaded,dict(original_document_id=original['id'],date='2026-06-04',document_type='Credit note',amount='25',reference='SELFTEST-CN-1',counter_account='1010',narration='Synthetic partial refund',confirmed=True))
            corrected=gst.summary(loaded,'2026-06-01','2026-06-30')
            if corrected['totals']['output_cgst']!=675 or any(corrected['unclassified'].values()):raise ValueError('Packaged GST note proof failed')
            v.save(loaded)
            if v.load()!=loaded:raise ValueError('Packaged GST note persistence failed')
            result['gst_notes_passed']=True
            purchase=gst.post(loaded,dict(date='2026-06-05',branch='JP',kind='Purchase',registration='29ABCDE1234F1Z5',place_of_supply='29',counterparty_gstin='29ZZZZZ1234Z1Z5',reference='SELFTEST-P-1',rule_key='SELFTEST-GST',amount='100',inclusive=False,account='4590',counter_account='1010',itc='Eligible',narration='Synthetic purchase',confirmed=True))
            credit=gst.post_note(loaded,dict(original_document_id=purchase['id'],date='2026-06-06',document_type='Credit note',amount='25',reference='SELFTEST-PC-1',counter_account='2000',narration='Synthetic purchase correction',confirmed=True))
            source=b'supplier_gstin,invoice,date,taxable,cgst,sgst,igst,itc_available,document_type,original_invoice,original_date\n29ZZZZZ1234Z1Z5,SELFTEST-PC-1,2026-06-06,25,2.25,2.25,0,Y,Credit note,SELFTEST-P-1,2026-06-05\n'
            statement=gst.import_2b(loaded,source,'selftest-2b.csv','29ABCDE1234F1Z5','2026-06')
            gst.match_2b(loaded,statement['id'],statement['rows'][0]['id'],credit['id'])
            review=gst.review_2b(loaded,statement['id'])
            if review['matched_signed_totals']['cgst']!=-225:raise ValueError('Packaged 2B note sign/link failed')
            gst.export_2b_csv(review,root/'2b-notes.csv');v.save(loaded)
            if v.load()!=loaded:raise ValueError('Packaged 2B note persistence failed')
            result['gst_2b_notes_passed']=True
            from . import rcm
            reverse=rcm.post(loaded,'Purchase',dict(date='2026-06-07',branch='JP',registration='29ABCDE1234F1Z5',reference='SELFTEST-RCM-1',narration='Synthetic RCM purchase',rule_key='SELFTEST-GST',rcm_basis='Invented test applicability only',supplier_name='Synthetic unregistered supplier',supplier_gstin='',supplier_state='29',place_of_supply='29',amount='100',account='4590',counter_account='2000',claim_expected=True,confirmed=True))
            rcm.post(loaded,'Deposit',dict(date='2026-06-08',branch='JP',registration='29ABCDE1234F1Z5',reference='SELFTEST-CIN-1',narration='Synthetic electronic cash funding',bank_account='1010',cgst='9',sgst='9',igst='0',confirmed=True))
            for stage,on,reference in [('Discharge','2026-06-09','SELFTEST-UTIL-1'),('Claim','2026-06-10','SELFTEST-RCM-ITC-1')]:
                rcm.post(loaded,stage,dict(date=on,purchase_id=reverse['id'],reference=reference,narration='Reviewed synthetic '+stage,eligibility_reviewed=True,confirmed=True))
            review=rcm.report(loaded,'2026-06-01','2026-06-30')
            if any(review['unclassified'].values()) or review['outstanding'][0]['itc_pending']:raise ValueError('Packaged RCM stage/proof failed')
            rcm.export(review,root/'rcm-working.json');v.save(loaded)
            if v.load()!=loaded:raise ValueError('Packaged RCM persistence failed')
            result['gst_rcm_passed']=True
            from . import financial_reports
            statement=financial_reports.snapshot(loaded,'Balance sheet','2026-04-01','2027-03-31')
            if any(statement['controls'].values()):raise ValueError('Packaged financial statement proof failed')
            financial_reports.export_xlsx(statement,root/'financial-report.xlsx')
            financial_reports.export_pdf(statement,root/'financial-report.pdf')
            workbook=load_workbook(root/'financial-report.xlsx',read_only=True)
            try:
                if workbook.active['A2'].value!='Balance sheet':raise ValueError('Packaged financial workbook mismatch')
            finally:workbook.close()
            if not (root/'financial-report.pdf').read_bytes().startswith(b'%PDF-'):raise ValueError('Packaged financial PDF export failed')
            result['financial_report_exports_passed']=True
            reports=financial_reports.report_pack(loaded,'2026-04-01','2027-03-31')
            financial_reports.export_pack_xlsx(reports,root/'financial-pack.xlsx')
            financial_reports.export_pdf(reports,root/'financial-pack.pdf')
            workbook=load_workbook(root/'financial-pack.xlsx',read_only=True)
            try:
                if workbook.sheetnames!=['Contents',*financial_reports.MODES]:raise ValueError('Packaged financial pack mismatch')
            finally:workbook.close()
            result['financial_report_pack_passed']=True
            from . import tds
            payee=tax.save_party(loaded,dict(name='Synthetic TDS company',effective='2026-04-01',category='Company',pan='ABCDE1234F'))
            tax.save_rule(loaded,dict(family='TDS',key='SELFTEST-J',label='Synthetic TDS rule',effective='2026-04-01',section='194J',rate='10',treatment='Withholding',party_category='Company',annual_threshold='1000',threshold_basis='Financial-year aggregate',applicability='Invented test only',source='Packaged test',reviewed=True))
            for day,reference in [('2026-06-11','SELFTEST-TDS-1'),('2026-06-12','SELFTEST-TDS-2')]:
                deduction=tds.post(loaded,'Expense',dict(date=day,branch='JP',reference=reference,narration='Synthetic professional fees',rule_key='SELFTEST-J',party_id=payee['party_id'],mode='Automatic',gross='600',base='600',account='4510',counter_account='2000',resident_reviewed=True,confirmed=True))
            if deduction['withheld']!=12000:raise ValueError('Packaged TDS catch-up failed')
            payment=tds.post(loaded,'Challan',dict(date='2026-06-13',branch='JP',reference='SELFTEST-TDS-UTR',narration='Synthetic TDS payment',bank_account='1010',rows=[dict(section='194J',branch='JP',year=2026,tax='120')],confirmed=True))
            tds.allocate(loaded,deduction['id'],payment['id'],0,'120','2026-06-13')
            tds.challan_details(loaded,payment['id'],dict(bsr='0123456',serial='00001',deposit_date='2026-06-13',verified=True))
            working=tds.report(loaded,'2026-04-01','2027-03-31')
            if working['outstanding']['rows'] or any(working['deductions']['controls'].values()):raise ValueError('Packaged TDS reconciliation failed')
            tds.export(loaded,'2026-04-01','2027-03-31','ALL',root/'tds.xlsx','xlsx');tds.export(loaded,'2026-04-01','2027-03-31','ALL',root/'tds.pdf','pdf');v.save(loaded)
            if v.load()!=loaded:raise ValueError('Packaged TDS persistence failed')
            result['tds_bookkeeping_passed']=True
            from . import cash_flow
            cfbook=core.empty_state()
            def cfpost(dr,cr,amount,opening=False):
                return core.post(cfbook,'2026-04-01','Synthetic cash-flow check',[dict(account=dr,branch='JP',debit=amount,credit=0),dict(account=cr,branch='JP',debit=0,credit=amount)],opening=opening)
            cfpost('1010','3000',100000,True)
            cfpost('1010','4000',20000);cfpost('1030','1010',10000);cfpost('1010','2400',30000);cfpost('4570','1050',2000)
            cash_flow.save_policy(cfbook,dict(cash_flow.defaults(cfbook),effective='2026-04-01',note='Synthetic reviewed policy only',reviewed=True))
            cfresult=cash_flow.analyse(cfbook,'2026-04-01','2026-04-30')
            if not cfresult['ready'] or cfresult['totals']['Investing']!=-10000 or cfresult['totals']['Financing']!=30000:raise ValueError('Packaged cash-flow classification/proof failed')
            for kind in ('xlsx','pdf'):cash_flow.export(cfbook,'2026-04-01','2026-04-30','ALL',root/('cash-flow.'+kind),kind)
            cfvault=Vault(root/'cash-flow.riju');cfvault.save(cfbook)
            if cfvault.load()!=cfbook:raise ValueError('Packaged cash-flow policy persistence failed')
            result['cash_flow_passed']=True
            from . import bank_templates
            templatebook=core.empty_state();bankraw=b'Booked,Movement,Ref\n02/04/2026,100,TEST-TEMPLATE-1\n'
            template=bank_templates.save(templatebook,bankraw,'bank.csv','1010','ALL','Synthetic packaged layout',mapping={'date':'Booked','amount':'Movement','reference':'Ref'},date_format='DD/MM/YYYY')
            from . import reconciliation
            imported=reconciliation.import_statement(templatebook,bankraw,'bank.csv','1010','ALL','2026-04-01','2026-04-30','0','100',template_id=template['id'])
            if imported['mapping_template']['id']!=template['id'] or len(templatebook['bank_lines'])!=1 or templatebook['vouchers']:raise ValueError('Packaged saved bank layout/provenance failed')
            templatevault=Vault(root/'bank-template.riju');templatevault.save(templatebook)
            if templatevault.load()!=templatebook:raise ValueError('Packaged saved bank layout persistence failed')
            result['bank_mapping_templates_passed']=True
            from . import fixed_assets
            assetbook=core.empty_state()
            registered=fixed_assets.register(assetbook,dict(reference='PACKAGED-ASSET',name='Synthetic medical equipment',branch='JP',category='Manual / specific Schedule II category',mode='Post purchase',acquired='2026-04-01',put_to_use='2026-04-01',baseline='2026-04-01',cost='1200',opening_accum='0',residual='0',months=12,basis='Synthetic reviewed book SLM and source only',reviewed=True))
            fixed_assets.post_depreciation(assetbook,'2026-04-30')
            fixed_assets.transfer(assetbook,registered['id'],'2026-05-01','KM','Synthetic approved transfer',True)
            for kind in ('xlsx','pdf'):fixed_assets.export(assetbook,'2026-04-01','2026-04-30','ALL',root/('fixed-assets.'+kind),kind)
            fixed_assets.dispose(assetbook,registered['id'],'2026-05-31','1050','1010','','','Synthetic disposal and separate tax review',True)
            if not fixed_assets.report(assetbook,'2026-04-01','2026-05-31')['ready']:raise ValueError('Packaged fixed-asset ledger/depreciation proof failed')
            avault=Vault(root/'fixed-assets.riju');avault.save(assetbook)
            if avault.load()!=assetbook:raise ValueError('Packaged fixed-asset persistence failed')
            result['book_fixed_assets_passed']=True
            from . import tax_depreciation
            taxbook=core.empty_state()
            taxblock=tax_depreciation.save_block(taxbook,'PACKAGED-TAX','Synthetic tax block',2024,{'JP':{'wdv':'1000','count':2}},'Synthetic reviewed opening tax WDV',True)
            tax_depreciation.save_rule(taxbook,taxblock['id'],2024,'15','Synthetic reviewed rate/regime',True)
            tax_depreciation.add_asset(taxbook,taxblock['id'],'TAX-ADD','JP','2024-12-01','500','Synthetic reviewed tax cost',True,put_to_use='2024-12-01')
            taxworking=tax_depreciation.calculate(taxbook,2024)
            if taxworking['normal']!=18750 or not taxworking['ready'] or taxbook['vouchers']:raise ValueError('Packaged tax WDV/short-use/no-journal proof failed')
            tax_depreciation.finalize(taxbook,2024,'Synthetic reviewed ended tax year',True)
            if tax_depreciation.calculate(taxbook,2025)['blocks'][0]['opening']!=131250:raise ValueError('Packaged tax carry-forward failed')
            for kind in ('xlsx','pdf'):tax_depreciation.export(taxbook,2024,'ALL',root/('income-tax-depreciation.'+kind),kind)
            tvault=Vault(root/'tax-depreciation.riju');tvault.save(taxbook)
            if tvault.load()!=taxbook:raise ValueError('Packaged tax depreciation persistence failed')
            result['tax_depreciation_passed']=True
            result.update(passed=True,revenue_modes=len(MODES),posted_vouchers=len(loaded['vouchers']),windows_vault_backup_roundtrip=True)
    except Exception as error:
        result['error']=str(error)
    Path(output).write_text(json.dumps(result,indent=2),encoding='utf-8')
    return result['passed']
