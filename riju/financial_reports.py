"""Shared read-only working-report snapshots for screen, Excel, PDF and CSV."""
import copy
import csv
from datetime import date,datetime,timedelta,timezone
from decimal import Decimal
import html
import os
from pathlib import Path
from . import core,foundation,ledgers

MODES=('Trial balance','Profit and loss','Balance sheet','Comparative P&L','Comparative balance sheet','Branch matrix','General ledger','Cash flow - direct','Cash flow - indirect','Cash flow - AP forecast')


def previous_year(day):
    from calendar import monthrange
    d=date.fromisoformat(day);return date(d.year-1,d.month,min(d.day,monthrange(d.year-1,d.month)[1])).isoformat()


def snapshot(state,mode,start,end,branch='ALL'):
    core.verify_audit(state)
    if mode not in MODES:raise ValueError('Choose a supported report.')
    if date.fromisoformat(start).isoformat()!=start or date.fromisoformat(end).isoformat()!=end or start>end:raise ValueError('Choose a valid YYYY-MM-DD reporting period.')
    names=foundation.branches(state)
    if branch!='ALL' and branch not in names:raise ValueError('Select a valid branch.')
    if mode in ('Cash flow - direct','Cash flow - indirect','Cash flow - AP forecast'):
        from . import cash_flow
        return cash_flow.snapshot(state,mode,start,end,branch)
    effective='ALL' if mode=='Branch matrix' else branch
    pack=dict(entity=state['settings']['entity'],registration=copy.deepcopy(state['settings'].get('registration',{})),mode=mode,start=start,end=end,branch=effective,scope='All operating units' if effective=='ALL' else names[effective],prepared=datetime.now(timezone(timedelta(hours=5,minutes=30))).isoformat(timespec='seconds'),currency='INR',columns=[],money_columns=[],rows=[],controls={},note='Accounting working report. Applicable Schedule III/Ind AS presentation, disclosures and final review remain pending.')
    rows=pack['rows'];closing=core.balances(state,end,effective);movement=core.balances(state,end,effective,start)
    if mode=='Trial balance':
        books=ledgers.ledger_pack(state,[a['code'] for a in state['accounts']],start,end,effective)['books'];pack['columns']=['Code','Ledger','Opening Dr','Opening Cr','Period debit','Period credit','Closing Dr','Closing Cr'];pack['money_columns']=list(range(2,8));totals=[0]*6
        for b in books:
            values=[max(b['opening'],0),max(-b['opening'],0),b['debit'],b['credit'],max(b['closing'],0),max(-b['closing'],0)];rows.append([b['code'],b['name'],*values]);totals=[x+y for x,y in zip(totals,values)]
        rows.append(['','TOTAL',*totals]);pack['controls']={'Opening Dr minus Cr':totals[0]-totals[1],'Period debit minus credit':totals[2]-totals[3],'Closing Dr minus Cr':totals[4]-totals[5]}
    elif mode=='General ledger':
        pack['columns']=['Date','Voucher','Unit','Account','Narration','Debit','Credit'];pack['money_columns']=[5,6];dr=cr=0
        for v in sorted(enumerate(state['vouchers']),key=lambda p:(p[1]['date'],p[0])):
            v=v[1]
            if not start<=v['date']<=end:continue
            for l in v['lines']:
                if effective!='ALL' and l['branch']!=effective:continue
                name=next(a['name'] for a in state['accounts'] if a['code']==l['account']);rows.append([v['date'],v['id'],names[l['branch']],l['account']+' '+name,core.display_narration(v,state),l['debit'],l['credit']]);dr+=l['debit'];cr+=l['credit']
        rows.append(['','','','','TOTAL',dr,cr]);pack['controls']={'Period debit minus credit':dr-cr}
    elif mode=='Branch matrix':
        pack['columns']=['Account','Classification',*names.values(),'Entity total'];pack['money_columns']=list(range(2,len(pack['columns'])));units={u:core.balances(state,end,u) for u in names};errors=0
        for a in state['accounts']:
            values=[units[u][a['code']] for u in names];rows.append([a['code']+' '+a['name'],a['kind'],*values,closing[a['code']]]);errors+=abs(sum(values)-closing[a['code']])
        pack['controls']={'Branch-to-entity difference':errors};pack['note']+=' Matrix includes all units regardless of the selected branch; values are debit-positive/credit-negative closing balances. No inter-branch eliminations are performed.'
    else:
        pl=mode in ('Profit and loss','Comparative P&L');comparative=mode.startswith('Comparative')
        pack['columns']=['Code','Classification','Ledger','Current INR']+(['Previous year INR','Change INR'] if comparative else []);pack['money_columns']=list(range(3,len(pack['columns'])))
        now=movement if pl else closing;old_start,old_end=previous_year(start),previous_year(end);prior=core.balances(state,old_end,effective,old_start if pl else None)
        if comparative:pack['comparison']={'start':old_start,'end':old_end}
        totals={'Revenue':[0,0],'Expense':[0,0],'Asset':[0,0],'Liability':[0,0],'Equity':[0,0]}
        def add(code,kind,label,n,o):rows.append([code,kind,label,n,*([o,n-o] if comparative else [])])
        for a in state['accounts']:
            if (a['kind'] in ('Revenue','Expense'))!=pl:continue
            sign=-1 if a['kind'] in ('Revenue','Liability','Equity') else 1;n=now[a['code']]*sign;o=prior[a['code']]*sign;totals[a['kind']][0]+=n;totals[a['kind']][1]+=o;add(a['code'],a['kind'],a['name'],n,o)
        if pl:
            for kind in ('Revenue','Expense'):add('',kind,'Total '+kind.lower(),*totals[kind])
            n,o=[totals['Revenue'][i]-totals['Expense'][i] for i in (0,1)];add('','Result','Profit / (loss)',n,o);pack['controls']={'Revenue minus expenses':n}
        else:
            earnings=[-sum(b[a['code']] for a in state['accounts'] if a['kind'] in ('Revenue','Expense')) for b in (now,prior)]
            add('','Equity','Unclosed accumulated result',*earnings)
            liabilities=[totals['Liability'][i]+totals['Equity'][i]+earnings[i] for i in (0,1)]
            add('','Total','Total assets',*totals['Asset']);add('','Total','Total liabilities and equity',*liabilities)
            pack['controls']={'Assets minus liabilities and equity':totals['Asset'][0]-liabilities[0]}
            if comparative:pack['controls']['Previous assets minus liabilities and equity']=totals['Asset'][1]-liabilities[1]
    return pack


def display_rows(pack):return [['' if v is None else ledgers.number(v,True) if i in pack['money_columns'] else str(v) for i,v in enumerate(row)] for row in pack['rows']]


def captions(pack):
    result=[pack['entity'],pack['mode'],pack['scope'],pack['start']+' to '+pack['end']+' | INR','Prepared India: '+pack['prepared']]
    if pack.get('comparison'):result.append('Previous year: '+pack['comparison']['start']+' to '+pack['comparison']['end'])
    if pack.get('status'):result.append('Status: '+pack['status'])
    result.extend(k.upper()+': '+v for k,v in pack['registration'].items() if v)
    return result


def _worksheet(wb,pack,title):
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    ws=wb.create_sheet(title)
    def append(values):
        ws.append(values)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
    for line in captions(pack)+[pack['note']]:
        append([line]);ws.merge_cells(start_row=ws.max_row,start_column=1,end_row=ws.max_row,end_column=len(pack['columns']));ws.cell(ws.max_row,1).alignment=Alignment(wrap_text=True,vertical='top');ws.row_dimensions[ws.max_row].height=34 if len(line)>100 else 22
    append(pack['columns']);header=ws.max_row
    for cell in ws[header]:cell.font=Font(bold=True,color='293B55');cell.fill=PatternFill('solid',fgColor='E3EEDF')
    for row in pack['rows']:
        append([Decimal(v)/100 if i in pack['money_columns'] and v is not None else v for i,v in enumerate(row)])
        for col in pack['money_columns']:ws.cell(ws.max_row,col+1).number_format='#,##0.00;[Red](#,##0.00)'
        for cell in ws[ws.max_row]:cell.alignment=Alignment(wrap_text=True,vertical='top')
    for label,value in pack['controls'].items():append([label,Decimal(value)/100]);ws.cell(ws.max_row,2).number_format='#,##0.00'
    for i,col in enumerate(pack['columns']):ws.column_dimensions[get_column_letter(i+1)].width=19 if i in pack['money_columns'] else 48 if col in ('Ledger','Narration','Account') else 23
    ws.freeze_panes=f'A{header+1}';ws.auto_filter.ref=f'A{header}:{get_column_letter(len(pack["columns"]))}{header+len(pack["rows"])}';ws.print_title_rows=f'1:{header}';ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.oddFooter.center.text='Page &P of &N'
    return ws


def export_xlsx(pack,path):
    from openpyxl import Workbook
    wb=Workbook();wb.remove(wb.active);_worksheet(wb,pack,'Financial report');wb.save(path);wb.close()


def report_pack(state,start,end,branch='ALL'):
    packs=[snapshot(state,mode,start,end,branch) for mode in MODES]
    for pack in packs:pack['prepared']=packs[0]['prepared']
    return packs


def export_pack_xlsx(packs,path):
    if not packs:raise ValueError('Select at least one report for the pack.')
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    wb=Workbook();ws=wb.active;ws.title='Contents'
    lines=[packs[0]['entity'],'Financial working report pack',packs[0]['start']+' to '+packs[0]['end'],'Prepared India: '+packs[0]['prepared'],packs[0]['note']]
    for line in lines:
        ws.append([line]);ws.cell(ws.max_row,1).data_type='s';ws.merge_cells(start_row=ws.max_row,start_column=1,end_row=ws.max_row,end_column=3);ws.cell(ws.max_row,1).alignment=Alignment(wrap_text=True);ws.row_dimensions[ws.max_row].height=35
    ws.append(['Report','Scope','Currency','Review status'])
    for cell in ws[ws.max_row]:cell.font=Font(bold=True);cell.fill=PatternFill('solid',fgColor='E3EEDF')
    for pack in packs:
        sheet=_worksheet(wb,pack,pack['mode'][:31]);ws.append([pack['mode'],pack['scope'],'INR',pack.get('status','Working report')])
        for cell in ws[ws.max_row]:cell.data_type='s'
        ws.cell(ws.max_row,1).hyperlink="#'"+sheet.title+"'!A1"
    for col in ('A','B','C','D'):ws.column_dimensions[col].width=38
    ws.freeze_panes='A7';wb.save(path);wb.close()


def export_csv(pack,path):
    with open(path,'w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f)
        for line in captions(pack)+[pack['note']]:w.writerow([ledgers.safe_csv(line)])
        w.writerow(pack['columns'])
        for row in pack['rows']:w.writerow([ledgers.number(v) if i in pack['money_columns'] else ledgers.safe_csv(v) for i,v in enumerate(row)])
        for label,value in pack['controls'].items():w.writerow([label,ledgers.number(value)])


def export_pdf(pack,path):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,LongTable,TableStyle,PageBreak
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4,landscape
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    font='Helvetica'
    for candidate in (Path(os.environ.get('WINDIR','C:/Windows'))/'Fonts'/'segoeui.ttf',Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')):
        if candidate.is_file():
            font='RijuFinancial'
            if font not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont(font,str(candidate)))
            break
    style=ParagraphStyle('Financial',fontName=font,fontSize=7,leading=10,wordWrap='CJK');heading=ParagraphStyle('Title',parent=style,fontSize=12,leading=16)
    p=lambda value:Paragraph(html.escape(str(value)).replace('\n','<br/>'),style)
    packs=pack if isinstance(pack,list) else [pack]
    if not packs:raise ValueError('Select at least one report for the pack.')
    first=packs[0];footer_scope=first['scope'] if len({section['scope'] for section in packs})==1 else 'Mixed scopes; see section headings';title=first['mode'] if len(packs)==1 else 'Financial working report pack'
    doc=SimpleDocTemplate(str(path),pagesize=landscape(A4),leftMargin=28,rightMargin=28,topMargin=26,bottomMargin=40,title=title,author=first['entity']);story=[]
    if len(packs)>1:
        for line in (first['entity'],title,first['start']+' to '+first['end'],'Prepared India: '+first['prepared']):story.append(Paragraph(html.escape(line),heading))
        story.extend([p(first['note']),Spacer(1,12)])
        for number,section in enumerate(packs,1):story.append(p(str(number)+'. '+section['mode']+' | '+section['scope']))
        story.append(PageBreak())
    for section_number,pack in enumerate(packs):
        if section_number:story.append(PageBreak())
        fixed=[i for i in range(len(pack['columns'])) if i not in pack['money_columns']];money=pack['money_columns']
        panels=[list(range(len(pack['columns'])))] if len(pack['columns'])<=9 or not money else [fixed+money[i:i+4] for i in range(0,len(money),4)]
        for panelno,indices in enumerate(panels):
            if panelno:story.append(PageBreak())
            for i,line in enumerate(captions(pack)):story.append(Paragraph(html.escape(line),heading if i<2 else style))
            story.extend([p(pack['note']),Spacer(1,10)])
            # Repeated column headings and wrapped paragraphs prevent clipped long narrations.
            weights=[4 if pack['columns'][i]=='Narration' else 2.5 if pack['columns'][i] in ('Ledger','Account') else 1.4 if i not in money else 1.3 for i in indices];widths=[doc.width*w/sum(weights) for w in weights]
            data=[[p(pack['columns'][i]) for i in indices]]+[[p('' if row[i] is None else ledgers.number(row[i],True) if i in money else row[i]) for i in indices] for row in pack['rows']]
            table=LongTable(data,colWidths=widths,repeatRows=1,splitInRow=1);table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E3EEDF')),('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),.25,colors.HexColor('#DCE2DC')),('BOTTOMPADDING',(0,0),(-1,-1),5)]));story.append(table)
        story.append(Spacer(1,10))
        for label,value in pack['controls'].items():story.append(p(label+': '+ledgers.number(value,True)))
        story.extend([Spacer(1,14),p('Prepared by: ____________________    Reviewed by: ____________________')])
    def footer(canvas,document):
        canvas.saveState();canvas.setFont(font,7);canvas.drawString(28,31,first['entity'][:100]);canvas.drawString(28,22,title+' | '+first['start']+' to '+first['end']+' | '+footer_scope);canvas.drawRightString(landscape(A4)[0]-28,22,'Page '+str(document.page));canvas.restoreState()
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
