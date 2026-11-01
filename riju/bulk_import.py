"""Mapped CSV/XLSX import with row review, duplicate identity and atomic posting."""
import copy
import csv
import hashlib
import io
import json
from pathlib import Path
import uuid
from . import core,foundation,activity

TYPES=('Journal','Expense','Petty cash','Opening')
FIELDS=('date','reference','branch','narration','debit_account','credit_account','amount','journal_id')


def _hash(value):return hashlib.sha256(json.dumps(value,sort_keys=True,ensure_ascii=False).encode()).hexdigest()
def _identity(r):return _hash([r['branch'],r['date'],''.join(c for c in r['reference'].upper() if c.isalnum())])
def _lines(r):return [dict(account=r['debit_account'],branch=r['branch'],debit=r['amount'],credit=0),dict(account=r['credit_account'],branch=r['branch'],debit=0,credit=r['amount'])]
def _queue(s,qid):
    q=next((q for q in s['queue'] if q['id']==qid and q.get('kind')=='bulk'),None)
    if not q:raise ValueError('Select a mapped bulk import.')
    return q

def _known(s,exclude=None):
    return {r['key'] for q in s['queue'] if q.get('kind')=='bulk' and q['id']!=exclude for r in q['rows'] if r.get('key') and r['review'] not in ('REJECTED','INVALID','DUPLICATE')}

def _validate(s,kind,r):
    from .tax import day
    day(r['date'])
    if not r['reference'].strip() or not r['narration'].strip():raise ValueError('Reference and narration are required.')
    if kind in ('Expense','Petty cash'):
        a=next((a for a in s['accounts'] if a['code']==r['debit_account']),None)
        if not a or a['kind']!='Expense':raise ValueError('Expense/petty cash must debit a genuine expense ledger.')
        group=s['account_details'].get(r['credit_account'],{}).get('group')
        if group not in ('Cash and bank','Trade payables') or kind=='Petty cash' and r['credit_account']!='1000':raise ValueError('Use cash/bank/payable for expenses; petty cash uses the cash ledger.')
    core.post(copy.deepcopy(s),r['date'],r['narration'],_lines(r),opening=kind=='Opening',metadata={'reference':r['reference']})


def _read(raw,filename,sheet):
    if len(raw)>25*1024*1024:raise ValueError('Import exceeds 25 MB.')
    if Path(filename).suffix.lower()=='.csv':
        rows=list(csv.reader(io.StringIO(raw.decode('utf-8-sig'))));label='CSV'
    elif Path(filename).suffix.lower()=='.xlsx':
        from zipfile import ZipFile
        from openpyxl import load_workbook
        with ZipFile(io.BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist())>100*1024*1024:raise ValueError('Expanded workbook exceeds 100 MB.')
        w=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
        try:
            if sheet not in w.sheetnames:raise ValueError('Choose the exact workbook sheet; none is silently selected.')
            ws=w[sheet]
            if ws.max_row>100000 or ws.max_column>100:raise ValueError('Workbook dimensions exceed import limits.')
            rows=list(ws.values);label=sheet
        finally:w.close()
    else:raise ValueError('Choose CSV or XLSX.')
    if not rows or len(rows)>100001:raise ValueError('Import is empty or exceeds 100000 source rows.')
    headers=[str(v or '').strip() for v in rows[0]]
    if not all(headers) or len(headers)!=len(set(headers)):raise ValueError('Headings must be nonblank and unique.')
    return headers,rows[1:],label


def stage(s,raw,filename,kind,mapping=None,sheet=''):
    if kind not in TYPES:raise ValueError('Choose journal, expense, petty cash or opening import.')
    headers,source,label=_read(raw,filename,sheet);mapping=mapping or {f:f for f in FIELDS if f in headers}
    if set(mapping)-set(FIELDS) or set(FIELDS[:-1])-set(mapping) or len(set(mapping.values()))!=len(mapping) or set(mapping.values())-set(headers):raise ValueError('Map each required field to a distinct exact heading.')
    digest=hashlib.sha256(raw).hexdigest()
    if any(q.get('kind')=='bulk' and q['digest']==digest and q['type']==kind and q['mapping']==mapping and q.get('sheet','')==sheet and q['status']!='CANCELLED' for q in s['queue']):raise ValueError('This file/mapping is already staged; review or cancel its prior batch.')
    q=dict(id=str(uuid.uuid4()),kind='bulk',type=kind,digest=digest,filename=Path(filename).name,mapping=copy.deepcopy(mapping),sheet=sheet,headers=headers,rows=[],status='PENDING_REVIEW')
    seen=_known(s)
    for number,values in enumerate(source,2):
        if not any(v not in (None,'') for v in values):continue
        r=dict(source_row=number,sheet=label,review='PENDING',reason='',error='',raw={h:(values[i].isoformat()[:10] if hasattr(values[i],'isoformat') else values[i]) if i<len(values) else '' for i,h in enumerate(headers)})
        try:
            if len(values)>len(headers) and any(v not in (None,'') for v in values[len(headers):]):raise ValueError('Source row contains extra values beyond the mapped headings.')
            for field in FIELDS:r[field]=r['raw'].get(mapping.get(field,''),'')
            if any(isinstance(v,str) and v.startswith('=') for v in r['raw'].values()):raise ValueError('Formula cells are not import values; provide reviewed literal data.')
            for f in FIELDS:
                if f not in ('amount','date'):r[f]=str(r[f] or '').strip()
            if hasattr(r['date'],'date'):r['date']=r['date'].date().isoformat()
            elif hasattr(r['date'],'isoformat'):r['date']=r['date'].isoformat()
            else:r['date']=str(r['date']).strip()
            from .sales_import import amount
            r['amount']=amount(r['amount'],'Amount');r['key']=_identity(r);_validate(s,kind,r)
            if r['key'] in seen:r.update(review='DUPLICATE',error='Source branch/date/reference already queued or posted, including changed-file copies.')
            else:seen.add(r['key'])
        except (ValueError,KeyError,TypeError) as error:r.update(review='INVALID',error=str(error))
        q['rows'].append(r)
    if not q['rows']:raise ValueError('No source rows found.')
    s['queue'].append(q);core.add_audit(s,'BULK_IMPORT_STAGED',dict(id=q['id'],filename=q['filename'],digest=digest,type=kind,mapping=mapping,rows=len(q['rows'])));return copy.deepcopy(q)


def review_row(s,qid,index,approved,reason):
    q=_queue(s,qid)
    if q['status']!='PENDING_REVIEW' or type(index)!=int or not 0<=index<len(q['rows']):raise ValueError('Select a pending source row.')
    r=q['rows'][index]
    if not str(reason).strip():raise ValueError('Enter the row review/rejection basis.')
    if approved:
        if r['review'] in ('INVALID','DUPLICATE'):raise ValueError('Correct/reimport invalid or duplicate source data; it cannot be approved.')
        _validate(s,q['type'],r)
    before=copy.deepcopy(r);r.update(review='APPROVED' if approved else 'REJECTED',reason=str(reason).strip());core.add_audit(s,'BULK_ROW_REVIEWED',dict(batch=qid,index=index,before=before,after=copy.deepcopy(r)))


def approve(s,qid):
    c=copy.deepcopy(s);q=_queue(c,qid)
    if q['status']!='PENDING_REVIEW' or any(r['review'] not in ('APPROVED','REJECTED') for r in q['rows']):raise ValueError('Approve or explicitly reject every row first; unresolved errors cannot be posted.')
    selected=[r for r in q['rows'] if r['review']=='APPROVED']
    if not selected:raise ValueError('No approved rows to post; cancel the batch instead.')
    known=_known(c,qid);groups={}
    for r in selected:
        if r['key'] in known:raise ValueError('An approved source was queued/posted elsewhere after this review.')
        known.add(r['key']);_validate(c,q['type'],r)
        key=(r['branch'],r['date'],'OPENING' if q['type']=='Opening' else r['journal_id'] or r['key'])
        groups.setdefault(key,[]).append(r)
    # An explicit compound group must not be partly rejected.
    for r in q['rows']:
        if r.get('journal_id') and any(a['journal_id']==r['journal_id'] and a['branch']==r['branch'] and a['date']==r['date'] for a in selected) and r['review']!='APPROVED':raise ValueError('Review every row of a compound journal consistently; partial group approval is blocked.')
    posted=[]
    for key,rows in groups.items():
        lines=[line for r in rows for line in _lines(r)];reference=' / '.join(r['reference'] for r in rows)
        v=core.post(c,key[1],' | '.join(r['narration'] for r in rows),lines,opening=q['type']=='Opening',metadata={'reference':reference},source=dict(kind='bulk_import',document_id=qid,record_id=_hash([r['key'] for r in rows]),row_count=len(rows)))
        for r in rows:r['voucher']=v['id']
        posted.append(v['id'])
    q['status']='PROCESSED';core.add_audit(c,'BULK_IMPORT_APPROVED',dict(batch=qid,vouchers=posted,approved=len(selected),rejected=len(q['rows'])-len(selected)));s.clear();s.update(c);return posted


def cancel(s,qid,reason):
    q=_queue(s,qid)
    if q['status']!='PENDING_REVIEW' or not str(reason).strip():raise ValueError('Only an unposted pending batch can be cancelled, with a reason.')
    q['status']='CANCELLED'
    for r in q['rows']:r.update(review='REJECTED',reason=reason)
    core.add_audit(s,'BULK_IMPORT_CANCELLED',dict(batch=qid,reason=reason))


def template(path,kind):
    if kind not in TYPES:raise ValueError('Choose an import type.')
    Path(path).write_text(','.join(FIELDS)+'\n2026-04-01,SYNTHETIC-1,JP,Synthetic review example,'+('1010,3000' if kind=='Opening' else '4590,1000' if kind=='Petty cash' else '4590,2000')+',100,\n',encoding='utf-8')


def export(s,qid,path,kind):
    from . import financial_reports as fr,history
    from datetime import datetime
    q=_queue(s,qid);dates=[r['date'] for r in q['rows'] if len(str(r.get('date','')))==10];start=min(dates) if dates else '1900-01-01';end=max(dates) if dates else start
    pack=dict(entity=s['settings']['entity'],registration=copy.deepcopy(s['settings']['registration']),mode='Bulk import review',start=start,end=end,branch='ALL',scope='Source rows; explicit unit per row',prepared=datetime.now(history.INDIA).isoformat(),currency='INR',status=q['status'],controls={},note=q['filename']+' | SHA256 '+q['digest']+' | '+q['type']+'. Row review does not post journals. Original field mapping and source rows are retained.',columns=['Source row','Reference','Unit','Date','Debit','Credit','Amount INR','Review','Reason/error','Voucher'],money_columns=[6],rows=[[r['source_row'],r.get('reference',''),r.get('branch',''),r.get('date',''),r.get('debit_account',''),r.get('credit_account',''),r.get('amount') if type(r.get('amount'))==int else None,r['review'],r['reason'] or r['error'],r.get('voucher','')] for r in q['rows']])
    if kind=='xlsx':fr.export_xlsx(pack,path)
    elif kind=='pdf':fr.export_pdf(pack,path)
    else:raise ValueError('Choose Excel or PDF.')
