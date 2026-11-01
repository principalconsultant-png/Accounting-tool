import copy
import json
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
from pypdf import PdfReader
from riju import core,foundation,tax,tds,history
from riju.storage import Vault


def pair(dr,cr,n,branch='JP'):return [dict(account=dr,branch=branch,debit=n,credit=0),dict(account=cr,branch=branch,debit=0,credit=n)]


class TdsTests(unittest.TestCase):
    def setUp(self):
        self.s=core.empty_state();self.rule=dict(family='TDS',key='TEST-J',effective='2024-04-01',label='Synthetic professional fees',section='194J',rate='10',treatment='Withholding',party_category='Company',annual_threshold='1000',single_threshold='500',threshold_basis='Financial-year aggregate',applicability='Synthetic threshold test only',source='Invented test reference',reviewed=True)
        tax.save_rule(self.s,self.rule);self.party=tax.save_party(self.s,dict(name='Synthetic supplier',effective='2024-04-01',category='Company',pan='ABCDE1234F'))
    def expense(self,amount='600',**extra):
        v=dict(date='2026-04-02',branch='JP',reference='INV-1',narration='Synthetic fees',rule_key='TEST-J',party_id=self.party['party_id'],mode='Automatic',gross=amount,base=amount,counter_account='2000',account='4510',resident_reviewed=True,confirmed=True);v.update(extra);return tds.post(self.s,'Expense',v)
    def challan(self,amount='120',**extra):
        v=dict(date='2026-04-05',branch='JP',reference='UTR-1',narration='Synthetic tax payment',bank_account='1010',rows=[dict(section='194J',branch='JP',year=2026,tax=amount)],confirmed=True);v.update(extra);return tds.post(self.s,'Challan',v)
    def test_annual_threshold_crossing_catches_up_all_base(self):
        a=self.expense();b=self.expense(reference='INV-2',date='2026-04-03')
        self.assertEqual(a['withheld'],0);self.assertEqual(b['withheld'],12000);self.assertEqual(b['assessment']['base_after'],120000)
        self.assertEqual(core.balances(self.s,'2026-04-30')['2100'],-12000);core.verify_audit(self.s)
    def test_aggregate_is_entity_wide_and_same_pan_ids_merge(self):
        self.expense();other=tax.save_party(self.s,dict(name='Same PAN alias',effective='2024-04-01',category='Company',pan='ABCDE1234F'))
        b=self.expense(reference='KM-1',branch='KM',party_id=other['party_id']);self.assertEqual(b['withheld'],12000)
    def test_earlier_years_separate_and_late_tax_point_uses_original_year(self):
        self.expense(date='2025-03-31',reference='OLD',amount='1100')
        fresh=self.expense();self.assertEqual(fresh['withheld'],0)
        late=self.expense(amount='200',reference='LATE',date='2026-04-06',tax_date='2025-03-31');self.assertEqual(late['year'],2024);self.assertEqual(late['withheld'],2000)
    def test_historical_rates_and_aggregation_groups_preserved(self):
        self.expense();tax.save_rule(self.s,dict(self.rule,effective='2026-04-03',rate='5'))
        r=self.expense(reference='NEW',date='2026-04-03');self.assertEqual(r['withheld'],9000)
        tax.save_rule(self.s,dict(self.rule,key='OTHER',threshold_group='Other nature',rate='2'))
        r=self.expense(reference='GROUP',date='2026-04-04',rule_key='OTHER');self.assertEqual(r['withheld'],0)
    def test_single_either_excess_and_rounding_settings(self):
        tax.save_rule(self.s,dict(self.rule,key='PER',threshold_basis='Per payment'))
        self.assertEqual(self.expense(rule_key='PER')['withheld'],6000)
        state=copy.deepcopy(self.s);tax.save_rule(self.s,dict(self.rule,key='EXCESS',threshold_group='Excess',deduction_method='Excess over annual threshold'))
        self.assertEqual(self.expense(rule_key='EXCESS',reference='EX1')['withheld'],0)
        self.assertEqual(self.expense(rule_key='EXCESS',reference='EX2',date='2026-04-03')['withheld'],2000)
        self.s=state;tax.save_rule(self.s,dict(self.rule,key='ROUND',threshold_group='Rounded',threshold_basis='Per payment',single_threshold='0',tds_rounding='Whole rupee'))
        self.assertEqual(self.expense('15.50',rule_key='ROUND',reference='ROUND')['withheld'],200)
    def test_trigger_equality_and_current_only_do_not_erase_prior_deductions(self):
        tax.save_rule(self.s,dict(self.rule,key='CURRENT',threshold_group='Current',threshold_basis='Per payment',single_threshold='100',threshold_trigger='At least',deduction_method='Current transaction only'))
        self.assertEqual(self.expense('100',rule_key='CURRENT')['withheld'],1000)
        self.assertEqual(self.expense('100',rule_key='CURRENT',reference='NEXT',date='2026-04-03')['withheld'],1000)
    def test_194c_and_194i_use_separate_section_controls(self):
        for section,code in [('194C','2110'),('194I','2120')]:
            tax.save_rule(self.s,dict(self.rule,key=section,section=section,threshold_basis='Per payment',single_threshold='0'))
            r=self.expense('100',rule_key=section,reference=section);self.assertEqual(r['tds_account'],code)
    def test_manual_exception_unknown_pan_keeps_source_and_does_not_get_reassessed(self):
        p=tax.save_party(self.s,dict(name='Unknown PAN',effective='2024-04-01',category='Company'))
        with self.assertRaises(ValueError):self.expense(party_id=p['party_id'])
        r=self.expense(party_id=p['party_id'],mode='Reviewed manual',manual_tax='60',review_reason='Reviewed synthetic exception');self.assertEqual(r['withheld'],6000)
        tax.save_party(self.s,dict(name='Unknown PAN',effective='2026-04-03',category='Company',pan='ZZZZZ1234Z'),p['party_id'])
        r=self.expense(reference='NEXT',date='2026-04-03',party_id=p['party_id']);self.assertEqual(r['withheld'],6000)
    def test_wrong_category_unknown_threshold_unreviewed_and_future_tax_date_atomic(self):
        for changes in ({'resident_reviewed':False},{'tax_date':'2027-01-01'},{'mode':'Reviewed manual','review_reason':''},{'gross':'500','base':'600'}):
            before=copy.deepcopy(self.s)
            with self.assertRaises(ValueError):self.expense(**changes)
            self.assertEqual(before,self.s)
        tax.save_rule(self.s,dict(self.rule,key='BAD',party_category='Individual / HUF'))
        with self.assertRaises(ValueError):self.expense(rule_key='BAD')
    def test_existing_payable_does_not_record_expense_twice_and_zero_record_voids(self):
        source=core.post(self.s,'2026-04-01','Existing expense',pair('4510','2000',60000))
        v=dict(date='2026-04-02',branch='JP',reference='ATTACH-1',narration='Attach source tax treatment',rule_key='TEST-J',party_id=self.party['party_id'],mode='Automatic',gross='600',base='600',counter_account='2000',source_voucher=source['id'],resident_reviewed=True,confirmed=True)
        r=tds.post(self.s,'Existing payable',v);self.assertIsNone(r['voucher_id']);self.assertEqual(len(self.s['vouchers']),1)
        with self.assertRaises(ValueError):core.reverse(self.s,source['id'],'2026-04-03')
        tds.void_record(self.s,r['id'],'2026-04-03','Correction');core.reverse(self.s,source['id'],'2026-04-03')
    def test_existing_source_gross_cannot_be_reused_even_by_backdating(self):
        source=core.post(self.s,'2026-04-01','Existing expense',pair('4510','2000',60000))
        v=dict(date='2026-04-05',branch='JP',reference='ATTACH-1',narration='Source',rule_key='TEST-J',party_id=self.party['party_id'],mode='Automatic',gross='600',base='600',counter_account='2000',source_voucher=source['id'],resident_reviewed=True,confirmed=True)
        tds.post(self.s,'Existing payable',v);before=copy.deepcopy(self.s)
        with self.assertRaises(ValueError):tds.post(self.s,'Existing payable',dict(v,date='2026-04-03',reference='BACK'))
        self.assertEqual(before,self.s)
    def test_catchup_adjustment_against_payable(self):
        self.expense('600',mode='Reviewed manual',manual_tax='0',review_reason='Reviewed pending treatment')
        self.expense('600',reference='AUTO',date='2026-04-03')
        # Current automatic tax respects the prior manual target rather than overriding it.
        self.assertEqual(self.s['tds_records'][-1]['withheld'],6000)
        r=tds.post(self.s,'Adjustment',dict(date='2026-04-04',branch='JP',reference='CORR',narration='Reviewed additional tax',rule_key='TEST-J',party_id=self.party['party_id'],mode='Reviewed manual',manual_tax='10',review_reason='Verified adjustment',counter_account='2000',resident_reviewed=True,confirmed=True))
        self.assertEqual(r['base'],0);self.assertEqual(r['withheld'],1000)
    def test_challan_payment_interest_and_partial_allocations_reconcile(self):
        self.expense();d=self.expense(reference='INV-2',date='2026-04-03');p=self.challan(interest='5',fee='2',charge_account='4580')
        self.assertEqual(p['total'],12700)
        tds.allocate(self.s,d['id'],p['id'],0,'50','2026-04-05');tds.allocate(self.s,d['id'],p['id'],0,'70','2026-04-06')
        report=tds.report(self.s,'2026-04-01','2027-03-31');self.assertEqual(report['outstanding']['rows'],[]);self.assertFalse(any(report['deductions']['controls'].values()));self.assertEqual(report['challans']['rows'][0][7],0)
    def test_cannot_allocate_interest_or_reuse_tax_or_wrong_section_branch_year(self):
        self.expense();d=self.expense(reference='INV-2',date='2026-04-03');p=self.challan('100',interest='20',charge_account='4580')
        with self.assertRaises(ValueError):tds.allocate(self.s,d['id'],p['id'],0,'120','2026-04-05')
        tds.allocate(self.s,d['id'],p['id'],0,'100','2026-04-05')
        with self.assertRaises(ValueError):tds.allocate(self.s,d['id'],p['id'],0,'1','2026-04-05')
        other=self.challan(reference='WRONG',rows=[dict(section='194C',branch='JP',year=2026,tax='20')])
        with self.assertRaises(ValueError):tds.allocate(self.s,d['id'],other['id'],0,'20','2026-04-06')
    def test_multi_unit_challan_balances_and_reversal_requires_dated_unallocation(self):
        self.expense();d=self.expense(reference='KM-1',branch='KM',date='2026-04-03')
        p=self.challan(rows=[dict(section='194J',branch='KM',year=2026,tax='120'),dict(section='194C',branch='JP',year=2026,tax='10')]);a=tds.allocate(self.s,d['id'],p['id'],0,'120','2026-04-05')
        for r in (d,p):
            with self.assertRaises(ValueError):core.reverse(self.s,r['voucher_id'],'2026-04-06')
        tds.unallocate(self.s,a['id'],'2026-04-07','Reviewed correction');self.assertEqual(tds.report(self.s,'2026-04-01','2026-04-06')['deductions']['rows'][-1][11],12000)
        core.reverse(self.s,p['voucher_id'],'2026-04-07');core.reverse(self.s,d['voucher_id'],'2026-04-07');core.verify_audit(self.s)
    def test_cin_optional_then_revision_validated_and_duplicate_cin_rejected(self):
        p=self.challan();tds.challan_details(self.s,p['id'],{})
        info=dict(bsr='1234567',serial='00001',deposit_date='2026-04-05',verified=True)
        tds.challan_details(self.s,p['id'],info);q=self.challan(reference='UTR-2')
        with self.assertRaises(ValueError):tds.challan_details(self.s,q['id'],info)
        with self.assertRaises(ValueError):tds.challan_details(self.s,q['id'],dict(bsr='bad'))
    def test_advance_challan_and_allocation_late_year_preserve_liability_year(self):
        p=self.challan(date='2026-04-01',rows=[dict(section='194J',branch='JP',year=2025,tax='110')])
        d=self.expense('1100',date='2026-04-02',tax_date='2026-03-31')
        self.assertEqual(d['year'],2025);tds.allocate(self.s,d['id'],p['id'],0,'110','2026-04-02')
    def test_reports_excel_pdf_readonly_and_unknown_tan_allowed(self):
        self.expense();self.expense(reference='INV-2',date='2026-04-03');self.challan();before=copy.deepcopy(self.s)
        with tempfile.TemporaryDirectory() as d:
            root=Path(d)
            for kind in ('xlsx','pdf','json'):tds.export(self.s,'2026-04-01','2027-03-31','ALL',root/('tds.'+kind),kind)
            w=load_workbook(root/'tds.xlsx');self.assertEqual(len(w.sheetnames),6);w.close();self.assertIn('TDS',PdfReader(root/'tds.pdf').pages[0].extract_text())
        self.assertEqual(before,self.s)
    def test_format_nine_migration_preupgrade_backup_and_preservation(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);ext=root/'external';ext.mkdir();old=copy.deepcopy(self.s);old['version']=9;old['schema_revision']=9
            for key in ('tds_records','tds_allocations','tds_challan_details','tds_ledgers'):old.pop(key)
            history.checkpoint(old);p=root/'old.json';p.write_text(json.dumps(old));raw=p.read_bytes();v=Vault(p,demo=True);new=v.load();new['settings']['backup_dir']=str(ext);v.save(new)
            self.assertEqual(v.last_upgrade_backup.read_bytes(),raw);self.assertEqual(new['version'],19);self.assertEqual(new['tax_config'],old['tax_config']);self.assertEqual(v.load(),new)

    def test_automatic_catchup_uses_recorded_under_deduction_target(self):
        self.expense(mode='Reviewed manual',manual_tax='0',manual_target='60',review_reason='Verified liability recorded before deduction')
        r=tds.post(self.s,'Adjustment',dict(date='2026-04-03',branch='JP',reference='AUTO-CATCH',narration='Deduct recorded arrears',rule_key='TEST-J',party_id=self.party['party_id'],mode='Automatic',counter_account='2000',resident_reviewed=True,confirmed=True))
        self.assertEqual(r['withheld'],6000);self.assertEqual(r['base'],0)
        with self.assertRaises(ValueError):tds.post(self.s,'Adjustment',dict(date='2026-04-04',branch='JP',reference='REPEAT',narration='No remaining arrears',rule_key='TEST-J',party_id=self.party['party_id'],mode='Automatic',counter_account='2000',resident_reviewed=True,confirmed=True))

    def test_register_legacy_deduction_adds_metadata_without_double_posting(self):
        source=core.post(self.s,'2026-04-01','Already deducted',[dict(account='4510',branch='JP',debit=60000,credit=0),dict(account='2000',branch='JP',debit=0,credit=54000),dict(account='2100',branch='JP',debit=0,credit=6000)])
        before=len(self.s['vouchers']);r=tds.post(self.s,'Register existing deduction',dict(date='2026-04-02',branch='JP',reference='LEGACY-1',narration='Register source TDS',rule_key='TEST-J',party_id=self.party['party_id'],mode='Reviewed manual',manual_tax='60',review_reason='Verified source ledger',base='600',gross='600',counter_account='2000',tds_account='2100',source_voucher=source['id'],resident_reviewed=True,confirmed=True))
        self.assertIsNone(r['voucher_id']);self.assertEqual(len(self.s['vouchers']),before)
        self.assertFalse(any(tds.report(self.s,'2026-04-01','2027-03-31')['deductions']['controls'].values()))
        p=self.challan('60');tds.allocate(self.s,r['id'],p['id'],0,'60','2026-04-05')
        self.assertEqual(tds.report(self.s,'2026-04-01','2027-03-31')['outstanding']['rows'],[])
        self.assertEqual(len(history.trace_voucher(self.s,source['id'])['tds_context']),2)
    def test_unregistered_manual_ledger_movements_are_exposed(self):
        core.post(self.s,'2026-04-01','Unregistered legacy tax',pair('4510','2100',1000))
        r=tds.report(self.s,'2026-04-01','2027-03-31');self.assertEqual(r['deductions']['controls']['2100 unclassified/manual movement'],-1000)

    def test_period_register_shows_dated_reversal_instead_of_erasing_history(self):
        r=self.expense('1200');core.reverse(self.s,r['voucher_id'],'2026-05-01')
        april=tds.report(self.s,'2026-04-01','2026-04-30');may=tds.report(self.s,'2026-05-01','2026-05-31')
        self.assertEqual(april['sections']['194J']['deducted'],12000);self.assertEqual(may['sections']['194J']['deducted'],-12000)
        self.assertEqual(may['outstanding']['rows'],[]);self.assertFalse(any(may['deductions']['controls'].values()))

    def test_mixed_challan_charges_belong_to_their_posting_unit(self):
        self.challan(rows=[dict(section='194J',branch='KM',year=2026,tax='100')],interest='5',fee='2',charge_account='4580')
        km=tds.report(self.s,'2026-04-01','2027-03-31','KM')['challans']['rows'];jp=tds.report(self.s,'2026-04-01','2027-03-31','JP')['challans']['rows']
        self.assertEqual(sum(r[9]+r[10] for r in km),0);self.assertEqual(sum(r[9]+r[10] for r in jp),700)
    def test_source_snapshots_and_cin_leading_zeros_export_as_literal_text(self):
        p=self.challan();tds.challan_details(self.s,p['id'],dict(bsr='0123456',serial='00001',deposit_date='2026-04-05',verified=True))
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'tds.xlsx';tds.export(self.s,'2026-04-01','2027-03-31','ALL',path,'xlsx');w=load_workbook(path)
            ws=w['TDS challan reconciliation'];values=[c.value for row in ws for c in row];self.assertIn('0123456',values);self.assertIn('00001',values);w.close()

    def test_late_backfilled_tax_point_includes_later_known_year_bases(self):
        first=self.expense('600',date='2026-05-03',tax_date='2026-05-03',reference='MAY-3')
        self.assertEqual(first['withheld'],0)
        late=self.expense('600',date='2026-06-02',tax_date='2026-05-02',reference='LATE-MAY-2')
        self.assertEqual(late['withheld'],12000);self.assertEqual(late['assessment']['base_after'],120000)
