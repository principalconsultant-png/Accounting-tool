import copy
import json
from datetime import date,timedelta
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
from pypdf import PdfReader
from riju import core,foundation,history,fixed_assets as book,tax_depreciation as td
from riju.storage import Vault

class TaxDepreciationTests(unittest.TestCase):
    def setUp(self):
        self.s=core.empty_state();self.b=td.save_block(self.s,'PLANT','Synthetic tax block',2024,{'JP':dict(wdv='1000',count=2)},'Reviewed opening tax WDV/count source',True);self.rule()
    def rule(self,**extra):
        opts=dict(effective_year=2024,rate='15',basis='Reviewed synthetic Appendix I rate and regime',reviewed=True);opts.update(extra);return td.save_rule(self.s,self.b['id'],**opts)
    def add(self,**extra):
        opts=dict(reference='A-1',branch='JP',acquired='2024-04-01',cost='1000',basis='Reviewed qualifying actual cost',reviewed=True,put_to_use='2024-04-01');opts.update(extra);return td.add_asset(self.s,self.b['id'],**opts)
    def result(self,y=2024):return td.calculate(self.s,y)
    def item(self,y=2024):return self.result(y)['blocks'][0]
    def sale(self,**extra):
        opts=dict(reference='SALE-1',branch='JP',on='2025-03-31',consideration='300',expenses='0',basis='Reviewed gross moneys payable and eligible transfer expenses',reviewed=True,count=1);opts.update(extra);return td.dispose(self.s,self.b['id'],**opts)
    def test_opening_wdv_annual_allowance_and_no_book_post(self):
        r=self.item();self.assertEqual(r['normal'],15000);self.assertEqual(r['closing'],85000);self.assertTrue(self.result()['ready']);self.assertFalse(self.s['vouchers']);core.verify_audit(self.s)
    def test_exact_180_days_is_full_and_179_is_short(self):
        end=date(2025,3,31);full=(end-timedelta(days=179)).isoformat();short=(end-timedelta(days=178)).isoformat()
        self.add(acquired=full,put_to_use=full);self.add(reference='A-2',acquired=short,put_to_use=short)
        r=self.item();self.assertEqual(r['full'],100000);self.assertEqual(r['short'],100000);self.assertEqual(r['normal'],37500);self.assertEqual([row[4] for row in r['facts']],[180,179])
    def test_short_use_counts_actual_days_until_sale_not_just_year_end(self):
        a=self.add();self.sale(addition_id=a['id'],on='2024-06-01',consideration='200');r=self.item();self.assertEqual(r['short'],100000);self.assertEqual(r['eligible_short'],80000);self.assertEqual(r['normal'],21000)
    def test_disposal_reduces_actual_short_cohort_first(self):
        a=self.add(acquired='2024-12-01',put_to_use='2024-12-01',cost='500');self.sale(addition_id=a['id'],consideration='200');r=self.item()
        self.assertEqual(r['eligible_full'],100000);self.assertEqual(r['eligible_short'],30000);self.assertEqual(r['normal'],17250);self.assertEqual(r['closing'],112750)
    def test_excess_full_realizations_offset_short_pool_without_negative_depreciation(self):
        self.add(acquired='2024-12-01',put_to_use='2024-12-01',cost='500');self.sale(consideration='1200');r=self.item();self.assertEqual(r['eligible_full'],0);self.assertEqual(r['eligible_short'],30000);self.assertEqual(r['normal'],2250);self.assertEqual(r['closing'],27750)
    def test_excess_short_realizations_offset_opening_pool(self):
        a=self.add(acquired='2024-12-01',put_to_use='2024-12-01',cost='500');self.sale(addition_id=a['id'],consideration='700');r=self.item();self.assertEqual(r['eligible_full'],80000);self.assertEqual(r['eligible_short'],0);self.assertEqual(r['normal'],12000)
    def test_block_ceases_no_depreciation_and_section50_loss_includes_transfer_expenses(self):
        self.sale(consideration='800',expenses='50',count=2);r=self.item();self.assertFalse(r['continued']);self.assertEqual(r['normal'],0);self.assertEqual(r['loss'],25000);self.assertEqual(r['closing'],0)
    def test_continuing_block_excess_section50_gain_uses_expenses_separately(self):
        self.sale(consideration='1500',expenses='100');r=self.item();self.assertTrue(r['continued']);self.assertEqual(r['gain'],40000);self.assertEqual(r['closing'],0);self.assertEqual(r['normal'],0)
    def test_opening_sale_cannot_consume_current_addition_counts(self):
        td.revise_opening(self.s,self.b['id'],{},'No pre-existing tax block assets',True);self.add();before=copy.deepcopy(self.s)
        with self.assertRaisesRegex(ValueError,'count'):self.sale()
        self.assertEqual(before,self.s)
    def test_unused_assets_excluded_and_later_use_of_prior_acquisition_is_full_rate(self):
        td.revise_opening(self.s,self.b['id'],{},'Zero opening WDV',True);a=self.add(acquired='2024-12-01',put_to_use='');r=self.item();self.assertEqual(r['normal'],0);self.assertEqual(r['pending'][0][4],100000)
        frozen=td.finalize(self.s,2024,'Reviewed prior tax year',True);td.commission(self.s,a['id'],'2025-12-01','Actual commissioning certificate',True)
        self.assertEqual(td.calculate(self.s,2024)['blocks'],frozen['report']['blocks']);self.assertEqual(self.item(2025)['full'],100000);self.assertEqual(self.item(2025)['short'],0);self.assertEqual(self.item(2025)['normal'],15000)
    def test_pre_use_disposal_stays_outside_depreciable_pool(self):
        a=self.add(put_to_use='');self.sale(addition_id=a['id'],consideration='800');r=self.item();self.assertEqual(r['realizations'],0);self.assertEqual(r['normal'],15000);self.assertIn('outside',r['pending'][0][-1])
        with self.assertRaises(ValueError):td.commission(self.s,a['id'],'2025-04-01','After sale',True)
    def test_reviewed_additional_half_and_next_year_original_entitlement(self):
        td.revise_opening(self.s,self.b['id'],{},'No opening assets',True);self.rule(additional_rate='20',additional_enabled=True)
        self.add(acquired='2024-12-01',put_to_use='2024-12-01',additional_eligible=True,additional_basis='Synthetic eligible new manufacturing machinery, approved business/regime')
        r=self.item();self.assertEqual(r['normal'],7500);self.assertEqual(r['additional'],10000);self.assertEqual(r['closing'],82500)
        td.finalize(self.s,2024,'All synthetic sources reviewed',True);self.rule(effective_year=2025,additional_rate='10',additional_enabled=True)
        r=self.item(2025);self.assertEqual(r['normal'],12375);self.assertEqual(r['additional'],10000);self.assertEqual(r['closing'],60125);self.assertFalse(r['carry'])
    def test_disallowed_regime_carry_requires_explicit_forfeiture(self):
        self.rule(additional_rate='20',additional_enabled=True);a=self.add(acquired='2024-12-01',put_to_use='2024-12-01',additional_eligible=True,additional_basis='Reviewed eligible source');td.finalize(self.s,2024,'Synthetic review',True);self.rule(effective_year=2025,additional_enabled=False)
        self.assertFalse(self.result(2025)['ready'])
        with self.assertRaises(ValueError):td.additional_claim(self.s,a['id'],2025,'100','0','Cannot bypass regime',True)
        td.additional_claim(self.s,a['id'],2025,'0','0','Explicitly reviewed regime forfeiture',True);self.assertTrue(self.result(2025)['ready'])
    def test_disposed_eligible_asset_requires_explicit_additional_review(self):
        self.rule(additional_rate='20',additional_enabled=True);a=self.add(additional_eligible=True,additional_basis='Reviewed eligible source');self.sale(addition_id=a['id'],consideration='500');self.assertFalse(self.result()['ready'])
        td.additional_claim(self.s,a['id'],2024,'0','0','Reviewed disposal means no additional claim in this working',True);self.assertTrue(self.result()['ready'])
    def test_missing_policy_and_invalid_inputs_do_not_post(self):
        before=copy.deepcopy(self.s)
        for action in [lambda:td.save_rule(self.s,self.b['id'],2024,'NaN','Bad',True),lambda:td.save_rule(self.s,self.b['id'],2024,'101','Bad',True),lambda:self.add(reviewed=False),lambda:self.add(put_to_use='2023-04-01'),lambda:self.add(additional_eligible=True),lambda:td.save_block(self.s,'BAD','Bad',2024,{'JP':dict(wdv='100',count=0)},'Bad',True)]:
            with self.assertRaises(ValueError):action()
        self.assertEqual(before,self.s);self.assertFalse(self.s['vouchers'])
        other=td.save_block(self.s,'FURN','No rate yet',2024,{},'Review',True);self.assertFalse(self.result()['ready'])
    def test_branch_pooling_bridge_reconciles_entity_and_does_not_create_unit_capital_loss(self):
        td.revise_opening(self.s,self.b['id'],{'JP':dict(wdv='800',count=1),'KM':dict(wdv='200',count=1)},'Reviewed two-unit pool',True);self.sale(branch='JP',consideration='100');r=self.item()
        self.assertTrue(r['continued']);self.assertEqual(r['loss'],0);self.assertEqual(sum(row[9] for row in r['branches']),r['closing']);jp=next(row for row in r['branches'] if row[1]=='JP');self.assertEqual(jp[9],0)
        for row in r['branches']:self.assertEqual(row[2]+row[3]+row[4]-row[5]+row[6]-row[7]-row[8],row[9])
    def test_prior_years_must_be_finalized_before_carry_is_accepted(self):
        self.assertFalse(self.result(2025)['ready'])
        with self.assertRaises(ValueError):td.finalize(self.s,2025,'Prior missing',True)
        td.finalize(self.s,2024,'2024 reviewed',True);self.assertTrue(self.result(2025)['ready']);self.assertEqual(self.item(2025)['opening'],85000)
    def test_finalize_freezes_report_and_latest_first_reopening_protects_chain(self):
        first=td.finalize(self.s,2024,'Reviewed source',True);td.finalize(self.s,2025,'Reviewed next year',True);before=copy.deepcopy(self.s)
        with self.assertRaises(ValueError):td.reopen(self.s,2024,'Too early')
        with self.assertRaises(ValueError):self.rule(rate='20')
        with self.assertRaises(ValueError):td.revise_opening(self.s,self.b['id'],{},'Too late',True)
        self.assertEqual(before,self.s);td.reopen(self.s,2025,'Correct carry');td.reopen(self.s,2024,'Correct opening');td.revise_opening(self.s,self.b['id'],{'JP':dict(wdv='2000',count=2)},'Corrected source',True)
        self.assertEqual(self.item()['opening'],200000);self.assertEqual(td.data(self.s)['reviews'][0]['report'],first['report'])
    def test_current_or_future_year_cannot_be_finalized_early(self):
        self.s=core.empty_state();b=td.save_block(self.s,'P','Future',2090,{},'Test',True);td.save_rule(self.s,b['id'],2090,'15','Test',True)
        with self.assertRaisesRegex(ValueError,'Finish'):td.finalize(self.s,2090,'Too early',True)
    def test_source_voiding_retains_history_and_duplicate_guards(self):
        a=self.add(put_to_use='');before=copy.deepcopy(self.s)
        with self.assertRaises(ValueError):self.add()
        self.assertEqual(before,self.s);td.void_record(self.s,'additions',a['id'],'Duplicate cost source corrected');self.assertFalse(td.data(self.s)['additions'][0]['active']);self.add();self.assertEqual(len(td.data(self.s)['additions']),2)
    def test_book_link_is_snapshot_only_and_final_tax_year_protects_backdated_source_reversal(self):
        a=book.register(self.s,dict(reference='BOOK',name='Synthetic',branch='JP',mode='Post purchase',category='Manual / specific Schedule II category',acquired='2024-04-01',put_to_use='2024-04-01',baseline='2024-04-01',cost='1000',residual='0',months=12,basis='Reviewed book source',reviewed=True))
        self.add(reference='LINK',book_asset_id=a['id']);count=len(self.s['vouchers']);td.finalize(self.s,2024,'All sources reviewed',True)
        with self.assertRaises(ValueError):core.reverse(self.s,a['cost_voucher'],'2025-03-31')
        self.assertEqual(count,len(self.s['vouchers']));self.assertEqual(td.data(self.s)['additions'][0]['book_source']['id'],a['id'])
    def test_rounding_and_largest_remainder_allocations_are_exact(self):
        td.revise_opening(self.s,self.b['id'],{'JP':dict(wdv='1000.01',count=1),'KM':dict(wdv='1000.01',count=1)},'Rounded allocation',True);self.rule(rounding='Whole rupee');r=self.item();self.assertEqual(r['normal'],30000);self.assertEqual(sum(row[9] for row in r['branches']),r['closing']);self.assertEqual(sum(td.allocate(101,{'JP':1,'KM':1,'VK':1}).values()),101)
    def test_branch_export_keeps_entity_summary_and_readonly_excel_pdf_source_policy(self):
        self.add(acquired='2024-12-01',put_to_use='2024-12-01');before=copy.deepcopy(self.s)
        with tempfile.TemporaryDirectory() as d:
            root=Path(d)
            for kind in ('xlsx','pdf','json'):td.export(self.s,2024,'JP',root/('tax.'+kind),kind)
            w=load_workbook(root/'tax.xlsx');self.assertIn('Tax rates and annual policy',w.sheetnames);self.assertTrue(any(c.data_type=='n' for row in w['Entity tax WDV blocks'] for c in row));w.close();text=' '.join(p.extract_text() for p in PdfReader(root/'tax.pdf').pages);self.assertIn('Tax rates and annual policy',text);self.assertIn('management allocation',text);self.assertIn('ReviewedsyntheticAppendixIrateandregime',''.join(text.split()))
        self.assertEqual(before,self.s)
    def test_new_schema_preserves_original_format13_backup_and_reviews(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);ext=root/'external';ext.mkdir();old=core.empty_state();old['version']=13;old['schema_revision']=13;old.pop('tax_depreciation');history.checkpoint(old);path=root/'old.json';path.write_text(json.dumps(old),encoding='utf-8');raw=path.read_bytes();v=Vault(path,demo=True);self.s=v.load();self.s['settings']['backup_dir']=str(ext);b=td.save_block(self.s,'P','Test',2024,{'JP':dict(wdv='100',count=1)},'Test',True);td.save_rule(self.s,b['id'],2024,'15','Test',True);td.finalize(self.s,2024,'Test',True);v.save(self.s);self.assertEqual(v.last_upgrade_backup.read_bytes(),raw);self.assertEqual(v.load(),self.s);self.assertEqual(self.s['version'],19)
    def test_fully_written_down_block_scrapping_expenses_are_a_cessation_loss(self):
        td.revise_opening(self.s,self.b['id'],{'JP':dict(wdv='0',count=2)},'Fully written-down assets remain in the block',True);self.sale(consideration='0',expenses='50',count=2);self.assertEqual(self.item()['loss'],5000)
    def test_additional_claim_cap_must_be_explicitly_resolved(self):
        self.rule(additional_rate='20',additional_enabled=True);a=self.add(additional_eligible=True,additional_basis='Reviewed eligible new plant');self.sale(consideration='1900');self.assertFalse(self.result()['ready'])
        td.additional_claim(self.s,a['id'],2024,'85','0','Reduced allowance reviewed against remaining pool basis',True);self.assertTrue(self.result()['ready']);self.assertEqual(self.item()['additional'],8500)
    def test_void_commissioning_and_unused_block_deactivation_are_audited(self):
        a=self.add(put_to_use='');u=td.commission(self.s,a['id'],'2024-12-01','Actual use',True);td.void_record(self.s,'uses',u['id'],'Incorrect commissioning document');self.assertEqual(len(self.item()['pending']),1)
        td.void_record(self.s,'additions',a['id'],'Incorrect capital source');td.deactivate_block(self.s,self.b['id'],'Wrong unused setup');self.assertFalse(td.data(self.s)['blocks'][0]['active']);self.assertFalse(self.result()['ready'])
    def test_unused_capital_with_block_cessation_requires_explicit_status_review(self):
        self.add(put_to_use='');self.sale(count=2,consideration='800');self.assertFalse(self.result()['ready'])
        td.review_block_status(self.s,self.b['id'],2024,'Ceased; unused costs excluded','Tax review confirms unused capital outside this statutory block',True);self.assertTrue(self.result()['ready']);self.assertEqual(self.item()['loss'],20000)
    def test_reviewed_continuance_without_allowance_preserves_basis_and_unit_proof(self):
        self.add(put_to_use='');self.sale(count=2,consideration='800');td.review_block_status(self.s,self.b['id'],2024,'Continuing; no current allowance','Tax review confirms block exists; no supported current use allowance',True)
        r=self.item();self.assertTrue(self.result()['ready']);self.assertEqual(r['normal'],0);self.assertEqual(r['closing'],20000);self.assertEqual(r['loss'],0);self.assertEqual(sum(row[9] for row in r['branches']),20000)
    def test_later_regime_change_does_not_silently_validate_an_old_positive_override(self):
        self.rule(additional_rate='20',additional_enabled=True);a=self.add(additional_eligible=True,additional_basis='Reviewed eligible source');td.additional_claim(self.s,a['id'],2024,'200','0','Reviewed supported allowance',True)
        self.rule(additional_rate='20',additional_enabled=False);self.assertFalse(self.result()['ready']);td.additional_claim(self.s,a['id'],2024,'0','0','Reviewed regime correction / forfeiture',True);self.assertTrue(self.result()['ready'])
