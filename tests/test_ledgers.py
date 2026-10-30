import copy
import csv
import io
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
from pypdf import PdfReader
from riju import core,foundation,ledgers


def lines(debit,credit,amount,branch='JP'):
    return [dict(account=debit,branch=branch,debit=amount,credit=0),dict(account=credit,branch=branch,debit=0,credit=amount)]


class LedgerTests(unittest.TestCase):
    def setUp(self):
        self.s=core.empty_state()
        core.post(self.s,'2023-04-01','Initial bank and capital',lines('1010','3000',100000),opening=True)
        core.post(self.s,'2023-04-02','Consultation sales',lines('1010','4000',50000),metadata={'type':'Receipt','reference':'REC-1'})
        core.post(self.s,'2023-04-03','Rent',lines('4530','1010',20000),metadata={'type':'Payment','reference':'RENT-1'})
        core.post(self.s,'2023-04-03','Other branch sale',lines('1000','4000',7000,'KM'))

    def pack(self,codes=('1010',),start='2023-04-01',end='2023-04-30',branch='JP'):
        return ledgers.ledger_pack(self.s,codes,start,end,branch)

    def test_opening_movements_running_and_closing(self):
        book=self.pack()['books'][0]
        self.assertEqual((book['opening'],book['debit'],book['credit'],book['closing']),(100000,50000,20000,130000))
        self.assertEqual([r['balance'] for r in book['rows']],[150000,130000])
        self.assertEqual(book['rows'][0]['reference'],'REC-1')
        self.assertEqual(len(book['opening_refs']),1)
        rows=list(ledgers.book_rows(book,'2023-04-01','2023-04-30'))
        self.assertEqual(rows[0][3],'Opening balance');self.assertEqual(rows[-1][3],'Closing balance')
        self.assertEqual(rows[-1][-2:],(130000,'Dr'))
        later=self.pack(start='2023-04-03')['books'][0]
        self.assertEqual((later['opening'],later['closing']),(150000,130000))

    def test_range_containing_initial_opening_and_empty_period(self):
        before=self.pack(start='2023-03-01')['books'][0]
        self.assertEqual(before['opening'],0);self.assertEqual(len(before['rows']),3)
        empty=self.pack(start='2024-04-01',end='2024-04-30')['books'][0]
        self.assertEqual(empty['opening'],130000);self.assertEqual(empty['closing'],130000);self.assertFalse(empty['rows'])
        unused=self.pack(codes=('1050',))['books'][0];self.assertEqual(unused['closing'],0)

    def test_branch_entity_multiple_same_account_lines_and_stable_order(self):
        core.post(self.s,'2023-04-03','Two bank debits',[dict(account='1010',branch='JP',debit=100,credit=0),dict(account='1010',branch='JP',debit=200,credit=0),dict(account='3000',branch='JP',debit=0,credit=300)])
        book=self.pack()['books'][0]
        self.assertEqual([r['debit'] for r in book['rows'][-2:]],[100,200])
        revenue=self.pack(codes=('4000',),branch='ALL')['books'][0]
        self.assertEqual(revenue['closing'],-57000)
        self.assertEqual(self.pack(codes=('4000',),branch='KM')['books'][0]['closing'],-7000)
        self.assertEqual(ledgers.balance(-57000),'570.00 Cr')
        for branch in ('JP','KM','VK','ALL'):
            books=self.pack(codes=tuple(a['code'] for a in self.s['accounts']),branch=branch)['books']
            self.assertEqual(sum(b['opening'] for b in books),0)
            self.assertEqual(sum(b['debit']-b['credit'] for b in books),0)
            self.assertEqual(sum(b['closing'] for b in books),0)

    def test_reversal_year_close_and_new_year_openings(self):
        voucher=self.s['vouchers'][2];core.reverse(self.s,voucher['id'],'2023-04-04')
        foundation.close_year(self.s,2023)
        revenue=self.pack(codes=('4000',),end='2024-03-31')['books'][0]
        self.assertEqual(revenue['closing'],0);self.assertTrue(revenue['rows'][-1]['closing_entry'])
        retained=self.pack(codes=('3100',),start='2024-04-01',end='2024-04-30',branch='ALL')['books'][0]
        self.assertEqual(retained['opening'],-57000);self.assertEqual(retained['closing'],-57000)
        self.assertEqual(core.balances(self.s,'2024-03-31','ALL','2023-04-01')['4000'],-57000)

    def test_read_only_no_drafts_duplicates_and_validation(self):
        foundation.save_draft(self.s,'2023-04-05','Unposted sale',lines('1010','4000',9900))
        original=copy.deepcopy(self.s)
        pack=self.pack(codes=('1010','1010'));self.assertEqual(len(pack['books']),1)
        self.assertEqual(self.s,original);self.assertEqual(pack['books'][0]['closing'],130000)
        for codes,start,end,branch in [((),'2023-04-01','2023-04-30','JP'),(('9999',),'2023-04-01','2023-04-30','JP'),(('1010',),'2023-05-01','2023-04-30','JP'),(('1010',),'20230401','2023-04-30','JP'),(('1010',),'2023-04-01','2023-04-30','BAD')]:
            with self.assertRaises(ValueError):self.pack(codes,start,end,branch)
        self.s['vouchers'][0]['narration']='Tampered'
        with self.assertRaises(ValueError):self.pack()

    def test_all_export_formats_preserve_values_and_escape_text(self):
        text='=HYPERLINK("https://invalid.example","fake") <script>alert(1)</script> & text'
        core.post(self.s,'2023-04-04',text,lines('1010','4000',12345),metadata={'reference':'+SUM(1,2)'})
        original=copy.deepcopy(self.s);pack=self.pack(codes=('1010','4000'))
        with tempfile.TemporaryDirectory() as directory:
            p=Path(directory)
            ledgers.export_csv(pack,p/'books.csv');ledgers.export_xlsx(pack,p/'books.xlsx');ledgers.export_html(pack,p/'books.html');ledgers.export_pdf(pack,p/'books.pdf')
            csvtext=(p/'books.csv').read_text(encoding='utf-8-sig')
            self.assertIn("'=HYPERLINK",csvtext);self.assertIn("'+SUM",csvtext);self.assertIn('1423.45',csvtext)
            html=(p/'books.html').read_text();self.assertNotIn('<script>alert(1)</script>',html);self.assertIn('&lt;script&gt;',html)
            self.assertIn('window.print()',html);self.assertIn('break-before:page',html)
            workbook=load_workbook(p/'books.xlsx',data_only=False)
            try:
                self.assertEqual(workbook.sheetnames,['1010','4000'])
                found=[c for row in workbook['1010'] for c in row if c.value==text]
                self.assertEqual(found[0].data_type,'s')
                closing=next(row for row in workbook['1010'] if row[3].value=='Closing balance')
                self.assertEqual(closing[8].value,1423.45);self.assertEqual(closing[9].value,'Dr')
                self.assertIsNotNone(workbook['1010'].print_title_rows)
            finally:workbook.close()
            reader=PdfReader(p/'books.pdf');self.assertEqual(len(reader.pages),2)
            pdftext=reader.pages[0].extract_text()
            self.assertIn('1,423.45',pdftext);self.assertIn('Closing balance',pdftext);self.assertIn('1010',pdftext)
            self.assertIn('4000',reader.pages[1].extract_text())
        self.assertEqual(self.s,original)

    def test_large_ledger_long_narration_pdf_and_small_paise(self):
        for i in range(90):core.post(self.s,'2023-04-04','Long narration '+('clinical service ' * (700 if i==0 else 3)),lines('1010','4000',1))
        pack=self.pack()
        self.assertEqual(pack['books'][0]['closing'],130090)
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'long.pdf';ledgers.export_pdf(pack,path)
            reader=PdfReader(path);self.assertGreater(len(reader.pages),2)
            texts=[page.extract_text() for page in reader.pages]
            self.assertTrue(all('1010' in t for t in texts))
            self.assertIn('1,300.90',texts[-1]);self.assertIn('Closing balance',texts[-1])
