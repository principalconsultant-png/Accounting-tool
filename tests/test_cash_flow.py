import copy
import json
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
from pypdf import PdfReader
from riju import core,foundation,cash_flow as cf,fx,history
from riju.storage import Vault


def line(code,n,dr=True,branch='JP'):return dict(account=code,branch=branch,debit=n if dr else 0,credit=0 if dr else n)
def pair(dr,cr,n,branch='JP'):return [line(dr,n,True,branch),line(cr,n,False,branch)]


class CashFlowTests(unittest.TestCase):
    def setUp(self):
        self.s=core.empty_state();core.post(self.s,'2026-04-01','Initial balances',pair('1010','3000',100000),opening=True);self.approve()
    def approve(self,**extra):
        values=dict(cf.defaults(self.s),effective='2026-04-01',note='Reviewed synthetic cash/classification basis',reviewed=True);values.update(extra);self.policy=cf.save_policy(self.s,values);return self.policy
    def result(self,branch='ALL'):return cf.analyse(self.s,'2026-04-01','2026-04-30',branch)
    def post(self,narration,lines):return core.post(self.s,'2026-04-02',narration,lines)
    def test_operating_receipts_payments_and_indirect_profit_match(self):
        self.post('Service',pair('1010','4000',20000));self.post('Rent',pair('4530','1010',5000));r=self.result()
        self.assertEqual(r['opening'],100000);self.assertEqual(r['totals']['Operating'],15000);self.assertEqual(r['profit'],15000);self.assertEqual(r['closing'],115000);self.assertTrue(r['ready'])
    def test_accrual_and_partial_patient_collection_working_capital(self):
        patient=self.s['patient_control'];self.post('Bill',pair(patient,'4000',20000));self.post('Part receipt',pair('1010',patient,10000));r=self.result()
        self.assertEqual(r['profit'],20000);self.assertEqual(r['working_capital'][patient],-10000);self.assertEqual(r['indirect_operating'],10000)
    def test_payable_expense_tax_withholding_and_payment_cash_only(self):
        self.post('Expense', [line('4510',60000),line('2000',54000,False),line('2100',6000,False)])
        self.post('Supplier payment',pair('2000','1010',54000));self.post('Tax payment',pair('2100','1010',6000));r=self.result()
        self.assertEqual(r['totals']['Operating'],-60000);self.assertEqual(r['indirect_operating'],-60000);self.assertTrue(r['ready'])
    def test_depreciation_is_separate_noncash_adjustment(self):
        self.post('Depreciation',pair('4570','1050',2000));r=self.result();self.assertEqual(r['profit'],-2000);self.assertEqual(r['noncash_profit']['4570'],2000);self.assertEqual(r['indirect_operating'],0)
    def test_investing_financing_and_interest_separate(self):
        self.post('Equipment',pair('1030','1010',10000));self.post('Loan',pair('1010','2400',30000));self.post('Interest',pair('4580','1010',1000));r=self.result()
        self.assertEqual(r['totals']['Investing'],-10000);self.assertEqual(r['totals']['Financing'],30000);self.assertEqual(r['totals']['Operating'],-1000);self.assertEqual(r['indirect_difference'],0)
    def test_noncash_capex_does_not_fabricate_cash_flow(self):
        self.post('Unpaid equipment',pair('1030','2000',10000));r=self.result();self.assertEqual(sum(r['totals'].values()),0);self.assertEqual(r['indirect_operating'],0);self.assertTrue(r['noncash_support'])
    def test_capital_creditor_payment_requires_investing_review(self):
        self.post('Unpaid equipment',pair('1030','2000',10000));payment=self.post('Pay capital creditor',pair('2000','1010',10000));self.assertFalse(self.result()['ready'])
        cf.save_allocation(self.s,self.policy['id'],payment['id'],'ALL',[dict(category='Investing',amount='-100',label='Capital creditor payment')],'Verified original equipment invoice')
        r=self.result();self.assertTrue(r['ready']);self.assertEqual(r['totals']['Investing'],-10000);self.assertEqual(r['totals']['Operating'],0)
    def test_disposal_loss_and_cash_proceeds_reconcile_after_review(self):
        self.post('Buy equipment',pair('1030','1010',10000));self.post('Depreciate',pair('4570','1050',3000))
        sale=self.post('Dispose equipment',[line('1010',5000),line('1050',3000),line('4590',2000),line('1030',10000,False)])
        self.assertFalse(self.result()['ready']);cf.save_allocation(self.s,self.policy['id'],sale['id'],'ALL',[dict(category='Investing',amount='50',label='Actual disposal proceeds')],'Verified sale settlement')
        r=self.result();self.assertEqual(r['totals']['Investing'],-5000);self.assertEqual(r['indirect_operating'],0);self.assertTrue(r['ready'])
    def test_cash_transfer_and_interbranch_entity_elimination(self):
        self.post('Cash withdrawal',pair('1000','1010',10000))
        self.post('Inter-unit funding',[line('1130',5000),line('1010',5000,False),line('1010',5000,True,'KM'),line('1130',5000,False,'KM')])
        entity=self.result();jp=self.result('JP');km=self.result('KM');self.assertEqual(sum(entity['totals'].values()),0);self.assertEqual(jp['totals']['Financing'],-5000);self.assertEqual(km['totals']['Financing'],5000)
        self.assertEqual(jp['closing']+km['closing'],entity['closing'])
    def test_cross_unit_capex_cash_is_investing_at_entity_and_branch_reconciliation_explains_it(self):
        self.post('JP pays KM equipment',[line('1130',10000),line('1010',10000,False),line('1030',10000,True,'KM'),line('1130',10000,False,'KM')])
        self.assertEqual(self.result()['totals']['Investing'],-10000);self.assertEqual(self.result('JP')['totals']['Financing'],-10000)
        pack=cf.branch_snapshot(self.s,'2026-04-01','2026-04-30');self.assertFalse(any(pack['controls'].values()));self.assertTrue(any(r[-1] for r in pack['rows'][3:]))
    def test_mixed_funding_and_gross_cash_limits(self):
        v=self.post('Part noncash loan funding',[line('1010',10000),line('1030',10000),line('2400',20000,False)]);self.assertFalse(self.result()['ready'])
        gross=[dict(category='Financing',amount='200',label='Gross loan receipt'),dict(category='Investing',amount='-100',label='Gross equipment payment')]
        with self.assertRaises(ValueError):cf.save_allocation(self.s,self.policy['id'],v['id'],'ALL',gross,'Review gross sources')
        cf.save_allocation(self.s,self.policy['id'],v['id'],'ALL',[dict(category='Financing',amount='100',label='Cash portion')],'Verified noncash portion of asset funding')
        r=self.result();self.assertTrue(r['ready']);self.assertEqual(r['totals']['Financing'],10000);self.assertEqual(r['indirect_difference'],0)
    def test_composite_zero_net_requires_review_despite_zero_proof(self):
        v=self.post('Cash offset composite',[line('1010',1000),line('1010',1000,False),line('4000',1000,False),line('4530',1000)])
        r=self.result();self.assertEqual(r['cash_difference'],0);self.assertFalse(r['ready'])
        cf.save_allocation(self.s,self.policy['id'],v['id'],'ALL',[dict(category='Operating',amount='10',label='Receipt'),dict(category='Operating',amount='-10',label='Payment')],'Verified separate underlying cash items')
        self.assertTrue(self.result()['ready'])
    def test_offset_control_does_not_get_mistaken_for_internal_transfer(self):
        v=self.post('Supplier refund/payment composite',[line('1010',1000),line('1010',1000,False),line('2000',1000),line('2000',1000,False)])
        self.assertFalse(self.result()['ready']);cf.save_allocation(self.s,self.policy['id'],v['id'],'ALL',[],'Actual bank transfer with offsetting control correction');self.assertTrue(self.result()['ready'])
    def test_internal_transfer_cannot_be_manually_fabricated_as_external_flows(self):
        v=self.post('Internal cash transfer',pair('1000','1010',1000))
        with self.assertRaises(ValueError):cf.save_allocation(self.s,self.policy['id'],v['id'],'ALL',[dict(category='Financing',amount='10'),dict(category='Financing',amount='-10')],'Unsupported external gross flows')
    def test_reviewed_allocations_inherit_exact_reversal_and_single_unit_scope(self):
        v=self.post('Mixed funding',[line('1010',10000),line('1030',10000),line('2400',20000,False)])
        cf.save_allocation(self.s,self.policy['id'],v['id'],'JP',[dict(category='Financing',amount='100')],'Verified JP cash funding')
        self.assertTrue(self.result()['ready']);core.reverse(self.s,v['id'],'2026-04-03');r=self.result();self.assertEqual(sum(r['totals'].values()),0);self.assertTrue(r['ready'])
    def test_fx_hold_conversion_exchange_effect_separate_from_operating_fees(self):
        fx.save_rate(self.s,'2026-04-02','USD','80','Invented test rate');fx.create_foreign_account(self.s,'1012','Synthetic FX cash','USD')
        fx.post_receipt(self.s,payment_date='2026-04-02',code='USD',foreign_amount='100',credit_account='4040',reference='FX-HOLD',mode='held',foreign_account='1012')
        fx.post_conversion(self.s,payment_date='2026-04-03',foreign_account='1012',foreign_amount='100',bank_account='1010',gross_inr='8500',fees_inr='10',reference='FX-CONVERT')
        self.approve();r=self.result();self.assertEqual(r['totals']['Exchange effects'],50000);self.assertEqual(r['totals']['Operating'],799000);self.assertEqual(r['cash_difference'],0);self.assertEqual(r['indirect_difference'],0)
    def test_policy_missing_and_unknown_accounts_remain_draft(self):
        other=core.empty_state();self.assertFalse(cf.analyse(other,'2026-04-01','2026-04-30')['ready'])
        foundation.save_account(self.s,'1195','Unclassified advance','Asset','Other assets');self.post('Advance',pair('1195','1010',1000));self.assertFalse(self.result()['ready'])
        proposed=cf.defaults(self.s);proposed['mapping']['1195']='Investing';self.approve(mapping=proposed['mapping']);self.assertTrue(self.result()['ready'])
    def test_opening_count_once_introductions_separate_and_year_close_excluded(self):
        self.assertEqual(self.result()['opening'],100000);self.assertEqual(self.result()['totals']['Financing'],0)
        result=cf.analyse(self.s,'2026-01-01','2026-04-30');self.assertEqual(result['introduced'],100000);self.assertFalse(result['ready']);self.assertEqual(result['cash_difference'],0)
        self.s=core.empty_state();core.post(self.s,'2024-04-01','Opening',pair('1010','3000',100000),opening=True);self.approve(effective='2024-04-01');core.post(self.s,'2024-04-02','Revenue',pair('1010','4000',5000));foundation.close_year(self.s,2024);r=cf.analyse(self.s,'2024-04-01','2025-03-31');self.assertEqual(r['profit'],5000);self.assertEqual(r['totals']['Operating'],5000);self.assertEqual(r['indirect_difference'],0)
    def test_readonly_outputs_numeric_excel_pdf_and_noncash_policy_support(self):
        self.post('Revenue',pair('1010','4000',5000));self.post('Depreciation',pair('4570','1050',1000));before=copy.deepcopy(self.s)
        with tempfile.TemporaryDirectory() as d:
            root=Path(d)
            for kind in ('xlsx','pdf','json'):cf.export(self.s,'2026-04-01','2026-04-30','ALL',root/('cash.'+kind),kind)
            w=load_workbook(root/'cash.xlsx');self.assertIn(cf.DIRECT,w.sheetnames);self.assertIn('Noncash and mixed voucher notes',w.sheetnames);self.assertIn('Cash-flow classification policy',w.sheetnames);w.close()
            text=' '.join(p.extract_text() for p in PdfReader(root/'cash.pdf').pages);self.assertIn('REVIEWED',text);self.assertIn('Opening cash',text)
        self.assertEqual(before,self.s)
    def test_format_ten_upgrade_backup_and_policy_allocation_persistence(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);ext=root/'external';ext.mkdir();old=copy.deepcopy(self.s);old['version']=10;old['schema_revision']=10;old.pop('cash_flow_policies');old.pop('cash_flow_allocations');history.checkpoint(old)
            path=root/'old.json';path.write_text(json.dumps(old),encoding='utf-8');raw=path.read_bytes();v=Vault(path,demo=True);self.s=v.load();self.s['settings']['backup_dir']=str(ext);v.save(self.s);self.assertEqual(v.last_upgrade_backup.read_bytes(),raw);self.assertEqual(self.s['version'],20)
            self.approve();v.save(self.s);self.assertEqual(v.load(),self.s)

    def test_manual_fx_revaluation_effect_and_fee_reversal_keep_bridge(self):
        fx.save_rate(self.s,'2026-04-02','USD','80','Invented');fx.create_foreign_account(self.s,'1012','Synthetic held cash','USD')
        fx.post_receipt(self.s,payment_date='2026-04-02',code='USD',foreign_amount='10',credit_account='4040',reference='FXR',mode='held',foreign_account='1012')
        converted=fx.post_conversion(self.s,payment_date='2026-04-03',foreign_account='1012',foreign_amount='10',bank_account='1010',gross_inr='850',fees_inr='1',reference='FXC')
        self.approve();core.reverse(self.s,converted['id'],'2026-04-04');r=self.result();self.assertEqual(r['totals']['Exchange effects'],0);self.assertEqual(r['totals']['Operating'],80000);self.assertEqual(r['indirect_difference'],0)
    def test_cash_policy_revisions_do_not_overwrite_earlier_period_policy(self):
        old=self.policy;self.approve(effective='2026-05-01',cash_accounts=['1000'])
        self.assertEqual(cf.policy(self.s,'2026-04-30')['id'],old['id']);self.assertEqual(cf.policy(self.s,'2026-05-31')['cash_accounts'],['1000'])
    def test_invalid_policy_and_allocation_leave_history_unchanged(self):
        v=self.post('Mixed',[line('1010',10000),line('1030',10000),line('2400',20000,False)]);before=copy.deepcopy(self.s)
        with self.assertRaises(ValueError):cf.save_policy(self.s,dict(cf.defaults(self.s),effective='2026-04-01',cash_accounts=['4000'],note='Invalid cash definition',reviewed=True))
        with self.assertRaises(ValueError):cf.save_allocation(self.s,self.policy['id'],v['id'],'ALL',[dict(category='Financing',amount='99')],'Does not balance')
        self.assertEqual(before,self.s)
    def test_policy_referenced_master_cannot_be_deleted(self):
        foundation.save_account(self.s,'1196','Synthetic asset','Asset','Other assets');self.approve();foundation.save_account(self.s,'1196','Synthetic asset','Asset','Other assets',active=False)
        with self.assertRaises(ValueError):foundation.delete_account(self.s,'1196')
