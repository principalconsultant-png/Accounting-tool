import copy
import csv
import io
import tempfile
import time
import unittest
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader
from riju import core,receipt_import as r,activity,financial_reports as f,sales_import


def row(**kw):
    v=dict.fromkeys(r.HEADERS,'');v.update(receipt_date='2026-11-01',branch='JP',receipt_no='R-001',record_type='earned',tax_mode='no_separate_tax',consultation='500',treatment='300',cash='200',upi='600',total='800');v.update(kw);return v

def raw(rows):
    out=io.StringIO();w=csv.DictWriter(out,fieldnames=r.HEADERS);w.writeheader();w.writerows(rows);return out.getvalue().encode()

class ReceiptImportTests(unittest.TestCase):
    def setUp(self):self.s=core.empty_state()
    def stage(self,rows):return r.stage(self.s,raw(rows),'receipts.csv')
    def approve(self,q,indices=None):r.review(self.s,q['id'],list(range(len(q['rows']))) if indices is None else indices,True,'Reviewed actual earned receipts/tax/payments',True)
    def test_balanced_splits_bank_and_no_post_on_upload(self):
        q=self.stage([row(party='Source label only'),row(receipt_no='R-002',bank='800',cash='',upi='')]);self.assertFalse(self.s['vouchers']);self.approve(q);self.assertEqual(sum(x[3] for x in r.totals_snapshot(self.s,q['id'])['rows'] if x[0]=='Revenue'),160000);self.assertEqual(r.post(self.s,q['id']),2)
        b=core.balances(self.s,'2026-11-30');self.assertEqual((b['1000'],b['1010'],b['1110'],b['4000'],b['4010']),(20000,80000,60000,-100000,-60000));self.assertEqual(activity.report(self.s)['metrics']['import_rows_processed'],2)
        with self.assertRaises(ValueError):r.post(self.s,q['id'])
    def test_bad_rows_retained_bulk_decisions_and_corrected_reimport(self):
        q=self.stage([row(),row(receipt_no='BAD',total='799'),row(receipt_no='BAD2',tax_mode='Taxable'),row(receipt_no='R-001')]);self.assertEqual([x['validation'] for x in q['rows']],['VALID','INVALID','INVALID','DUPLICATE'])
        before=copy.deepcopy(self.s)
        with self.assertRaises(ValueError):self.approve(q,[0,1])
        self.assertEqual(before,self.s);self.approve(q,[0])
        with self.assertRaises(ValueError):r.post(self.s,q['id'])
        r.review(self.s,q['id'],[1,2,3],False,'Rejected: source report correction required',True);self.assertEqual(r.post(self.s,q['id']),1)
        corrected=self.stage([row(receipt_no='BAD'),row(receipt_no='BAD2')]);self.assertTrue(all(x['validation']=='VALID' for x in corrected['rows']))
    def test_duplicate_identity_across_changed_date_files_units_fiscal_year(self):
        q=self.stage([row()]);self.approve(q);r.post(self.s,q['id'])
        edited=self.stage([row(receipt_date='2026-11-02',consultation='501',total='801',upi='601')]);self.assertEqual(edited['rows'][0]['validation'],'DUPLICATE')
        unique=self.stage([row(branch='KM'),row(receipt_date='2027-11-01')]);self.assertTrue(all(x['validation']=='VALID' for x in unique['rows']))
        with self.assertRaises(ValueError):r.stage(self.s,raw([row()]),'renamed.csv')
    def test_rejects_formula_rounding_negative_advance_and_missing_method(self):
        inputs=[dict(consultation='=SUM(1,2)'),dict(cash='-1'),dict(consultation='500.001'),dict(record_type='advance'),dict(branch='UNKNOWN'),dict(receipt_no=''),dict(receipt_date='2026-02-30'),dict(upi=''),dict(total='NaN')]
        q=self.stage([row(receipt_no=str(i),**v) if 'receipt_no' not in v else row(**v) for i,v in enumerate(inputs)]);self.assertTrue(all(x['validation']=='INVALID' for x in q['rows']));self.assertFalse(self.s['vouchers'])
    def test_atomic_post_after_period_lock_and_rejected_source_no_false_reservation(self):
        q=self.stage([row(),row(receipt_no='R2',receipt_date='2026-12-01')]);self.approve(q);self.s['locked_through']='2026-11-30';before=copy.deepcopy(self.s)
        with self.assertRaises(ValueError):r.post(self.s,q['id'])
        self.assertEqual(before,self.s)
    def test_csv_xlsx_template_and_excel_date_roundtrip(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);r.template(root/'template.csv');self.assertEqual(next(csv.reader(io.StringIO((root/'template.csv').read_text(encoding='utf-8-sig')))),list(r.HEADERS))
            r.template(root/'template.xlsx');wb=load_workbook(root/'template.xlsx');self.assertEqual(wb.sheetnames,['Receipts','Instructions']);wb['Receipts'].append([row()[h] for h in r.HEADERS]);wb.save(root/'source.xlsx');wb.close()
            q=r.stage(self.s,(root/'source.xlsx').read_bytes(),'source.xlsx');self.assertEqual(q['rows'][0]['validation'],'VALID')
            pack=r.snapshot(self.s,q['id']);f.export_xlsx(pack,root/'review.xlsx');f.export_pdf(pack,root/'review.pdf');self.assertIn('R-001',' '.join(p.extract_text() for p in PdfReader(root/'review.pdf').pages))
    def test_legacy_and_standard_cross_duplicate_guard(self):
        from test_sales import workbook
        q=sales_import.stage_sales(self.s,workbook(),'legacy.xlsx','JP');source=q['rows'][0]
        new=self.stage([row(receipt_date=source['date'],branch=source['branch'],receipt_no=source['receipt'])]);self.assertEqual(new['rows'][0]['validation'],'DUPLICATE')
        self.s=core.empty_state();self.stage([row(receipt_date=source['date'],branch=source['branch'],receipt_no=source['receipt'])])
        with self.assertRaises(ValueError):sales_import.stage_sales(self.s,workbook(),'legacy.xlsx','JP')
    def test_600_receipts_bulk_post_source_counts_and_cpu_budget(self):
        started=time.perf_counter();q=self.stage([row(receipt_no=f'BULK-{i:04}') for i in range(600)]);self.approve(q);self.assertEqual(r.post(self.s,q['id']),600);self.assertEqual(len(self.s['vouchers']),600);self.assertLess(time.perf_counter()-started,15)
        self.assertEqual(sum(core.balances(self.s,'2026-11-30').values()),0);self.assertEqual(activity.report(self.s)['metrics']['import_rows_processed'],600)
