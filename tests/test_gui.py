"""Run with a graphical display: DISPLAY=:99 python -m unittest discover -s tests -v."""
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from riju.storage import Vault

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class DesktopSmokeTests(unittest.TestCase):
    def test_screens_and_review_post_flow(self):
        from riju.app import App
        with tempfile.TemporaryDirectory() as d:
            app=App(Vault(Path(d)/'demo.json',demo=True))
            try:
                for page in app.pages:
                    app.show(page);app.update()
                app.show('Journal studio');app.update()
                app.jdate.set('2026-04-01')
                app.narration.insert('1.0','Paid ₹5000 rent by bank')
                app.prepare()
                self.assertEqual(len(app.draft),2)
                with patch('riju.app.messagebox.askyesno',return_value=True),patch('riju.app.messagebox.showinfo'):
                    app.post_draft()
                self.assertEqual(len(app.state['vouchers']),1)
                app.show('Financial reports');app.update()
                restored=Vault(Path(d)/'demo.json',demo=True).load()
                self.assertEqual(restored['vouchers'][0]['lines'][0]['debit'],500000)
                for description in ('Medicines purchased from our Pharmacy by Patient Hema R in Cash INR1291', 'Received cash INR1291 from Patient Hema R for Medicines.'):
                    app.show('Journal studio');app.update()
                    app.jdate.set('2026-10-29')
                    app.narration.insert('1.0',description)
                    app.prepare()
                    self.assertEqual([l['account'] for l in app.draft],['1000','4030'])
                    self.assertEqual(app.draft[0]['debit'],129100)
                    with patch('riju.app.messagebox.askyesno',return_value=True),patch('riju.app.messagebox.showinfo'):
                        app.post_draft()
                self.assertEqual(len(app.state['vouchers']),3)
                from test_sales import workbook
                from riju import core
                import tkinter as tk
                def widgets(parent):
                    for w in parent.winfo_children():
                        yield w
                        yield from widgets(w)
                def invoke(label):
                    next(w for w in widgets(app.content) if w.winfo_class()=='TButton' and w.cget('text')==label).invoke()
                app.show('Sales receipt');app.update()
                entries=[w for w in widgets(app.content) if w.winfo_class()=='TEntry']
                self.assertEqual(len(entries),21)
                for field in (entries[2],entries[-4]):field.delete(0,'end');field.insert(0,'100')
                invoke('Prepare compound journal for review')
                self.assertEqual(app.current,'Journal studio')
                self.assertEqual([l['account'] for l in app.draft],['1000','4010'])
                source=Path(d)/'synthetic.xlsx';source.write_bytes(workbook())
                app.show('Import queue');app.update()
                with patch('riju.app.filedialog.askopenfilename',return_value=str(source)),patch('riju.app.messagebox.showinfo'):
                    invoke('Upload daily sales XLSX')
                self.assertEqual(len(app.state['queue']),1)
                tree=next(w for w in widgets(app.content) if w.winfo_class()=='Treeview')
                tree.selection_set(tree.get_children()[0])
                invoke('Review rows');app.update()
                for child in app.winfo_children():
                    if isinstance(child,tk.Toplevel):child.destroy()
                with patch('riju.app.messagebox.askyesno',return_value=True):invoke('Approve & post entire file')
                self.assertEqual(app.state['queue'][0]['status'],'PROCESSED')
                self.assertEqual(len(app.state['vouchers']),4)
                core.verify_audit(app.state)
            finally: app.destroy()

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class FoundationDesktopTests(unittest.TestCase):
    def buttons(self,widget,text):
        from tkinter import ttk
        for child in widget.winfo_children():
            if isinstance(child,ttk.Button) and child.cget('text')==text:yield child
            yield from self.buttons(child,text)

    def click(self,app,text):
        next(self.buttons(app.content,text)).invoke();app.update()

    def test_optional_entity_branch_ledger_and_saved_draft_roundtrip(self):
        from riju.app import App
        from riju import core
        from tkinter import ttk
        with tempfile.TemporaryDirectory() as directory:
            app=App(Vault(Path(directory)/'demo.json',demo=True))
            errors=[];app.report_callback_exception=lambda *args:errors.append(args)
            try:
                with patch('riju.app.messagebox.showinfo'),patch('riju.foundation_ui.messagebox.showinfo'),patch('riju.app.messagebox.askyesno',return_value=True):
                    app.show('Setup & periods');app.update()
                    app.entity_name.set('Test hospital');self.click(app,'Save entity details')
                    self.assertEqual(app.state['settings']['entity'],'Test hospital')
                    app.entity_fields['tan'].set('BLRA12345B');self.click(app,'Save entity details')
                    self.assertEqual(app.state['settings']['registration']['tan'],'BLRA12345B')
                    app.branch_fields['code'].set('NEW');app.branch_fields['name'].set('New branch');self.click(app,'Save branch')
                    self.assertIn('New branch',app.branch_names)
                    app.setup_year.set('2022');self.click(app,'Add financial year')
                    self.assertEqual(app.year.get(),'2022')
                    app.show('Chart of accounts');app.update()
                    app.ledger_code.set('1011');app.ledger_name.set('Second bank');app.ledger_kind.set('Asset');app.ledger_group.set('Cash and bank')
                    self.click(app,'Save ledger')
                    self.assertTrue(any(a['code']=='1011' for a in app.state['accounts']))
                    app.show('Journal studio');app.update()
                    app.jdate.set('2022-06-01');app.jbranch.set('New branch');app.jreference.set('R-42');app.jtype.set('Receipt')
                    app.narration.insert('1.0','Received cash INR100 for consultation');app.prepare();app.save_current_draft()
                    self.assertFalse(app.state['vouchers']);self.assertEqual(len(app.state['drafts']),1)
                    app.show('Saved drafts');app.update()
                    # Choose the saved draft in its tree, then use the actual desktop edit/post actions.
                    def trees(w):
                        for c in w.winfo_children():
                            if isinstance(c,ttk.Treeview):yield c
                            yield from trees(c)
                    tree=next(trees(app.content));tree.selection_set(tree.get_children()[0]);self.click(app,'Open & edit draft')
                    self.assertEqual(app.jreference.get(),'R-42');self.assertEqual(app.jbranch.get(),'New branch')
                    app.post_draft();self.assertFalse(app.state['drafts']);self.assertEqual(len(app.state['vouchers']),1)
                    self.assertEqual(app.state['vouchers'][0]['lines'][0]['branch'],'NEW')
                    self.assertEqual(app.state['vouchers'][0]['metadata']['reference'],'R-42')
                    app.geometry('1000x680')
                    for page in app.pages:
                        app.show(page);app.update();app.page_canvas.yview_moveto(1);app.update()
                    self.assertFalse(errors,errors)
                    core.verify_audit(app.state)
            finally:app.destroy()

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class LedgerDesktopTests(unittest.TestCase):
    def test_navigation_account_filters_exports_and_voucher(self):
        from riju.app import App
        from riju import core
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            p=Path(directory);app=App(Vault(p/'test.json',demo=True))
            errors=[];app.report_callback_exception=lambda *args:errors.append(args)
            try:
                app.commit(lambda s:core.post(s,'2023-04-01','Synthetic opening',[dict(account='1010',branch='JP',debit=100000,credit=0),dict(account='3000',branch='JP',debit=0,credit=100000)],opening=True))
                app.commit(lambda s:core.post(s,'2023-04-02','Consultation receipts',[dict(account='1010',branch='JP',debit=50000,credit=0),dict(account='4000',branch='JP',debit=0,credit=50000)],metadata={'reference':'R-TEST'}))
                self.assertEqual(list(app.pages)[:2],['Overview','Ledger A/c'])
                app.year.set('2023');app.nav_buttons['Ledger A/c'].invoke();app.update()
                self.assertIn('Patient A/C',' '.join(app.ledger_list.get(0,'end')))
                app.ledger_category.set('Bank books');app.update();self.assertIn('1010',app.selected_ledger_codes())
                self.assertIn('1,500.00 Dr',app.ledger_summary.cget('text'))
                transaction=next(item for item in app.ledger_table.get_children() if item in app.ledger_row_vouchers)
                app.ledger_table.selection_set(transaction)
                with patch('riju.ledger_ui.messagebox.showinfo') as info:
                    app.ledger_voucher();self.assertIn('R-TEST',info.call_args.args[1])
                with patch('riju.ledger_ui.filedialog.asksaveasfilename',return_value=str(p/'ledger.xlsx')),patch('riju.ledger_ui.messagebox.showinfo'):
                    app.export_ledgers('xlsx')
                wb=load_workbook(p/'ledger.xlsx');self.assertEqual(wb.sheetnames,['1010']);wb.close()
                with patch('riju.ledger_ui.filedialog.asksaveasfilename',return_value=str(p/'ledger.html')),patch('riju.ledger_ui.webbrowser.open',return_value=True) as browser:
                    app.export_ledgers('html');browser.assert_called_once();self.assertTrue((p/'ledger.html').is_file())
                app.ledger_category.set('All ledgers');app.select_all_ledgers();self.assertEqual(len(app.selected_ledger_pack()['books']),len(app.state['accounts']))
                app.ledger_search.set('does not exist');app.update();self.assertFalse(app.selected_ledger_codes())
                with self.assertRaises(ValueError):app.export_ledgers('pdf')
                app.ledger_search.set('1010');app.ledger_from.set('2024-04-01');app.ledger_to.set('2024-04-30');app.render_ledger()
                self.assertIn('0 lines',app.ledger_summary.cget('text'));self.assertIn('1,500.00 Dr',app.ledger_summary.cget('text'))
                app.branch.set('Kammanahalli');app.show('Ledger A/c');app.ledger_category.set('Bank books');app.update()
                self.assertEqual(app.ledger_from.get(),'2024-04-01')
                self.assertEqual(app.selected_ledger_codes(),['1010'])
                self.assertIn('0.00',app.ledger_summary.cget('text'))
                self.assertFalse(errors,errors)
            finally:app.destroy()

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class FxWorkDesktopTests(unittest.TestCase):
    def test_fx_rate_preview_post_and_work_statistics(self):
        from riju.app import App
        from riju import fx,activity,core,ledgers
        with tempfile.TemporaryDirectory() as directory:
            p=Path(directory);app=App(Vault(p/'test.json',demo=True))
            errors=[];app.report_callback_exception=lambda *args:errors.append(args)
            try:
                app.commit(lambda s:fx.save_rate(s,'2026-06-01','USD','83.25','Synthetic bank advice'))
                app.show('FX receipts');app.update()
                app.fx_input['payment_date'].set('2026-06-01');app.fx_input['foreign_amount'].set('100')
                app.fx_input['reference'].set('FX-DESKTOP-1')
                self.assertTrue(app.fx_input['credit_account'].get().startswith('2300'))
                app.fx_preview_receipt();self.assertEqual(len(app.state['vouchers']),0)
                self.assertIn('8,325.00',app.fx_summaries['receipt'].cget('text'))
                app.fx_input['foreign_amount'].set('101')
                with self.assertRaises(ValueError):app.fx_post_receipt()
                app.fx_input['foreign_amount'].set('100')
                with patch('riju.operations_ui.messagebox.askyesno',return_value=True),patch('riju.operations_ui.messagebox.showinfo'):
                    app.fx_post_receipt()
                self.assertEqual(len(app.state['vouchers']),1)
                self.assertEqual(app.state['vouchers'][0]['lines'][0]['branch'],'VK')
                result=activity.report(app.state)['metrics'];self.assertEqual(result['fx_receipts_posted'],1)
                self.assertEqual(result['journal_vouchers_created'],1);self.assertEqual(result['journal_vouchers_posted'],1)
                pack=ledgers.ledger_pack(app.state,['1010'],'2026-06-01','2026-06-30','VK')
                self.assertIn('USD 100',pack['books'][0]['rows'][0]['narration'])
                app.show('Work volume');app.update();self.assertTrue(app.work_table.get_children())
                app.branch.set('JP Nagar');app.show('Work volume');app.update()
                self.assertEqual(activity.report(app.state,branch='JP')['metrics']['journal_vouchers_posted'],0)
                app.show('Overview');app.update();self.assertIn('0 journal',app.work_overview_label.cget('text'))
                app.branch.set('All operating units');app.show('Overview');app.update();self.assertIn('1 journal',app.work_overview_label.cget('text'))
                self.assertFalse(errors,errors);core.verify_audit(app.state)
            finally:app.destroy()

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class ForeignHoldingDesktopTests(unittest.TestCase):
    def test_holding_conversion_and_reload_statistics(self):
        from riju.app import App
        from riju import fx,activity
        with tempfile.TemporaryDirectory() as directory:
            p=Path(directory)/'test.json';app=App(Vault(p,demo=True))
            try:
                app.commit(lambda s:fx.save_rate(s,'2026-06-01','USD','83','Synthetic bank'))
                app.commit(lambda s:fx.create_foreign_account(s,'1012','USD bank Vaikom','USD'))
                app.show('FX receipts');app.update()
                for key,value in dict(payment_date='2026-06-01',foreign_amount='100',reference='HELD-UI-1',mode='Hold foreign currency').items():app.fx_input[key].set(value)
                app.fx_preview_receipt();self.assertIn('No INR bank conversion',app.fx_summaries['receipt'].cget('text'))
                with patch('riju.operations_ui.messagebox.askyesno',return_value=True),patch('riju.operations_ui.messagebox.showinfo'):
                    app.fx_post_receipt()
                    for key,value in dict(payment_date='2026-06-02',foreign_amount='50',gross_inr='4200',fees_inr='10',reference='CONV-UI-1').items():app.fx_conversion[key].set(value)
                    app.fx_preview_conversion();app.fx_post_conversion()
                self.assertEqual(fx.holding_balance(app.state,'1012')['foreign_amount'],'50')
                self.assertEqual(fx.holding_balance(app.state,'1012')['inr_minor'],415000)
            finally:app.destroy()
            reopened=App(Vault(p,demo=True))
            try:
                m=activity.report(reopened.state)['metrics'];self.assertEqual(m['journal_vouchers_posted'],2);self.assertEqual(m['fx_conversions_posted'],1)
                reopened.show('Work volume');reopened.update()
                self.assertTrue(reopened.work_table.get_children())
            finally:reopened.destroy()

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class BankReconciliationDesktopTests(unittest.TestCase):
    def test_bank_review_finalize_and_reload(self):
        from riju.app import App
        from riju import core,reconciliation as r
        with tempfile.TemporaryDirectory() as d:
            vault=Vault(Path(d)/'demo.json',demo=True);app=App(vault)
            try:
                def prepare(state):
                    v=core.post(state,'2026-04-02','Synthetic receipt',[dict(account='1010',branch='JP',debit=10000,credit=0),dict(account='4000',branch='JP',debit=0,credit=10000)])
                    statement=r.import_statement(state,b'date,narration,reference,amount\n2026-04-02,Receipt,T1,100\n','bank.csv','1010','ALL','2026-04-01','2026-04-30','0','100')
                    return v,statement
                # Exercise the same durable mutation path used by the desktop dialogs.
                app.commit(prepare)
                app.show('Bank reconciliation');app.update()
                self.assertEqual(len(app.brs_bank_table.get_children()),1)
                self.assertEqual(len(app.brs_book_table.get_children()),1)
                sid=app.state['bank_statements'][0]['id'];app.brs_selected_statement=sid
                app.show('Bank reconciliation');app.update()
                app.brs_bank_table.selection_set(app.brs_bank_table.get_children()[0])
                app.brs_book_table.selection_set(app.brs_book_table.get_children()[0])
                with patch('riju.bank_ui.simpledialog.askstring',return_value=''),patch('riju.bank_ui.messagebox.askyesno',return_value=True):
                    app.match_bank_books()
                self.assertTrue(r.brs(app.state,sid)['ready'])
                with patch('riju.bank_ui.messagebox.askyesno',return_value=True):app.finalize_bank_statement()
                self.assertIn('RECONCILED',app.brs_report_label.cget('text'))
            finally:app.destroy()
            loaded=vault.load();self.assertEqual(loaded['bank_statements'][0]['status'],'RECONCILED')
            self.assertEqual(len(loaded['vouchers']),1)

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class HistoryDesktopTests(unittest.TestCase):
    def test_history_filter_details_export_and_durable_change_set(self):
        from riju.app import App
        from riju import foundation,history as h
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);vault=Vault(root/'demo.json',demo=True);app=App(vault)
            try:
                app.commit(lambda s:foundation.save_entity(s,'History test hospital',{}))
                app.show('Change history');app.update()
                app.history_action.set('CHANGE_SET');app.refresh_history()
                self.assertEqual(len(app.history_rows),1)
                item=app.history_table.get_children()[0];app.history_table.selection_set(item);app.show_history_detail()
                self.assertIn('History test hospital',app.history_detail.get('1.0','end'))
                self.assertIn('before',app.history_detail.get('1.0','end'))
                app.history_query.set('History test hospital')
                with patch('riju.history_ui.filedialog.asksaveasfilename',return_value=str(root/'history.csv')),patch('riju.history_ui.messagebox.showinfo'):
                    app.export_history('csv')
                self.assertTrue((root/'history.csv').exists())
                with patch.object(app,'_history_dialog') as dialog:
                    app.check_history_integrity();self.assertTrue(dialog.call_args.args[1]['passed'])
                self.assertTrue(h.verify_saved_state(app.state))
            finally:app.destroy()
            restored=vault.load();self.assertTrue(h.integrity_report(restored)['passed'])

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class TaxSetupDesktopTests(unittest.TestCase):
    def test_create_version_party_preview_export_and_reload(self):
        import tkinter as tk
        from tkinter import ttk
        from riju.app import App
        from riju import tax
        def children(widget):
            for child in widget.winfo_children():
                yield child
                yield from children(child)
        def set_field(win,label,value):
            item=next(w for w in children(win) if isinstance(w,ttk.Label) and w.cget('text')==label)
            entry=next(w for w in item.master.winfo_children() if isinstance(w,ttk.Entry))
            if isinstance(entry,ttk.Combobox):entry.set(value)
            else:entry.delete(0,'end');entry.insert(0,value)
        def press(win,text):next(w for w in children(win) if isinstance(w,ttk.Button) and w.cget('text')==text).invoke()
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);vault=Vault(root/'demo.json',demo=True);app=App(vault)
            try:
                app.show('Tax setup');app.update();app.edit_tax_rule();app.update()
                win=next(w for w in app.winfo_children() if isinstance(w,tk.Toplevel))
                for label,value in [('Stable rule key','GUI-TEST'),('New effective date','2024-04-01'),('Rule label','Invented service'),('Total rate (%)','18'),('Specific applicability / conditions','Synthetic test only'),('HSN/SAC (TDS: leave blank)','999311')]:set_field(win,label,value)
                with patch('riju.app.messagebox.showerror') as errors:press(win,'Save new version');errors.assert_not_called()
                app.update();self.assertEqual(len(app.tax_rule_table.get_children()),1)
                app.tax_rule_table.selection_set(app.tax_rule_table.get_children()[0]);app.version_tax_rule();app.update()
                win=next(w for w in app.winfo_children() if isinstance(w,tk.Toplevel))
                set_field(win,'New effective date','2025-04-01');set_field(win,'Total rate (%)','12')
                with patch('riju.app.messagebox.showerror') as errors:press(win,'Save new version');errors.assert_not_called()
                app.update();self.assertEqual(len(app.tax_rule_table.get_children()),2)
                self.assertEqual(tax.resolve_rule(app.state,'GUI-TEST','2024-05-01')['rate'],'18')
                app.edit_tax_party();app.update();win=next(w for w in app.winfo_children() if isinstance(w,tk.Toplevel))
                set_field(win,'Party / payee name','Synthetic supplier');set_field(win,'New effective date','2024-04-01')
                with patch('riju.app.messagebox.showerror') as errors:press(win,'Save new version');errors.assert_not_called()
                app.update();self.assertEqual(len(app.tax_party_table.get_children()),1)
                app.preview_tax_rule();app.update();win=next(w for w in app.winfo_children() if isinstance(w,tk.Toplevel))
                set_field(win,'Transaction date','2025-04-01');press(win,'Preview')
                text=next(w for w in children(win) if isinstance(w,tk.Text));self.assertIn('12',text.get('1.0','end'));self.assertIn('not been marked reviewed',text.get('1.0','end'));win.destroy()
                with patch('riju.tax_ui.filedialog.asksaveasfilename',return_value=str(root/'tax.json')),patch('riju.tax_ui.messagebox.showinfo'):app.export_taxes()
                self.assertTrue((root/'tax.json').exists());self.assertEqual(len(app.state['vouchers']),0)
            finally:app.destroy()
            loaded=vault.load();self.assertEqual(len(tax.rules(loaded)),2);self.assertEqual(len(tax.parties(loaded)),1)

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class GstDesktopTests(unittest.TestCase):
    def test_review_post_summary_2b_and_reload(self):
        import tkinter as tk
        from tkinter import ttk
        from riju.app import App
        from riju import tax,gst
        def walk(w):
            for child in w.winfo_children():
                yield child
                yield from walk(child)
        def set_field(win,label,value):
            item=next(w for w in walk(win) if isinstance(w,ttk.Label) and w.cget('text')==label)
            e=next(w for w in item.master.winfo_children() if isinstance(w,ttk.Entry))
            if isinstance(e,ttk.Combobox):e.set(value)
            else:e.delete(0,'end');e.insert(0,value)
        def press(win,text):next(w for w in walk(win) if isinstance(w,ttk.Button) and w.cget('text')==text).invoke()
        with tempfile.TemporaryDirectory() as d:
            vault=Vault(Path(d)/'demo.json',demo=True);app=App(vault)
            try:
                app.commit(lambda s:tax.save_rule(s,dict(family='GST',key='GUI-GST',label='Synthetic test GST',effective='2026-04-01',rate='18',treatment='Taxable',hsn_sac='999311',applicability='Invented tests only',source='Synthetic test',reviewed=True)))
                app.show('GST accounting');app.update();app.new_gst_invoice();app.update()
                win=next(w for w in app.winfo_children() if isinstance(w,tk.Toplevel))
                for label,value in [('Invoice date','2026-04-02'),('Document kind','Purchase'),('Our GSTIN for this invoice','29ABCDE1234F1Z5'),('Place-of-supply state code','29'),('Customer/supplier GSTIN (optional B2C)','29ZZZZZ1234Z1Z5'),('Invoice/reference number','GUI-P1'),('INR amount','100'),('Revenue/cost ledger code','4590'),('Purchase ITC assessment','Eligible'),('Narration','Invented GST purchase')]:set_field(win,label,value)
                with patch('riju.app.messagebox.showerror') as errors:
                    press(win,'Prepare journal preview');errors.assert_not_called()
                    # Changing a field invalidates the reviewed preview.
                    set_field(win,'INR amount','200');press(win,'Approve & post');self.assertTrue(errors.called)
                    set_field(win,'INR amount','100');press(win,'Prepare journal preview')
                with patch('riju.gst_ui.messagebox.askyesno',return_value=True):press(win,'Approve & post')
                app.update();self.assertEqual(len(app.state['gst_documents']),1);self.assertEqual(len(app.gst_docs.get_children()),1)
                self.assertFalse(any(app.gst_report['unclassified'].values()))
                raw=b'supplier_gstin,invoice,date,taxable,cgst,sgst,igst,itc_available\n29ZZZZZ1234Z1Z5,GUI-P1,2026-04-02,100,9,9,0,Y\n'
                app.commit(lambda s:gst.import_2b(s,raw,'2b.csv','29ABCDE1234F1Z5','2026-04'));app.show('GST accounting');app.update()
                app.gst_import_table.selection_set(app.gst_import_table.get_children()[0]);app.review_gst_2b();app.update()
                win=next(w for w in app.winfo_children() if isinstance(w,tk.Toplevel));table=next(w for w in walk(win) if isinstance(w,ttk.Treeview) and 'Match status' in w['columns'])
                table.selection_set(table.get_children()[0])
                with patch('riju.gst_ui.messagebox.askyesno',return_value=True):press(win,'Approve exact match')
                self.assertEqual(len(app.state['gst_2b_matches']),1);self.assertEqual(len(app.state['vouchers']),1);win.destroy()
                with patch('riju.gst_ui.filedialog.asksaveasfilename',return_value=str(Path(d)/'gst.csv')):app.export_gst_summary()
                self.assertIn('GUI-P1',(Path(d)/'gst.csv').read_text(encoding='utf-8-sig'))
            finally:app.destroy()
            loaded=vault.load();self.assertEqual(len(loaded['gst_documents']),1);self.assertEqual(len(loaded['gst_2b_matches']),1)

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class GstNoteDesktopTests(unittest.TestCase):
    def test_refund_note_preview_approval_and_reload(self):
        import tkinter as tk
        from tkinter import ttk
        from riju.app import App
        from riju import tax,gst
        from test_gst import values
        def walk(w):
            for child in w.winfo_children():
                yield child
                yield from walk(child)
        def set_field(win,label,value):
            item=next(w for w in walk(win) if isinstance(w,ttk.Label) and w.cget('text')==label)
            entry=next(w for w in item.master.winfo_children() if isinstance(w,ttk.Entry))
            if isinstance(entry,ttk.Combobox):entry.set(value)
            else:entry.delete(0,'end');entry.insert(0,value)
        def press(win,text):next(w for w in walk(win) if isinstance(w,ttk.Button) and w.cget('text')==text).invoke()
        with tempfile.TemporaryDirectory() as d:
            vault=Vault(Path(d)/'demo.json',demo=True);app=App(vault)
            try:
                app.commit(lambda s:tax.save_rule(s,dict(family='GST',key='TEST-GST',effective='2026-04-01',label='Test',rate='18',treatment='Taxable',hsn_sac='999311',applicability='Synthetic tests',source='Invented',reviewed=True)))
                app.commit(lambda s:gst.post(s,values()));app.show('GST accounting');app.update();app.new_gst_note();app.update()
                win=next(w for w in app.winfo_children() if isinstance(w,tk.Toplevel))
                for label,value in [('Note date','2026-05-01'),('Unique note reference','GUI-CN'),('Adjustment base INR (excludes GST)','25'),('Refund / balance-adjustment ledger','1010'),('Reason / narration','Invented partial refund')]:set_field(win,label,value)
                with patch('riju.app.messagebox.showerror') as error:
                    press(win,'Approve & post note');self.assertTrue(error.called);error.reset_mock()
                    press(win,'Prepare note preview');error.assert_not_called()
                with patch('riju.gst_ui.messagebox.askyesno',return_value=False):press(win,'Approve & post note')
                self.assertEqual(len(app.state['gst_documents']),1)
                with patch('riju.gst_ui.messagebox.askyesno',return_value=True):press(win,'Approve & post note')
                app.update();self.assertEqual(len(app.state['gst_documents']),2)
                self.assertEqual(app.gst_report['totals']['sales_base'],7500);self.assertFalse(any(app.gst_report['unclassified'].values()))
                self.assertTrue(any('Credit note' in app.gst_docs.item(i,'values') for i in app.gst_docs.get_children()))
            finally:app.destroy()
            loaded=vault.load();self.assertEqual(loaded['gst_documents'][-1]['original_reference'],'TEST-1');self.assertEqual(len(loaded['vouchers']),2)

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class Gst2bNoteDesktopTests(unittest.TestCase):
    def test_typed_note_import_match_signed_export_and_reopen(self):
        import copy
        import tkinter as tk
        from tkinter import ttk
        from riju.app import App
        from riju import gst
        from test_gst_2b_notes import Gst2bNoteTests
        def walk(w):
            for child in w.winfo_children():
                yield child
                yield from walk(child)
        def press(win,text):next(w for w in walk(win) if isinstance(w,ttk.Button) and w.cget('text')==text).invoke()
        fixture=Gst2bNoteTests();fixture.setUp()
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);source=root/'2b-notes.csv';source.write_bytes(fixture.raw([fixture.row()]))
            vault=Vault(root/'demo.json',demo=True);app=App(vault)
            try:
                app.commit(lambda s:(s.clear(),s.update(copy.deepcopy(fixture.s))));app.show('GST accounting');app.update()
                count=len(app.state['vouchers'])
                with patch('riju.gst_ui.filedialog.askopenfilename',return_value=str(source)),patch('riju.gst_ui.simpledialog.askstring',side_effect=['29ABCDE1234F1Z5','2026-05']),patch('riju.gst_ui.messagebox.askyesno',return_value=True):app.import_gst_2b()
                app.update();app.gst_import_table.selection_set(app.gst_import_table.get_children()[0]);app.review_gst_2b();app.update()
                win=next(w for w in app.winfo_children() if isinstance(w,tk.Toplevel));table=next(w for w in walk(win) if isinstance(w,ttk.Treeview) and 'Match status' in w['columns'])
                self.assertIn('Credit note',table.item(table.get_children()[0],'values'));self.assertIn('P-1',table.item(table.get_children()[0],'values'))
                table.selection_set(table.get_children()[0])
                with patch('riju.gst_ui.messagebox.askyesno',return_value=True):press(win,'Approve exact match')
                self.assertEqual(gst.review_2b(app.state,app.state['gst_2b_imports'][0]['id'])['matched_signed_totals']['cgst'],-450)
                with patch('riju.gst_ui.filedialog.asksaveasfilename',return_value=str(root/'review.csv')):press(win,'Export review CSV / JSON')
                self.assertIn('-50.00',(root/'review.csv').read_text(encoding='utf-8-sig'));self.assertEqual(len(app.state['vouchers']),count);win.destroy()
            finally:app.destroy()
            loaded=vault.load();self.assertEqual(loaded['gst_2b_matches'][0]['document_type'],'Credit note')

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class RcmDesktopTests(unittest.TestCase):
    def test_review_all_rcm_stages_and_reopen(self):
        import tkinter as tk
        from tkinter import ttk
        from riju.app import App
        from riju import tax,gst
        def walk(w):
            for child in w.winfo_children():
                yield child
                yield from walk(child)
        def fill(win,label,value):
            item=next(w for w in walk(win) if isinstance(w,ttk.Label) and w.cget('text')==label);entry=next(w for w in item.master.winfo_children() if isinstance(w,ttk.Entry))
            if isinstance(entry,ttk.Combobox):entry.set(value)
            else:entry.delete(0,'end');entry.insert(0,value)
        def press(win,text):next(w for w in walk(win) if isinstance(w,ttk.Button) and w.cget('text')==text).invoke()
        with tempfile.TemporaryDirectory() as d:
            vault=Vault(Path(d)/'demo.json',demo=True);app=App(vault)
            try:
                app.commit(lambda s:tax.save_rule(s,dict(family='GST',key='GUI-RCM',effective='2026-04-01',label='Invented',rate='18',treatment='Taxable',hsn_sac='999311',applicability='Synthetic tests',source='Invented source',reviewed=True)))
                app.show('GST reverse charge');app.update()
                stage_fields=[('Purchase',[('Accounting/action date','2026-04-02'),('Invoice / challan / utilisation reference','GUI-RCM-1'),('Narration / supporting review','Test purchase'),('Recipient GSTIN','29ABCDE1234F1Z5'),('Verified RCM applicability basis','Invented test only'),('Supplier name (required if unregistered)','Test supplier'),('Supplier state code','29'),('Reviewed place-of-supply state','29'),('Purchase base INR (excludes RCM)','100')]),('Deposit',[('Accounting/action date','2026-04-03'),('Invoice / challan / utilisation reference','GUI-CIN'),('Narration / supporting review','Test deposit'),('Recipient GSTIN','29ABCDE1234F1Z5'),('CGST cash deposit INR','9'),('SGST cash deposit INR','9')]),('Discharge',[('Accounting/action date','2026-04-04'),('Invoice / challan / utilisation reference','GUI-UTIL'),('Narration / supporting review','Test actual cash utilisation')]),('Claim',[('Accounting/action date','2026-04-05'),('Invoice / challan / utilisation reference','GUI-ITC'),('Narration / supporting review','Reviewed test credit')])]
                for stage,fields in stage_fields:
                    app.new_rcm(stage);app.update();win=next(w for w in app.winfo_children() if isinstance(w,tk.Toplevel))
                    for label,value in fields:fill(win,label,value)
                    if stage in ('Purchase','Claim'):next(w for w in walk(win) if isinstance(w,ttk.Checkbutton)).invoke()
                    with patch('riju.app.messagebox.showerror') as errors:
                        press(win,'Prepare RCM preview');errors.assert_not_called()
                        if stage=='Purchase':
                            app.commit(lambda s:s['settings'].update(entity='Changed synthetic entity'))
                            press(win,'Approve & post RCM');self.assertTrue(errors.called);self.assertFalse(app.state['gst_rcm_records'])
                            errors.reset_mock();press(win,'Prepare RCM preview');errors.assert_not_called()
                    with patch('riju.rcm_ui.messagebox.askyesno',return_value=True):press(win,'Approve & post RCM')
                    app.update()
                self.assertEqual(len(app.state['gst_rcm_records']),4);self.assertEqual(len(app.rcm_table.get_children()),4);self.assertFalse(any(app.rcm_report['unclassified'].values()))
                self.assertEqual(gst.summary(app.state,'2026-04-01','2026-04-30')['totals']['input_cgst'],900)
                with patch('riju.rcm_ui.filedialog.asksaveasfilename',return_value=str(Path(d)/'rcm.json')):app.export_rcm()
                self.assertTrue((Path(d)/'rcm.json').exists())
            finally:app.destroy()
            loaded=vault.load();self.assertEqual(len(loaded['gst_rcm_records']),4)

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class FinancialReportExportDesktopTests(unittest.TestCase):
    def test_all_report_exports_match_the_displayed_snapshot(self):
        from riju.app import App
        from riju import core,financial_reports as f
        from openpyxl import load_workbook
        from pypdf import PdfReader
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);vault=Vault(root/'demo.json',demo=True);app=App(vault)
            try:
                app.commit(lambda s:core.post(s,'2026-04-02','Invented report test',[dict(account='1010',branch='JP',debit=12345,credit=0),dict(account='4000',branch='JP',debit=0,credit=12345)]))
                app.show('Financial reports');app.update()
                for index,mode in enumerate(f.MODES):
                    app.report_mode.set(mode);app.report_from.set('2026-04-01');app.report_through.set('2027-03-31')
                    # Export must refresh edited selectors even when Refresh was not clicked.
                    for kind in ('xlsx','pdf'):
                        path=root/f'{index}.{kind}'
                        with patch('riju.app.filedialog.asksaveasfilename',return_value=str(path)),patch('riju.app.messagebox.showinfo'):app.export_financial_report(kind)
                        app.update();self.assertEqual(app.financial_snapshot['mode'],mode)
                        self.assertEqual(len(app.report_table.get_children()),len(app.financial_snapshot['rows']))
                        if kind=='xlsx':
                            w=load_workbook(path);self.assertEqual(w.active['A2'].value,mode);w.close()
                        else:self.assertIn(mode,PdfReader(path).pages[0].extract_text())
                self.assertEqual(len(app.state['vouchers']),1)
            finally:app.destroy()
            self.assertEqual(len(vault.load()['vouchers']),1)

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class FinancialPackDesktopTests(unittest.TestCase):
    def test_customer_pack_buttons_export_all_reports_without_posting(self):
        from riju.app import App
        from riju import financial_reports as f
        from openpyxl import load_workbook
        from pypdf import PdfReader
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);vault=Vault(root/'demo.json',demo=True);app=App(vault)
            try:
                app.show('Financial reports');app.update()
                for kind in ('xlsx','pdf'):
                    path=root/('pack.'+kind)
                    with patch('riju.app.filedialog.asksaveasfilename',return_value=str(path)),patch('riju.app.messagebox.showinfo'):app.export_financial_pack(kind)
                    if kind=='xlsx':
                        wb=load_workbook(path);self.assertEqual(wb.sheetnames,['Contents',*f.MODES]);wb.close()
                    else:self.assertIn('Financial working report pack',PdfReader(path).pages[0].extract_text())
                self.assertEqual(app.state['vouchers'],[])
            finally:app.destroy()

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class TdsDesktopTests(unittest.TestCase):
    def test_complete_tds_expense_challan_allocation_cin_export_and_reopen(self):
        import tkinter as tk
        from tkinter import ttk
        from riju.app import App
        from riju import tax
        def walk(w):
            for child in w.winfo_children():
                yield child
                yield from walk(child)
        def fill(win,label,value):
            item=next(w for w in walk(win) if isinstance(w,ttk.Label) and w.cget('text')==label);entry=next(w for w in item.master.winfo_children() if isinstance(w,ttk.Entry))
            if isinstance(entry,ttk.Combobox):entry.set(value)
            else:entry.delete(0,'end');entry.insert(0,value)
        def press(win,text):next(w for w in walk(win) if isinstance(w,ttk.Button) and w.cget('text')==text).invoke()
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);vault=Vault(root/'demo.json',demo=True);app=App(vault)
            try:
                app.commit(lambda s:tax.save_rule(s,dict(family='TDS',key='GUI-J',effective='2024-04-01',label='Synthetic',section='194J',rate='10',treatment='Withholding',party_category='Company',annual_threshold='1000',threshold_basis='Financial-year aggregate',applicability='Tests only',source='Invented',reviewed=True)))
                app.commit(lambda s:tax.save_party(s,dict(name='Synthetic payee',effective='2024-04-01',category='Company',pan='ABCDE1234F')))
                app.show('TDS accounting');app.update()
                for index in (1,2):
                    app.new_tds('Expense');app.update();win=next(w for w in app.winfo_children() if isinstance(w,tk.Toplevel))
                    for label,value in [('Posting date','2026-04-0'+str(index+1)),('Credit/payment tax-point date','2026-04-0'+str(index+1)),('Unique invoice/payment reference','GUI-INV-'+str(index)),('Narration','Synthetic fees'),('Gross expense/payable INR','600'),('Deduction base INR','600')]:fill(win,label,value)
                    next(w for w in walk(win) if isinstance(w,ttk.Checkbutton)).invoke()
                    with patch('riju.app.messagebox.showerror') as error:press(win,'Prepare TDS preview');error.assert_not_called()
                    with patch('riju.tds_ui.messagebox.askyesno',return_value=True):press(win,'Approve & record TDS')
                    app.update()
                app.new_tds('Challan');app.update();win=next(w for w in app.winfo_children() if isinstance(w,tk.Toplevel))
                for label,value in [('Posting date','2026-04-05'),('Unique invoice/payment reference','GUI-UTR'),('Narration','Synthetic tax payment')]:fill(win,label,value)
                source=next(w for w in walk(win) if isinstance(w,tk.Text) and int(w.cget('height'))==4);source.delete('1.0','end');source.insert('1.0','section,branch,year,tax\n194J,JP,2026,120\n')
                with patch('riju.app.messagebox.showerror') as error:press(win,'Prepare TDS preview');error.assert_not_called()
                with patch('riju.tds_ui.messagebox.askyesno',return_value=True):press(win,'Approve & record TDS')
                app.update();count=len(app.state['vouchers']);app.allocate_tds();app.update();win=next(w for w in app.winfo_children() if isinstance(w,tk.Toplevel))
                fill(win,'Tax to allocate INR','120');fill(win,'Allocation date','2026-04-06')
                with patch('riju.tds_ui.messagebox.askyesno',return_value=True):press(win,'Approve allocation')
                app.update();self.assertEqual(len(app.state['vouchers']),count);self.assertEqual(app.tds_report['outstanding']['rows'],[])
                app.edit_tds_challan();app.update();win=next(w for w in app.winfo_children() if isinstance(w,tk.Toplevel))
                for label,value in [('BSR code (optional)','1234567'),('Challan serial (optional)','00001'),('Actual deposit date (optional)','2026-04-05')]:fill(win,label,value)
                next(w for w in walk(win) if isinstance(w,ttk.Checkbutton)).invoke();press(win,'Save detail revision');app.update()
                self.assertEqual(app.tds_report['challans']['rows'][0][8],'Reviewed')
                for kind in ('xlsx','pdf'):
                    with patch('riju.tds_ui.filedialog.asksaveasfilename',return_value=str(root/('tds.'+kind))):app.export_tds(kind)
                    self.assertTrue((root/('tds.'+kind)).exists())
            finally:app.destroy()
            reopened=vault.load();self.assertEqual(len(reopened['tds_records']),3);self.assertEqual(len(reopened['tds_allocations']),1)

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class CashFlowDesktopTests(unittest.TestCase):
    def test_policy_review_mixed_voucher_allocation_export_and_reopen(self):
        import tkinter as tk
        from tkinter import ttk
        from riju.app import App
        from riju import core,cash_flow as cf
        def walk(w):
            for child in w.winfo_children():
                yield child
                yield from walk(child)
        def fill(win,label,value):
            item=next(w for w in walk(win) if isinstance(w,ttk.Label) and w.cget('text')==label);entry=next(w for w in item.master.winfo_children() if isinstance(w,ttk.Entry))
            if isinstance(entry,ttk.Combobox):entry.set(value)
            else:entry.delete(0,'end');entry.insert(0,value)
        def press(win,text):next(w for w in walk(win) if isinstance(w,ttk.Button) and w.cget('text')==text).invoke()
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);vault=Vault(root/'demo.json',demo=True);app=App(vault)
            try:
                app.commit(lambda s:core.post(s,'2026-04-02','Synthetic loan and noncash equipment',[dict(account='1010',branch='JP',debit=10000,credit=0),dict(account='1030',branch='JP',debit=10000,credit=0),dict(account='2400',branch='JP',debit=0,credit=20000)]))
                app.show('Cash flow setup');app.update();self.assertFalse(app.cash_flow_result['ready']);app.new_cash_flow_policy();app.update()
                win=next(w for w in app.winfo_children() if isinstance(w,tk.Toplevel));fill(win,'Effective report-policy date','2026-04-01');fill(win,'Verified definition/classification basis','Synthetic cash policy reviewed')
                next(w for w in walk(win) if isinstance(w,ttk.Checkbutton)).invoke();press(win,'Save policy version');app.update()
                app.review_cash_flow_voucher();app.update();win=next(w for w in app.winfo_children() if isinstance(w,tk.Toplevel));fill(win,'Supporting classification reason','Verified cash portion of loan, remaining equipment funded noncash')
                text=next(w for w in walk(win) if isinstance(w,tk.Text));text.delete('1.0','end');text.insert('1.0','category,amount,label\nFinancing,100,Cash portion of loan\n');press(win,'Save reviewed cash allocation');app.update()
                self.assertTrue(app.cash_flow_result['ready']);count=len(app.state['vouchers'])
                for kind in ('xlsx','pdf'):
                    with patch('riju.cash_flow_ui.filedialog.asksaveasfilename',return_value=str(root/('cash.'+kind))):app.export_cash_flow(kind)
                    self.assertTrue((root/('cash.'+kind)).exists())
                self.assertEqual(len(app.state['vouchers']),count)
                app.show('Financial reports');app.report_mode.set(cf.INDIRECT);app.refresh_financial_report();self.assertEqual(app.financial_snapshot['cash_flow']['indirect_difference'],0)
            finally:app.destroy()
            loaded=vault.load();self.assertEqual(len(loaded['cash_flow_policies']),1);self.assertEqual(len(loaded['cash_flow_allocations']),1)

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class BankTemplateDesktopTests(unittest.TestCase):
    def test_save_reuse_preview_approve_disable_and_reload_layout(self):
        import tkinter as tk
        from tkinter import ttk
        from riju.app import App
        def walk(w):
            for child in w.winfo_children():
                yield child
                yield from walk(child)
        def field(win,label):
            item=next(w for w in walk(win) if isinstance(w,ttk.Label) and w.cget('text')==label)
            return next(w for w in item.master.winfo_children() if isinstance(w,ttk.Entry))
        def fill(win,label,value):
            w=field(win,label)
            if isinstance(w,ttk.Combobox):w.set(value)
            else:w.delete(0,'end');w.insert(0,value)
        def press(win,label):next(w for w in walk(win) if isinstance(w,ttk.Button) and w.cget('text')==label).invoke()
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);vault=Vault(root/'demo.json',demo=True);app=App(vault)
            try:
                app.show('Bank reconciliation');app.update()
                april=root/'april.csv';april.write_text('When,Movement,Ref\n02/04/2026,100,A1\n',encoding='utf-8')
                with patch('riju.bank_ui.filedialog.askopenfilename',return_value=str(april)):app.import_bank_statement()
                win=next(w for w in app.winfo_children() if isinstance(w,tk.Toplevel));win.geometry('900x450');app.update()
                canvas=next(w for w in walk(win) if isinstance(w,tk.Canvas));self.assertGreater(canvas.bbox('all')[3],canvas.winfo_height());canvas.yview_moveto(1)
                for label,value in [('Date format','DD/MM/YYYY'),('Date column','When'),('Amount column','Movement'),('Reference column','Ref'),('Layout name','Synthetic bank layout'),('Statement from','2026-04-01'),('Through','2026-04-30'),('Bank opening INR (optional)','0'),('Bank closing INR (optional)','100')]:fill(win,label,value)
                press(win,'Validate & save layout');app.update();self.assertEqual(len(app.state['bank_mapping_templates']),1)
                press(win,'Validate & preview rows');app.update();preview=next(w for w in win.winfo_children() if isinstance(w,tk.Toplevel));press(preview,'Approve statement import');app.update()
                self.assertEqual(len(app.state['bank_statements']),1)
                may=root/'may.csv';may.write_text('When,Movement,Ref\n02/05/2026,50,M1\n',encoding='utf-8')
                with patch('riju.bank_ui.filedialog.askopenfilename',return_value=str(may)):app.import_bank_statement()
                win=next(w for w in app.winfo_children() if isinstance(w,tk.Toplevel));choice=field(win,'Saved layout');choice.set(choice.cget('values')[1]);press(win,'Apply saved layout')
                self.assertEqual(field(win,'Date column').get(),'When');self.assertEqual(field(win,'Bank opening INR (optional)').get(),'')
                for label,value in [('Statement from','2026-05-01'),('Through','2026-05-31'),('Bank opening INR (optional)','100'),('Bank closing INR (optional)','150')]:fill(win,label,value)
                press(win,'Validate & preview rows');app.update();preview=next(w for w in win.winfo_children() if isinstance(w,tk.Toplevel));press(preview,'Approve statement import');app.update()
                self.assertEqual(len(app.state['bank_statements']),2);self.assertEqual(len(app.state['vouchers']),0)
                with patch('riju.bank_ui.filedialog.askopenfilename',return_value=str(may)):app.import_bank_statement()
                win=next(w for w in app.winfo_children() if isinstance(w,tk.Toplevel));choice=field(win,'Saved layout');choice.set(choice.cget('values')[1])
                with patch('riju.bank_ui.simpledialog.askstring',return_value='Obsolete layout'):press(win,'Disable layout')
                self.assertFalse(app.state['bank_mapping_templates'][0]['active']);win.destroy()
            finally:app.destroy()
            loaded=vault.load();self.assertEqual(len(loaded['bank_statements']),2);self.assertFalse(loaded['bank_mapping_templates'][0]['active']);self.assertTrue(loaded['bank_statements'][0]['mapping_template']['active'])

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class FixedAssetsDesktopTests(unittest.TestCase):
    def walk(self,w):
        for child in w.winfo_children():
            yield child
            yield from self.walk(child)
    def fill(self,win,label,value):
        from tkinter import ttk
        item=next(w for w in self.walk(win) if isinstance(w,ttk.Label) and w.cget('text')==label);entry=next(w for w in item.master.winfo_children() if isinstance(w,ttk.Entry))
        if isinstance(entry,ttk.Combobox):entry.set(value)
        else:entry.delete(0,'end');entry.insert(0,value)
    def press(self,win,label):
        from tkinter import ttk
        next(w for w in self.walk(win) if isinstance(w,ttk.Button) and w.cget('text')==label).invoke()
    def top(self,parent):
        import tkinter as tk
        return [w for w in parent.winfo_children() if isinstance(w,tk.Toplevel)][-1]
    def test_register_depreciate_transfer_dispose_export_and_reload(self):
        from tkinter import ttk
        from riju.app import App
        from riju import fixed_assets as fa
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);vault=Vault(root/'demo.json',demo=True);app=App(vault)
            try:
                app.year.set('2026');app.show('Assets & depreciation');app.asset_to.set('2026-04-30');app.show('Assets & depreciation');app.register_book_asset();app.update();form=self.top(app)
                for label,value in [('Unique asset reference','GUI-1'),('Asset / component name','Synthetic diagnostic equipment'),('Cost source mode','Post purchase'),('Acquired date','2026-04-01'),('Put-to-use date','2026-04-01'),('Book baseline date','2026-04-01'),('Original capitalized cost INR','1200'),('Useful / remaining life months','12'),('Reviewed classification/source basis','Synthetic reviewed medical category, life and cost source')]:self.fill(form,label,value)
                next(w for w in self.walk(form) if isinstance(w,ttk.Checkbutton)).invoke();self.press(form,'Validate & preview asset');app.update();preview=self.top(app)
                # Editing the form after preview must not alter the approved snapshot.
                self.fill(form,'Original capitalized cost INR','9999');self.press(preview,'Approve asset action');app.update();self.assertEqual(app.state['book_assets'][0]['cost'],120000);self.assertEqual(len(app.state['vouchers']),1)
                app.book_depreciation();form=self.top(app);self.press(form,'Preview depreciation journal');app.update();self.press(self.top(app),'Approve asset action');app.update();self.assertTrue(app.asset_report['ready'])
                for kind in ('xlsx','pdf'):
                    with patch('riju.fixed_assets_ui.filedialog.asksaveasfilename',return_value=str(root/('assets.'+kind))):app.export_assets(kind)
                    self.assertTrue((root/('assets.'+kind)).exists())
                app.asset_table.selection_set(app.asset_table.get_children()[0]);app.asset_transfer();form=self.top(app)
                for label,value in [('Transfer date','2026-05-01'),('Destination unit code','KM'),('Transfer source basis','Reviewed physical move')]:self.fill(form,label,value)
                next(w for w in self.walk(form) if isinstance(w,ttk.Checkbutton)).invoke();self.press(form,'Preview transfer journal');app.update();self.press(self.top(app),'Approve asset action');app.update()
                app.asset_to.set('2026-05-31');app.show('Assets & depreciation');app.asset_table.selection_set(app.asset_table.get_children()[0]);app.asset_disposal();form=self.top(app)
                for label,value in [('Disposal date','2026-05-31'),('Proceeds excluding taxes INR','1050'),('Disposal source / tax-review basis','Reviewed sale with separately considered tax')]:self.fill(form,label,value)
                next(w for w in self.walk(form) if isinstance(w,ttk.Checkbutton)).invoke();self.press(form,'Preview disposal journal');app.update();self.press(self.top(app),'Approve asset action');app.update();self.assertTrue(app.asset_report['ready']);self.assertEqual(app.asset_report['register'][0][9],'Disposed')
            finally:app.destroy()
            loaded=vault.load();self.assertEqual(len(loaded['fixed_asset_events']),3);self.assertTrue(fa.disposed(loaded,loaded['book_assets'][0],'2026-05-31'))
    def test_changed_books_reject_stale_asset_preview_without_posting(self):
        from riju.app import App
        from riju import foundation,fixed_assets as fa
        with tempfile.TemporaryDirectory() as d:
            app=App(Vault(Path(d)/'demo.json',demo=True))
            try:
                app.show('Assets & depreciation');app.commit(lambda s:fa.register(s,dict(reference='STALE',name='Synthetic',branch='JP',category='Manual / specific Schedule II category',mode='Post purchase',acquired='2026-04-01',put_to_use='2026-04-01',baseline='2026-04-01',cost='1200',months=12,residual='0',basis='Reviewed test',reviewed=True)))
                app._asset_preview('Synthetic stale preview',lambda s:fa.post_depreciation(s,'2026-04-30'));preview=self.top(app);count=len(app.state['vouchers']);app.commit(lambda s:foundation.save_entity(s,'Changed entity',{}))
                with patch('riju.app.messagebox.showerror') as error:self.press(preview,'Approve asset action');self.assertTrue(error.called)
                self.assertEqual(len(app.state['vouchers']),count);self.assertFalse(app.state['fixed_asset_events'])
            finally:app.destroy()

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class TaxDepreciationDesktopTests(unittest.TestCase):
    def walk(self,w):
        for c in w.winfo_children():
            yield c
            yield from self.walk(c)
    def form(self,app):
        import tkinter as tk
        return [w for w in app.winfo_children() if isinstance(w,tk.Toplevel)][-1]
    def fill(self,win,label,value):
        from tkinter import ttk
        item=next(w for w in self.walk(win) if isinstance(w,ttk.Label) and w.cget('text')==label);field=next(w for w in item.master.winfo_children() if isinstance(w,ttk.Entry))
        if isinstance(field,ttk.Combobox):field.set(value)
        else:field.delete(0,'end');field.insert(0,value)
    def press(self,win,label):
        from tkinter import ttk
        next(w for w in self.walk(win) if isinstance(w,ttk.Button) and w.cget('text')==label).invoke()
    def reviewed(self,win):
        from tkinter import ttk
        [w for w in self.walk(win) if isinstance(w,ttk.Checkbutton)][-1].invoke()
    def approve(self,app,form,label):
        self.press(form,label);app.update();self.press(self.form(app),'Approve tax working action');app.update()
    def select(self,app):app.tax_block_table.selection_set(app.tax_block_table.get_children()[0])
    def test_entity_tax_block_addition_finalize_exports_reopen_and_reload(self):
        import tkinter as tk
        from riju.app import App
        from riju import tax_depreciation as td
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);vault=Vault(root/'demo.json',demo=True);app=App(vault)
            try:
                app.year.set('2024');app.show('Tax depreciation');app.new_tax_block();form=self.form(app)
                for label,value in [('Block key','P15'),('Block name','Synthetic plant tax block'),('First tax financial year','2024'),('Opening WDV / control-count source basis','Reviewed opening WDV and asset counts')]:self.fill(form,label,value)
                text=next(w for w in self.walk(form) if isinstance(w,tk.Text));text.delete('1.0','end');text.insert('1.0','unit,wdv,count\nJP,1000,2\nKM,0,0\nVK,0,0\n');self.reviewed(form);self.approve(app,form,'Preview opening tax pool');self.select(app);app.tax_depreciation_rule();form=self.form(app)
                for label,value in [('Effective financial year','2024'),('WDV rate %','15'),('Reviewed year / regime / rate source','Reviewed synthetic Appendix I rate')]:self.fill(form,label,value)
                self.reviewed(form);self.approve(app,form,'Preview tax rate policy');self.select(app);app.tax_addition();form=self.form(app)
                for label,value in [('Unique tax addition reference','GUI-TAX-1'),('Acquired date','2024-12-01'),('Actual first use (blank if unused)','2024-12-01'),('Actual tax cost INR','1000'),('Reviewed cost / tax-source basis','Reviewed capital cost and actual commissioning')]:self.fill(form,label,value)
                self.reviewed(form);self.approve(app,form,'Preview tax addition');self.assertEqual(app.tax_dep_report['normal'],22500);self.assertTrue(app.tax_dep_report['ready']);self.assertFalse(app.state['vouchers'])
                for kind in ('xlsx','pdf'):
                    with patch('riju.tax_depreciation_ui.filedialog.asksaveasfilename',return_value=str(root/('tax.'+kind))):app.export_tax_depreciation(kind)
                    self.assertTrue((root/('tax.'+kind)).exists())
                app.finalize_tax_year();form=self.form(app);self.fill(form,'Entity year-end source / tax review basis','Reviewed complete annual source and regime working');self.reviewed(form);self.approve(app,form,'Preview final tax working');self.assertTrue(app.tax_dep_report.get('finalized_review_id'))
                with patch('riju.tax_depreciation_ui.simpledialog.askstring',return_value='Revisit original client source'):app.reopen_tax_year()
                self.press(self.form(app),'Approve tax working action');app.update();self.assertFalse(app.tax_dep_report.get('finalized_review_id'))
            finally:app.destroy()
            loaded=vault.load();self.assertEqual(td.data(loaded)['reviews'][0]['status'],'REOPENED');self.assertEqual(td.calculate(loaded,2024)['normal'],22500);self.assertFalse(loaded['vouchers'])
    def test_changed_books_reject_stale_tax_working(self):
        from riju.app import App
        from riju import tax_depreciation as td,foundation
        with tempfile.TemporaryDirectory() as d:
            app=App(Vault(Path(d)/'demo.json',demo=True))
            try:
                app.show('Tax depreciation');app._tax_work_preview('Synthetic preview',lambda s:td.save_block(s,'STALE','Synthetic',2024,{},'Reviewed',True));preview=self.form(app);app.commit(lambda s:foundation.save_entity(s,'Changed entity',{}))
                with patch('riju.app.messagebox.showerror') as error:self.press(preview,'Approve tax working action');self.assertTrue(error.called)
                self.assertFalse(td.data(app.state)['blocks']);self.assertFalse(app.state['vouchers'])
            finally:app.destroy()

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class CwipDesktopTests(unittest.TestCase):
    walk=TaxDepreciationDesktopTests.walk
    form=TaxDepreciationDesktopTests.form
    fill=TaxDepreciationDesktopTests.fill
    press=TaxDepreciationDesktopTests.press
    reviewed=TaxDepreciationDesktopTests.reviewed
    def approve(self,app,form):
        self.reviewed(form);self.press(form,'Validate & preview CWIP');app.update();self.press(self.form(app),'Approve CWIP action');app.update()
    def select_project(self,app):app.cwip_project_table.selection_set(app.cwip_project_table.get_children()[0])
    def select_cost(self,app):app.cwip_cost_table.selection_set(app.cwip_cost_table.get_children()[0])
    def test_project_plan_cost_retention_tds_commission_exports_and_reload(self):
        import tkinter as tk
        from riju.app import App
        from riju import cwip, tax, fixed_assets as fa
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);vault=Vault(root/'demo.json',demo=True);app=App(vault)
            try:
                app.year.set('2026');app.cwip_report_end='2026-05-31';app.show('CWIP & commissioning');app.new_cwip_project();f=self.form(app)
                for label,value in [('Unique project key','GUI-VK'),('Project name','Synthetic desktop construction'),('Project/source basis','Synthetic sources reviewed')]:self.fill(f,label,value)
                self.approve(app,f);self.select_project(app);app.cwip_plan();f=self.form(app)
                for label,value in [('Effective date','2026-04-01'),('Civil budget INR (blank unknown)','2000'),('Plan revision basis','Approved synthetic plan')]:self.fill(f,label,value)
                text=next(w for w in self.walk(f) if isinstance(w,tk.Text));text.delete('1.0','end');text.insert('1.0','key,name,due,completed\nREADY,Certificate review,2026-05-01,2026-05-01\n');self.approve(app,f)
                self.select_project(app);app.cwip_cost();f=self.form(app)
                for label,value in [('Cost / source date','2026-04-01'),('Supplier/source reference','GUI-BILL'),('Supplier/payee reference','Synthetic builder'),('Gross cost INR','1000'),('Eligible capital INR','900'),('Retention to hold INR','100'),('Capitalization / expense review basis','Reviewed 900 capital and 100 operating cost')]:self.fill(f,label,value)
                self.approve(app,f);self.assertTrue(app.cwip_report['ready']);self.select_cost(app);app.cwip_retention();f=self.form(app)
                for label,value in [('Release / payment date','2026-04-20'),('Retention amount INR','40'),('Retention release / payment basis','Approved partial release')]:self.fill(f,label,value)
                self.approve(app,f)
                def setup(s):
                    p=tax.save_party(s,dict(name='Synthetic builder',effective='2026-04-01',category='Company',pan='ABCDE1234F'))
                    tax.save_rule(s,dict(family='TDS',key='GUI-CWIP',label='Synthetic contractor rule',effective='2026-04-01',section='194C',rate='10',treatment='Withholding',party_category='Company',annual_threshold='0',threshold_basis='Financial-year aggregate',applicability='Synthetic only',source='Synthetic test',reviewed=True));return p
                app.commit(setup);app.show('CWIP & commissioning');self.select_cost(app);app.cwip_tds();f=self.form(app)
                self.fill(f,'Deduction posting date','2026-04-02');self.fill(f,'Verified applicability / exception basis','Reviewed synthetic resident contractor applicability');self.approve(app,f)
                self.select_cost(app);app.cwip_expense();f=self.form(app)
                for label,value in [('Reclassification date','2026-04-25'),('Ineligible cost INR','100'),('Reclassification source basis','Reviewed start-up cost, not capital')]:self.fill(f,label,value)
                self.approve(app,f);self.select_project(app);app.cwip_commission();f=self.form(app)
                for label,value in [('Commissioning transfer date','2026-05-01'),('Actual put-to-use date','2026-05-16'),('Unique book asset reference','GUI-CWIP-PPE'),('Asset / component name','Synthetic commissioned component'),('Reviewed useful life months','12'),('Ready-for-use / capitalization certificate basis','Synthetic reviewed readiness certificate')]:self.fill(f,label,value)
                self.approve(app,f);self.assertTrue(app.cwip_report['ready']);self.assertEqual(app.state['book_assets'][0]['cost'],80000)
                for kind in ('xlsx','pdf'):
                    with patch('riju.cwip_ui.filedialog.asksaveasfilename',return_value=str(root/('cwip.'+kind))):app.export_cwip(kind)
                    self.assertTrue((root/('cwip.'+kind)).exists())
            finally:app.destroy()
            loaded=vault.load();self.assertTrue(cwip.report(loaded,'2026-04-01','2026-05-31')['ready']);self.assertEqual(loaded['book_assets'][0]['put_to_use'],'2026-05-16');self.assertEqual(fa.target(loaded['book_assets'][0],'2026-05-15'),0)
    def test_cwip_changed_books_rejects_frozen_preview(self):
        from riju.app import App
        from riju import cwip, foundation
        with tempfile.TemporaryDirectory() as d:
            app=App(Vault(Path(d)/'demo.json',demo=True))
            try:
                app.show('CWIP & commissioning');app._cwip_preview('Synthetic stale preview',lambda s:cwip.save_project(s,'STALE','Synthetic','VK','Test',True));preview=self.form(app);app.commit(lambda s:foundation.save_entity(s,'Changed entity',{}))
                with patch('riju.app.messagebox.showerror') as err:self.press(preview,'Approve CWIP action');self.assertTrue(err.called)
                self.assertFalse(cwip.data(app.state)['projects']);self.assertFalse(app.state['vouchers'])
            finally:app.destroy()

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class ServiceRevenueDesktopTests(unittest.TestCase):
    walk=TaxDepreciationDesktopTests.walk
    form=TaxDepreciationDesktopTests.form
    fill=TaxDepreciationDesktopTests.fill
    press=TaxDepreciationDesktopTests.press
    reviewed=TaxDepreciationDesktopTests.reviewed
    def test_advance_earning_export_reload_and_stale_preview(self):
        from riju.app import App
        from riju import tax,service_revenue as sr,foundation
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);vault=Vault(root/'service.json',demo=True);app=App(vault)
            try:
                app.commit(lambda s:tax.save_rule(s,dict(family='GST',key='GUI-SERVICE',effective='2026-04-01',label='Synthetic',rate='0',treatment='Exempt',source='Test',applicability='Test',reviewed=True)))
                app.service_report_end='2026-04-30';app.show('Service revenue');app.service_action('Advance');f=self.form(app)
                for k,v in [('Action date','2026-04-01'),('Unique action reference','GUI-ADV'),('Service / allocated gross INR','1000'),('Actual payments JSON','{"1010":"1000"}'),('Service / tax / source review basis','Reviewed synthetic advance')]:self.fill(f,k,v)
                self.reviewed(f);self.press(f,'Validate service preview');app.update();self.press(self.form(app),'Approve service action');app.update()
                app.service_table.selection_set(app.service_table.get_children()[0]);app.service_action('Earn advance');f=self.form(app)
                for k,v in [('Action date','2026-04-02'),('Unique action reference','GUI-EARN'),('Service / allocated gross INR','600'),('Service / tax / source review basis','Reviewed synthetic earned service')]:self.fill(f,k,v)
                self.reviewed(f);self.press(f,'Validate service preview');app.update();self.press(self.form(app),'Approve service action');app.update();self.assertEqual(__import__('riju.core',fromlist=['balances']).balances(app.state,'2026-04-30')['2300'],-40000)
                for kind in ('xlsx','pdf'):
                    with patch('riju.service_revenue_ui.filedialog.asksaveasfilename',return_value=str(root/('service.'+kind))):app.export_service_revenue(kind)
                    self.assertTrue((root/('service.'+kind)).exists())
                app._service_preview('Stale',lambda s:sr.post(s,'Bill',dict(date='2026-04-03',reference='STALE',branch='JP',account='4010',rule_key='GUI-SERVICE',amount='1',basis='Review',reviewed=True)));f=self.form(app);app.commit(lambda s:foundation.save_entity(s,'Changed',{}))
                with patch('riju.app.messagebox.showerror') as error:self.press(f,'Approve service action');self.assertTrue(error.called)
                self.assertEqual(len(app.state['service_revenue']),2)
            finally:app.destroy()
            self.assertEqual(len(vault.load()['service_revenue']),2)

@unittest.skipUnless(os.environ.get('DISPLAY') or os.environ.get('RIJU_GUI_TEST'),'A graphical display is required')
class BulkImportDesktopTests(unittest.TestCase):
    def test_mapped_row_review_frozen_post_export_and_reload(self):
        from riju.app import App
        from riju import bulk_import as bulk
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);vault=Vault(root/'bulk.json',demo=True);app=App(vault)
            try:
                raw=b'date,reference,branch,narration,debit_account,credit_account,amount,journal_id\n2026-04-01,GUI-BULK,JP,Synthetic,4590,2000,100,\n'
                app.commit(lambda s:bulk.stage(s,raw,'gui.csv','Expense'));app.show('Bulk import review');app.bulk_batches.selection_set(app.bulk_batches.get_children()[0]);app.open_bulk_rows();app.bulk_rows.selection_set(app.bulk_rows.get_children()[0])
                with patch('riju.bulk_import_ui.simpledialog.askstring',return_value='Reviewed synthetic source'):app.review_bulk_row(True)
                app.post_bulk();app.update();helper=TaxDepreciationDesktopTests();helper.press(helper.form(app),'Approve service action');app.update();self.assertEqual(len(app.state['vouchers']),1)
                app.bulk_batches.selection_set(app.bulk_batches.get_children()[0])
                for kind in ('xlsx','pdf'):
                    with patch('riju.bulk_import_ui.filedialog.asksaveasfilename',return_value=str(root/('bulk.'+kind))):app.export_bulk(kind)
                    self.assertTrue((root/('bulk.'+kind)).exists())
            finally:app.destroy()
            self.assertEqual(vault.load()['queue'][0]['status'],'PROCESSED')
