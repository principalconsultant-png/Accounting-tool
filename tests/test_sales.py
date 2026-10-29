import copy
from datetime import datetime
import io
import json
from pathlib import Path
import tempfile
import unittest
from openpyxl import Workbook
from riju import core
from riju.revenue import MODES,ensure_accounts,receipt_lines
from riju.sales_import import read_sales_xlsx,stage_sales,approve_sales
from riju.storage import Vault


def workbook(records=None):
    w=Workbook();s=w.active;s.title='Sales'
    s.cell(3,2,'Date');s.cell(3,6,'Invoice No.')
    for col,(_,label,*_) in enumerate(MODES,8):s.cell(4,col,label)
    for col,label in enumerate(('Sub total','Card','UPI','Cash','Other source','Total'),23):s.cell(4,col,label)
    if records is None:
        records=[{'revenues':{m[0]:100*(i+1) for i,m in enumerate(MODES)},'payments':{'cash':3000,'card':2000,'upi':5000,'other':2000}}]
    for row,record in enumerate(records,5):
        s.cell(row,1,row-4);s.cell(row,2,record.get('date',datetime(2026,6,1)));s.cell(row,4,'Invented patient')
        s.cell(row,6,record.get('invoice',f'INV-{row}'));s.cell(row,7,record.get('receipt',f'R-{row}'))
        for col,(key,*_) in enumerate(MODES,8):s.cell(row,col,record['revenues'].get(key,0))
        s.cell(row,23,record.get('subtotal',sum(record['revenues'].values())))
        for col,key in enumerate(('card','upi','cash','other'),24):s.cell(row,col,record['payments'].get(key,0))
        s.cell(row,28,sum(record['payments'].values()))
    out=io.BytesIO();w.save(out);w.close();return out.getvalue()

class RevenueTests(unittest.TestCase):
    def setUp(self):self.s=core.empty_state()
    def test_all_fifteen_categories_all_payment_methods_all_units(self):
        self.assertEqual(len(MODES),15)
        for key,label,*_ in MODES:
            for payment,code in [('cash','1000'),('card','1120'),('upi','1110'),('bank','1010')]:
                for branch in core.BRANCHES:
                    text=f'Received {payment} INR1291 for {label}'
                    with self.subTest(text=text,branch=branch):
                        lines=core.parse_transaction(text,branch,'2026-06-01',self.s)
                        self.assertEqual(lines,[dict(account=code,branch=branch,debit=129100,credit=0),dict(account=self.s['revenue_mode_accounts'][key],branch=branch,debit=0,credit=129100)])
    def test_specific_phrases_do_not_double_classify(self):
        for text,key in [('Medicine Treatment','medicine_treatment'),('treatment medicines','medicine_treatment'),('treatment payment received','treatment_receipt'),('Food expenses','food'),('Duty doctor charges','duty_doctor')]:
            lines=core.parse_transaction(f'Patient paid cash INR100 for {text}','JP','2026-06-01',self.s)
            self.assertEqual(lines[1]['account'],self.s['revenue_mode_accounts'][key])
    def test_multi_category_language_needs_amount_split(self):
        with self.assertRaises(ValueError):core.parse_transaction('Received cash INR1591 for consultation and medicines','JP','2026-06-01',self.s)
    def test_compound_sales_with_split_payment(self):
        lines=receipt_lines(self.s,'KM',{'consultation':30000,'medicine_retail':129100},{'cash':59100,'upi':100000})
        core.post(self.s,'2026-06-01','Sales',lines)
        b=core.balances(self.s,'2026-06-30','KM')
        self.assertEqual(b['1000'],59100);self.assertEqual(b['1110'],100000)
        self.assertEqual(b['4000'],-30000);self.assertEqual(b['4030'],-129100)
        self.assertEqual(sum(b.values()),0)
    def test_treatment_receipts_earned_and_other_clearing(self):
        lines=receipt_lines(self.s,'JP',{'treatment_receipt':10000},{'other':10000})
        self.assertEqual(lines[0]['account'],self.s['revenue_mode_accounts']['other_receipts'])
        self.assertEqual(lines[1]['account'],self.s['revenue_mode_accounts']['treatment_receipt'])
        self.assertNotEqual(lines[1]['account'],'2300')
    def test_explicit_advance_remains_liability(self):
        lines=receipt_lines(self.s,'JP',{'treatment_receipt':10000},{'cash':10000},treatment_policy='advance')
        self.assertEqual(lines[1]['account'],'2300')
        with self.assertRaises(ValueError):receipt_lines(self.s,'JP',{'treatment_receipt':10000},{'cash':10000},treatment_policy='review')
    def test_invalid_compound_collections(self):
        for rev,pay in [({'consultation':100},{'cash':99}),({'food':-100},{'cash':-100}),({'unknown':100},{'cash':100}),({'food':0},{'cash':0})]:
            with self.assertRaises(ValueError):receipt_lines(self.s,'JP',rev,pay)
    def test_legacy_migration_preserves_vouchers_and_custom_code(self):
        old=core.empty_state();old.pop('revenue_mode_accounts')
        old['accounts']=[dict(code=c,name=n,kind=k) for c,n,k in core.ACCOUNTS]
        old['accounts'].append(dict(code='4050',name='My custom ledger',kind='Equity'))
        core.post(old,'2026-06-01','Old journal',[dict(account='1010',branch='JP',debit=100,credit=0),dict(account='3000',branch='JP',debit=0,credit=100)])
        original=copy.deepcopy(old['vouchers'])
        with tempfile.TemporaryDirectory() as d:
            v=Vault(Path(d)/'legacy.json',demo=True);v.load();v.save(old)
            migrated=v.load();core.verify_audit(migrated)
            self.assertEqual(migrated['vouchers'],original)
            self.assertNotEqual(migrated['revenue_mode_accounts']['eye_check'],'4050')
            self.assertIn(dict(code='4050',name='My custom ledger',kind='Equity'),migrated['accounts'])
            count=len(migrated['audit']);ensure_accounts(migrated,audit=True)
            self.assertEqual(len(migrated['audit']),count)
            v.save(migrated);self.assertEqual(v.load(),migrated)

class SalesImportTests(unittest.TestCase):
    def setUp(self):self.s=core.empty_state()
    def test_all_modes_including_treatment_compound_post(self):
        raw=workbook();rows,warnings=read_sales_xlsx(raw,'VK')
        self.assertEqual(len(rows),1);self.assertFalse(warnings)
        self.assertEqual(rows[0]['revenues']['treatment'],10000)
        self.assertEqual(sum(rows[0]['revenues'].values()),1200000)
        q=stage_sales(self.s,raw,'invented.xlsx','VK')
        self.assertFalse(self.s['vouchers'])
        approve_sales(self.s,q['id']);core.verify_audit(self.s)
        self.assertEqual(len(self.s['vouchers']),1)
        self.assertEqual(sum(core.balances(self.s,'2026-06-30','VK').values()),0)
        credits=[l for l in self.s['vouchers'][0]['lines'] if l['credit']]
        self.assertEqual(len(credits),15)
        self.assertTrue(all(l['branch']=='VK' for l in credits))
    def test_inconsistent_totals_rejected_including_omitted_treatment(self):
        raw=workbook([{'revenues':{'treatment':500,'consultation':300},'payments':{'cash':300},'subtotal':300}])
        with self.assertRaisesRegex(ValueError,'all 15 service columns'):read_sales_xlsx(raw,'JP')
        raw=workbook([{'revenues':{'treatment':500,'consultation':300},'payments':{'cash':800},'subtotal':300}])
        with self.assertRaisesRegex(ValueError,'Sub total disagrees'):read_sales_xlsx(raw,'JP')
    def test_duplicate_source_receipts_not_double_posted(self):
        raw=workbook();q=stage_sales(self.s,raw,'a.xlsx','JP');approve_sales(self.s,q['id'])
        for action in (lambda:stage_sales(self.s,raw,'renamed.xlsx','JP'),lambda:approve_sales(self.s,q['id'])):
            with self.assertRaises(ValueError):action()
        modified=workbook([{'revenues':{'consultation':900},'payments':{'cash':900}}])
        with self.assertRaisesRegex(ValueError,'already queued or posted'):stage_sales(self.s,modified,'different.xlsx','JP')
        self.assertEqual(len(self.s['vouchers']),1)
    def test_atomic_approval_if_later_date_is_locked(self):
        records=[{'date':datetime(2026,6,2),'revenues':{'food':100},'payments':{'cash':100}},
                 {'date':datetime(2026,6,1),'revenues':{'food':200},'payments':{'cash':200}}]
        q=stage_sales(self.s,workbook(records),'a.xlsx','JP');self.s['locked_through']='2026-06-01';before=copy.deepcopy(self.s)
        with self.assertRaises(ValueError):approve_sales(self.s,q['id'])
        self.assertEqual(self.s,before)
    def test_duplicate_rows_invalid_dates_formulas_and_negative_amounts(self):
        for records in [
            [{'receipt':'same','invoice':'same','revenues':{'food':100},'payments':{'cash':100}}]*2,
            [{'date':'bad date','revenues':{'food':100},'payments':{'cash':100}}],
            [{'revenues':{'food':-100},'payments':{'cash':-100}}],
        ]:
            with self.subTest(records=records):
                with self.assertRaises(ValueError):read_sales_xlsx(workbook(records),'JP')
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(workbook()));wb.active['H5']='=100';out=io.BytesIO();wb.save(out)
        with self.assertRaisesRegex(ValueError,'numeric values'):read_sales_xlsx(out.getvalue(),'JP')
    def test_zero_blank_formula_totals_and_empty_report(self):
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(workbook()));wb.active['W5']='=SUM(H5:V5)';wb.active['AB5']='=SUM(X5:AA5)'
        out=io.BytesIO();wb.save(out);rows,warnings=read_sales_xlsx(out.getvalue(),'JP')
        self.assertEqual(len(rows),1);self.assertEqual(len(warnings),2)
        with self.assertRaises(ValueError):read_sales_xlsx(workbook([]),'JP')
    def test_non_sales_sheet_not_silently_ignored(self):
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(workbook()));wb.create_sheet('Other')['A1']='unrecognized data';out=io.BytesIO();wb.save(out)
        with self.assertRaisesRegex(ValueError,'No sheets were silently skipped'):read_sales_xlsx(out.getvalue(),'JP')

if __name__=='__main__':unittest.main()
