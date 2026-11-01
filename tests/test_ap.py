import copy
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader
from riju import core,foundation,ap,cash_flow as cf,financial_reports as f,history
from riju.storage import Vault

def pair(dr,cr,n,unit='JP'):return [dict(account=dr,branch=unit,debit=n,credit=0),dict(account=cr,branch=unit,debit=0,credit=n)]

class ApTests(unittest.TestCase):
    def setUp(self):
        self.s=core.empty_state();self.vendor=ap.save_vendor(self.s,dict(name='Synthetic Vendor',credit_days='30'));core.post(self.s,'2026-04-01','Opening bank',pair('1010','3000',1000000),opening=True)
    def enter(self,**kw):
        v=dict(vendor=self.vendor['id'],branch='JP',reference='INV-01',invoice_date='2026-04-02',received_date='2026-04-03',accounting_date='2026-04-03',credit_days='30',due_basis='Receipt date',amount='1000',payable='2000',allocations=[dict(account='4590',amount='1000')],tax_treatment='No separate GST/TDS; reviewed',reviewed=True,basis='Synthetic reviewed source');v.update(kw);return ap.enter_invoice(self.s,v)
    def pay(self,r,**kw):
        v=dict(date='2026-05-03',reference='PAY-01',bank_account='1010',items=[dict(invoice=r['id'],amount='1000')],reviewed=True,basis='Synthetic actual bank evidence');v.update(kw);return ap.pay(self.s,v)
    def test_accrual_dates_credit_due_payment_journals_and_positions(self):
        r=self.enter();self.assertEqual(r['due_date'],'2026-05-03');self.assertEqual(len(self.s['vouchers']),2)
        self.assertEqual(core.balances(self.s,'2026-04-30')['2000'],-100000);self.assertEqual(core.balances(self.s,'2026-04-30')['1010'],1000000)
        self.assertTrue(ap.register(self.s,'2026-05-03')[0]['overdue']);p=self.pay(r);self.assertEqual(ap.position(self.s,r)['status'],'Paid');self.assertEqual(core.balances(self.s,'2026-05-31')['2000'],0);self.assertEqual(core.balances(self.s,'2026-05-31')['1010'],900000)
        self.assertEqual(history.trace_voucher(self.s,p['voucher'])['ap_context']['invoices'][0]['id'],r['id']);core.verify_audit(self.s)
    def test_partial_multiple_invoice_payment_and_atomic_capacity(self):
        a=self.enter();b=self.enter(reference='INV-02');self.pay(a,items=[dict(invoice=a['id'],amount='400')]);self.assertEqual(ap.position(self.s,a)['remaining'],60000)
        self.pay(a,reference='PAY-02',items=[dict(invoice=a['id'],amount='600'),dict(invoice=b['id'],amount='1000')]);self.assertEqual(ap.position(self.s,b)['remaining'],0)
        before=copy.deepcopy(self.s)
        with self.assertRaises(ValueError):self.pay(a,reference='PAY-03')
        self.assertEqual(before,self.s)
    def test_extensions_preserve_contractual_and_historical_forecast(self):
        r=self.enter();n=len(self.s['vouchers']);ap.extend(self.s,r['id'],'2026-05-02','2026-05-20','Vendor agreed',True)
        self.assertEqual(r['due_date'],'2026-05-03');self.assertEqual(ap.schedule(self.s,r,'2026-05-01'),'2026-05-03');self.assertEqual(ap.schedule(self.s,r,'2026-05-02'),'2026-05-20');self.assertEqual(n,len(self.s['vouchers']))
        self.assertEqual(ap.forecast(self.s,'2026-05-01','2026-05-01','2026-05-31')[0]['forecast_date'],'2026-05-03');self.assertEqual(ap.forecast(self.s,'2026-05-02','2026-05-02','2026-05-31')[0]['forecast_date'],'2026-05-20')
        with self.assertRaises(ValueError):ap.extend(self.s,r['id'],'2026-05-01','2026-06-01','Backdate',True)
    def test_existing_tax_accrual_and_existing_payment_link_without_repost(self):
        source=core.post(self.s,'2026-04-03','Reviewed GST and TDS source',[*pair('4590','2000',100000),*pair('1200','2000',18000),*pair('2000','2100',10000)])
        r=self.enter(amount='1080',existing_voucher=source['id']);self.assertEqual(len(self.s['vouchers']),2)
        paid=core.post(self.s,'2026-05-03','Actual previously recorded bank payment',pair('2000','1010',108000));n=len(self.s['vouchers']);self.pay(r,existing_voucher=paid['id'],items=[dict(invoice=r['id'],amount='1080')]);self.assertEqual(len(self.s['vouchers']),n);self.assertEqual(ap.position(self.s,r)['remaining'],0)
        before=copy.deepcopy(self.s)
        with self.assertRaises(ValueError):self.enter(reference='INV-02',amount='1',existing_voucher=source['id'])
        self.assertEqual(before,self.s)
    def test_payment_reversal_dependency_chain_and_future_overpay_guard(self):
        r=self.enter();p=self.pay(r)
        with self.assertRaises(ValueError):core.reverse(self.s,r['voucher'],'2026-05-04')
        reverse=core.reverse(self.s,p['voucher'],'2026-05-04');self.assertEqual(ap.position(self.s,r)['remaining'],100000);core.reverse(self.s,r['voucher'],'2026-05-05')
        with self.assertRaises(ValueError):core.reverse(self.s,reverse['id'],'2026-05-06')
        self.assertEqual(ap.register(self.s,'2026-05-03')[0]['status'],'Paid');self.assertEqual(ap.register(self.s,'2026-05-05')[0]['status'],'Cancelled')
        # Future booked payment reserves capacity against a second backdated payment.
        self.s=core.empty_state();self.vendor=ap.save_vendor(self.s,dict(name='Second Vendor',credit_days='30'));r=self.enter();self.pay(r,date='2026-06-01')
        before=copy.deepcopy(self.s)
        with self.assertRaises(ValueError):self.pay(r,date='2026-05-01',reference='EARLY')
        self.assertEqual(before,self.s)
    def test_forecast_no_actual_cash_until_paid_and_partial_then_remove(self):
        r=self.enter();proposed=cf.defaults(self.s);cf.save_policy(self.s,dict(proposed,effective='2026-04-01',reviewed=True,note='Synthetic policy'))
        self.assertEqual(sum(cf.analyse(self.s,'2026-04-01','2026-04-30')['totals'].values()),0);self.assertEqual(ap.forecast(self.s,'2026-04-30','2026-05-01','2026-05-31')[0]['amount'],-100000)
        self.pay(r,items=[dict(invoice=r['id'],amount='400')]);self.assertEqual(ap.forecast(self.s,'2026-05-04','2026-05-04','2026-05-31')[0]['amount'],-60000);self.assertEqual(cf.analyse(self.s,'2026-05-01','2026-05-31')['totals']['Operating'],-40000)
        self.pay(r,reference='PAY-02',items=[dict(invoice=r['id'],amount='600')]);self.assertFalse(ap.forecast(self.s,'2026-05-04','2026-05-04','2026-05-31'))
    def test_asset_and_expense_cashflow_categories_shared_control(self):
        r=self.enter(allocations=[dict(account='1030',amount='700'),dict(account='4590',amount='300')]);proposed=cf.defaults(self.s);cf.save_policy(self.s,dict(proposed,effective='2026-04-01',reviewed=True,note='Synthetic classifications'));p=self.pay(r)
        result=cf.analyse(self.s,'2026-04-01','2026-05-31');self.assertEqual(result['totals']['Investing'],-70000);self.assertEqual(result['totals']['Operating'],-30000);self.assertEqual(result['cash_difference'],0);self.assertEqual(result['indirect_difference'],0)
        core.reverse(self.s,p['voucher'],'2026-05-04');self.assertEqual(sum(cf.analyse(self.s,'2026-04-01','2026-05-31')['totals'].values()),0)
    def test_duplicate_date_amount_tax_and_payment_guards_atomic(self):
        self.enter();before=copy.deepcopy(self.s)
        for v in [dict(reference='inv01'),dict(reference='BAD',amount='1.001'),dict(reference='BAD',received_date='2026-04-01'),dict(reference='BAD',tax_treatment='Taxable'),dict(reference='BAD',credit_days='1.5'),dict(reference='BAD',allocations=[dict(account='1010',amount='1000')]),dict(reference='BAD',reviewed=False)]:
            with self.subTest(v=v),self.assertRaises(ValueError):self.enter(**v)
        self.assertEqual(before,self.s)
        r=ap.data(self.s)['invoices'][0]
        for v in [dict(bank_account='1110'),dict(date='2026-04-02'),dict(reviewed=False)]:
            with self.assertRaises(ValueError):self.pay(r,**v)
        self.assertEqual(before,self.s)
    def test_register_forecast_excel_pdf_and_format19_upgrade(self):
        self.enter();before=copy.deepcopy(self.s)
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);packs=ap.snapshots(self.s,'2026-04-01','2026-05-31')+[ap.forecast_snapshot(self.s,'2026-04-30','2026-05-01','2026-05-31')];f.export_pack_xlsx(packs,root/'ap.xlsx');f.export_pdf(packs,root/'ap.pdf');wb=load_workbook(root/'ap.xlsx');self.assertIn('AP payment history',wb.sheetnames);wb.close();self.assertIn('Synthetic Vendor',' '.join(p.extract_text() for p in PdfReader(root/'ap.pdf').pages));self.assertEqual(before,self.s)
            history.checkpoint(self.s);v=Vault(root/'ap.json',demo=True);v.save(self.s);self.assertEqual(v.load(),self.s)

    def test_actual_format19_migration_original_backup_and_shared_cancel_guard(self):
        import json
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);external=root/'external';external.mkdir();old=core.empty_state();old['version']=19;old['schema_revision']=19;old.pop('ap');old['settings']['backup_dir']=str(external);core.post(old,'2026-04-01','Existing source',pair('1010','3000',10000),opening=True);history.checkpoint(old);path=root/'old.json';path.write_text(json.dumps(old),encoding='utf-8');original=path.read_bytes();v=Vault(path,demo=True);new=v.load();self.assertEqual(new['version'],20);self.assertEqual(new['vouchers'],old['vouchers']);self.assertFalse(new['ap']['invoices']);v.save(new);self.assertEqual(v.last_upgrade_backup.read_bytes(),original);self.assertEqual(v.load(),new)
        source=core.post(self.s,'2026-04-03','Shared cost source',pair('4590','2000',200000));a=self.enter(existing_voucher=source['id']);self.enter(existing_voucher=source['id'],reference='INV-02')
        before=copy.deepcopy(self.s)
        with self.assertRaises(ValueError):ap.cancel_invoice(self.s,a['id'],'2026-04-04','Cannot cancel one part',True)
        self.assertEqual(before,self.s)
