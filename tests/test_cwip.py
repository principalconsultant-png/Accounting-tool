import copy
import json
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
from pypdf import PdfReader
from riju import core, foundation, cwip, fixed_assets as fa, tax, tds, history
from riju.storage import Vault


class CwipTests(unittest.TestCase):
    def setUp(self):
        self.s=core.empty_state();self.p=cwip.save_project(self.s,'VK-1','Synthetic construction','VK','Synthetic approved project',True)
    def cost(self,**kwargs):
        v=dict(reference='BILL-1',party='Synthetic builder',category='Civil',date='2026-04-01',gross='1000',capital='900',retention='100',expense='4590',payable='2000',basis='Synthetic reviewed direct capital and operating cost',reviewed=True)
        v.update(kwargs);return cwip.record_cost(self.s,self.p['id'],v)
    def commission(self,c,amount='900',on='2026-05-01',**kwargs):
        a=dict(reference='PPE-1',name='Synthetic commissioned component',put_to_use=on,months=12,residual='0',ppe='1030',accum='1050',expense='4570')
        a.update(kwargs);return cwip.commission(self.s,self.p['id'],on,{c['id']:amount},a,'Synthetic readiness and cost certificate',True)
    def report(self,end='2026-05-31',branch='ALL'):return cwip.report(self.s,'2026-04-01',end,branch)
    def rule(self,rate='10'):
        p=tax.save_party(self.s,dict(name='Synthetic builder',effective='2026-04-01',category='Company',pan='ABCDE1234F'))
        tax.save_rule(self.s,dict(family='TDS',key='CWIP-TDS',label='Synthetic rule',effective='2026-04-01',section='194C',rate=rate,treatment='Withholding',party_category='Company',annual_threshold='0',threshold_basis='Financial-year aggregate',applicability='Synthetic only',source='Synthetic test',reviewed=True))
        return p
    def deduct(self,c,p,**kwargs):
        v=dict(date='2026-04-01',tax_date='2026-04-01',rule_key='CWIP-TDS',party_id=p['party_id'],base='1000',mode='Automatic',review_reason='Synthetic reviewed resident capital contractor',confirmed=True,resident_reviewed=True,narration='Synthetic withholding')
        v.update(kwargs);return cwip.deduct_tds(self.s,c['id'],v)
    def unchanged(self,action):
        old=copy.deepcopy(self.s)
        with self.assertRaises((ValueError,KeyError)):action()
        self.assertEqual(self.s,old)
    def test_project_five_separate_ledgers_no_default_budget_or_date(self):
        self.assertEqual(set(self.p['ledgers']),set(cwip.CATEGORIES));self.assertEqual(len(set(self.p['ledgers'].values())),5)
        self.assertTrue(all(v is None for v in cwip.plan(self.p,'2026-04-01')['budgets'].values()))
        self.assertEqual(self.s['branches']['VK']['opening_date'],'');self.assertEqual(self.s['branches']['VK']['status'],'Under construction')
        self.assertTrue(self.report()['ready']);core.verify_audit(self.s)
    def test_duplicate_project_and_failed_creation_are_atomic(self):
        self.unchanged(lambda:cwip.save_project(self.s,'vk-1','Duplicate','VK','Test',True))
        self.unchanged(lambda:cwip.save_project(self.s,'P','Missing unit','BAD','Test',True))
        self.unchanged(lambda:cwip.save_project(self.s,'P','No review','VK','Test',False))
    def test_optional_budget_zero_unknown_revisions_and_milestones(self):
        pid=self.p['id'];cwip.save_plan(self.s,pid,'2026-04-01',{'Civil':'2000','MEP':'0'},[dict(key='READY',name='Readiness review',due='2026-07-01',completed='')],'Original budget source',True)
        cwip.save_plan(self.s,pid,'2026-06-01',{'Civil':'3000'},[dict(key='READY',name='Readiness review',due='2026-08-01',completed='2026-07-30')],'Approved revision',True,'Renamed project')
        p=cwip.project(self.s,pid);self.assertEqual(cwip.plan(p,'2026-05-31')['budgets']['Civil'],200000);self.assertEqual(cwip.plan(p,'2026-05-31')['budgets']['MEP'],0)
        self.assertIsNone(cwip.plan(p,'2026-05-31')['budgets']['Architecture']);self.assertEqual(cwip.plan(p,'2026-06-01')['name'],'Renamed project')
        self.assertEqual(cwip.report(self.s,'2026-04-01','2026-06-30')['milestones'][0][4],'')
        self.assertEqual(cwip.report(self.s,'2026-04-01','2026-08-01')['milestones'][0][4],'2026-07-30')
    def test_invalid_plan_does_not_mutate(self):
        self.unchanged(lambda:cwip.save_plan(self.s,self.p['id'],'2026-04-01',{'Civil':'-1'},[],'Review',True))
        self.unchanged(lambda:cwip.save_plan(self.s,self.p['id'],'2026-04-01',{},[dict(key='A',name='A',due='bad')],'Review',True))
    def test_mixed_eligible_and_hospital_expense_does_not_double_cost(self):
        c=self.cost();b=core.balances(self.s,'2026-04-30','VK')
        self.assertEqual(b[c['cwip']],90000);self.assertEqual(b['4590'],10000);self.assertEqual(b['2000'],-90000);self.assertEqual(b[self.p['retention']],-10000)
        self.assertTrue(self.report()['ready']);self.assertEqual(self.report()['costs'][0][6:12],[100000,90000,10000,90000,10000,0])
    def test_all_expense_cost_does_not_create_cwip(self):
        c=self.cost(capital='0',retention='0');self.assertEqual(cwip._remaining(self.s,c),0);self.assertTrue(self.report()['ready']);self.unchanged(lambda:self.commission(c,'1'))
    def test_cost_invalid_review_amount_category_or_date_rolls_back(self):
        for v in [dict(reviewed=False),dict(capital='1001'),dict(capital='-1'),dict(retention='1001'),dict(gross='0'),dict(category='Other'),dict(date='bad'),dict(party=''),dict(expense='4570')]:self.unchanged(lambda v=v:self.cost(**v))
    def test_duplicate_supplier_ref_year_and_branch_isolation(self):
        self.cost();self.unchanged(lambda:self.cost(reference='bill-1',party='SYNTHETIC BUILDER'))
        self.cost(reference='BILL-1',date='2027-04-01');self.assertFalse(self.report(branch='JP')['costs']);self.assertEqual(len(self.report()['costs']),1)
    def test_link_existing_allocations_capacity_and_no_duplicate_post(self):
        code=self.p['ledgers']['Civil'];v=core.post(self.s,'2026-04-01','Synthetic existing',fa._pair(code,'2000',100000,'VK'))
        c=self.cost(mode='Link existing',source=v['id'],capital='1000',retention='0');self.assertEqual(len(self.s['vouchers']),1);self.assertTrue(self.report()['ready'])
        self.unchanged(lambda:self.cost(mode='Link existing',source=v['id'],reference='DUP',gross='1',capital='1',retention='0'))
    def test_partial_source_two_costs_fit_debit_and_payable_once(self):
        code=self.p['ledgers']['Civil'];v=core.post(self.s,'2026-04-01','Grouped costs',fa._pair(code,'2000',100000,'VK'))
        for ref in ('A','B'):self.cost(mode='Link existing',source=v['id'],reference=ref,gross='500',capital='500',retention='0')
        self.assertTrue(self.report()['ready']);self.assertEqual(len(self.s['vouchers']),1)
    def test_retention_partial_release_pay_and_reversal_proofs(self):
        c=self.cost();e=cwip.settle_retention(self.s,c['id'],'2026-04-10','40','2000','Reviewed release',True)
        pay=cwip.settle_retention(self.s,c['id'],'2026-04-11','30','1010','Reviewed payment',True)
        self.assertEqual(cwip.retained(self.s,c),3000);self.assertTrue(self.report()['ready'])
        core.reverse(self.s,pay['voucher'],'2026-04-12');self.assertEqual(cwip.retained(self.s,c),6000);self.assertTrue(self.report()['ready'])
    def test_retention_overuse_invalid_counter_and_backdate_are_atomic(self):
        c=self.cost()
        for on,n,code in [('2026-04-01','101','1010'),('2026-03-31','1','1010'),('2026-04-01','1','4590')]:self.unchanged(lambda:cwip.settle_retention(self.s,c['id'],on,n,code,'Review',True))
    def test_future_retention_release_blocks_backdated_reuse(self):
        c=self.cost();cwip.settle_retention(self.s,c['id'],'2026-06-01','100','2000','Release',True)
        self.unchanged(lambda:cwip.settle_retention(self.s,c['id'],'2026-05-01','1','1010','Backdated',True))
    def test_tds_integrates_b09_without_duplicate_cost_and_allocates_challan(self):
        c=self.cost();p=self.rule();r=self.deduct(c,p)
        self.assertEqual(r['withheld'],10000);self.assertEqual(core.balances(self.s,'2026-04-30','VK')[c['cwip']],90000)
        self.assertEqual(self.report()['costs'][0][11],10000);self.assertTrue(self.report()['ready'])
        payment=tds.post(self.s,'Challan',dict(date='2026-04-10',branch='VK',reference='UTR-1',narration='Synthetic challan',bank_account='1010',rows=[dict(section='194C',branch='VK',year=2026,tax='100')],confirmed=True))
        tds.allocate(self.s,r['id'],payment['id'],0,'100','2026-04-10')
        self.unchanged(lambda:core.reverse(self.s,r['voucher_id'],'2026-04-11'))
    def test_tds_and_retention_cannot_exceed_gross(self):
        c=self.cost(retention='950');p=self.rule();self.unchanged(lambda:self.deduct(c,p))
    def test_tds_duplicate_no_confirm_or_earlier_source_are_atomic(self):
        c=self.cost();p=self.rule();self.unchanged(lambda:self.deduct(c,p,confirmed=False));self.unchanged(lambda:self.deduct(c,p,date='2026-03-31'))
        self.deduct(c,p);self.unchanged(lambda:self.deduct(c,p))
    def test_cost_reversal_requires_retention_and_commission_reversal(self):
        c=self.cost();self.unchanged(lambda:core.reverse(self.s,c['voucher'],'2026-04-02'))
        ret=cwip.data(self.s)['events'][0];core.reverse(self.s,ret['voucher'],'2026-04-02');e=self.commission(c)
        self.unchanged(lambda:core.reverse(self.s,c['voucher'],'2026-05-02'))
        core.reverse(self.s,e['voucher'],'2026-05-02');core.reverse(self.s,c['voucher'],'2026-05-02');self.assertTrue(self.report()['ready'])
    def test_commission_partial_source_multiple_categories_and_asset_depreciation(self):
        c=self.cost();m=self.cost(reference='MEP-1',category='MEP',gross='500',capital='500',retention='0')
        e=cwip.commission(self.s,self.p['id'],'2026-05-01',{c['id']:'400',m['id']:'500'},dict(reference='PPE',name='Synthetic components',put_to_use='2026-05-16',months=12,residual='0'),'Ready and reviewed',True)
        a=fa.asset(self.s,e['asset']);self.assertEqual(a['cost'],90000);self.assertEqual(fa.target(a,'2026-05-15'),0);self.assertEqual(fa.target(a,'2026-05-31'),3871)
        self.assertEqual(cwip._remaining(self.s,c),50000);self.assertEqual(cwip._remaining(self.s,m),0);self.assertTrue(self.report()['ready'])
        fa.post_depreciation(self.s,'2026-05-31',branch='VK');self.assertTrue(fa.report(self.s,'2026-04-01','2026-05-31','VK')['ready'])
        self.unchanged(lambda:core.reverse(self.s,e['voucher'],'2026-05-31'))
    def test_commission_invalid_life_duplicate_ref_and_insufficient_cost_roll_back(self):
        c=self.cost();self.unchanged(lambda:self.commission(c,months=0));self.unchanged(lambda:self.commission(c,amount='901'));self.unchanged(lambda:self.commission(c,put_to_use='2026-04-30'))
        self.commission(c,amount='400');self.unchanged(lambda:self.commission(c,amount='100'))
    def test_future_commission_prevents_backdated_double_use(self):
        c=self.cost();self.commission(c,on='2026-06-01');self.unchanged(lambda:self.commission(c,'1',on='2026-05-01',reference='EARLIER'))
    def test_commission_cannot_use_another_project(self):
        c=self.cost();p2=cwip.save_project(self.s,'P2','Synthetic second','VK','Test',True)
        self.unchanged(lambda:cwip.commission(self.s,p2['id'],'2026-05-01',{c['id']:'900'},dict(reference='PPE',name='Wrong project',months=12),'Test',True))
    def test_reclassification_cannot_expense_commissioned_cost(self):
        c=self.cost();e=cwip.expense_remaining(self.s,c['id'],'2026-04-10','100','4590','Ineligible start-up cost',True)
        self.assertEqual(cwip._remaining(self.s,c),80000);self.assertTrue(self.report()['ready']);self.commission(c,'800')
        self.unchanged(lambda:cwip.expense_remaining(self.s,c['id'],'2026-05-02','1','4590','Double use',True))
        self.assertEqual(self.report('2026-04-09')['costs'][0][9],90000)
    def test_manual_or_legacy_cwip_differences_visible_and_block_commission(self):
        c=self.cost();core.post(self.s,'2026-04-02','Unallocated legacy',fa._pair('1020','2000',100,'VK'))
        self.assertFalse(self.report()['ready']);self.unchanged(lambda:self.commission(c));self.assertTrue(any(r[0]=='Unallocated' for r in self.report()['proof']))
    def test_exact_reversal_restore_keeps_asof_sources_and_capacity(self):
        c=self.cost(retention='0');e=self.commission(c);rev=core.reverse(self.s,e['voucher'],'2026-05-10')
        self.assertEqual(self.report('2026-05-09')['costs'][0][9],0);self.assertEqual(self.report('2026-05-10')['costs'][0][9],90000)
        core.reverse(self.s,rev['id'],'2026-05-12');self.assertEqual(self.report()['costs'][0][9],0);self.assertTrue(self.report()['ready'])
    def test_period_lock_and_inactive_accounts_apply(self):
        c=self.cost();self.s['locked_through']='2026-05-01';self.unchanged(lambda:self.commission(c));self.s['locked_through']=''
        self.s['account_details']['1030']['active']=False;self.unchanged(lambda:self.commission(c))
    def test_historical_project_ledger_groups_and_deletion_protected(self):
        code=self.p['ledgers']['Civil'];self.unchanged(lambda:foundation.save_account(self.s,code,'Changed','Asset','PPE'))
        foundation.save_account(self.s,code,'Civil revised','Asset','CWIP',active=False);self.unchanged(lambda:foundation.delete_account(self.s,code))
    def test_readonly_excel_pdf_numeric_blanks_literal_text_and_pagination(self):
        self.cost(reference='=SYNTHETIC',party='+Builder');before=copy.deepcopy(self.s)
        with tempfile.TemporaryDirectory() as d:
            root=Path(d)
            for kind in ('xlsx','pdf','json'):cwip.export(self.s,'2026-04-01','2026-05-31','VK',root/('cwip.'+kind),kind)
            book=load_workbook(root/'cwip.xlsx');self.assertEqual(len(book.sheetnames),6)
            ws=book.worksheets[1];rows=list(ws.values);self.assertTrue(any(900.0 in row for row in rows));self.assertTrue(any(None in row for row in rows));book.close()
            text=''.join(p.extract_text() for p in PdfReader(root/'cwip.pdf').pages);self.assertIn('CWIP ledger reconciliation',text);self.assertIn('Vaikom',text);self.assertIn('900.00',text)
            pack=cwip.snapshots(self.s,'2026-04-01','2026-05-31','VK');pack[-1]['rows']*=40
            from riju import financial_reports as fr
            fr.export_pdf(pack,root/'long.pdf');self.assertGreater(len(PdfReader(root/'long.pdf').pages),2)
        self.assertEqual(self.s,before)
    def test_schema14_migration_original_backup_and_reload(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);ext=root/'external';ext.mkdir();old=core.empty_state();old['version']=14;old['schema_revision']=14;old.pop('cwip');history.checkpoint(old);path=root/'old.json';path.write_text(json.dumps(old),encoding='utf-8');raw=path.read_bytes()
            v=Vault(path,demo=True);s=v.load();s['settings']['backup_dir']=str(ext);cwip.save_project(s,'P','Test','VK','Test',True);v.save(s)
            self.assertEqual(s['version'],15);self.assertEqual(v.last_upgrade_backup.read_bytes(),raw);self.assertEqual(v.load(),s)

    def test_existing_legacy_cost_retention_tds_link_without_duplicate_journals(self):
        # All three historical amounts already exist in a compound journal.
        p=cwip.save_project(self.s,'OLD','Existing Vaikom','VK','Reviewed historic controls',True,{'Civil':'1020'},'2010');self.p=p
        party=self.rule();code=tds._control(self.s,'194C')
        sv=core.post(self.s,'2026-04-01','Historical contractor bill',[dict(account='1020',branch='VK',debit=100000,credit=0),dict(account='2000',branch='VK',debit=0,credit=80000),dict(account='2010',branch='VK',debit=0,credit=10000),dict(account=code,branch='VK',debit=0,credit=10000)])
        c=self.cost(mode='Link existing',source=sv['id'],capital='1000',retention='0')
        cwip.link_retention(self.s,c['id'],sv['id'],'100','Reviewed old retention',True)
        self.deduct(c,party,kind='Register existing deduction',manual_tax='100',mode='Reviewed manual')
        self.assertEqual(len(self.s['vouchers']),1);self.assertTrue(self.report()['ready']);self.assertEqual(self.report()['costs'][0][10:12],[10000,10000])
        self.unchanged(lambda:cwip.link_retention(self.s,c['id'],sv['id'],'1','Duplicate',True));self.commission(c,'1000');self.assertTrue(self.report()['ready'])
    def test_existing_ledger_ownership_not_shared_across_projects(self):
        self.unchanged(lambda:cwip.save_project(self.s,'P2','Duplicate mapping','VK','Test',True,{'Civil':self.p['ledgers']['Civil']}))
    def test_generic_b09_deduction_cannot_exceed_source_less_retention(self):
        c=self.cost(retention='950');p=self.rule()
        self.unchanged(lambda:tds.post(self.s,'Existing payable',dict(date=c['date'],branch='VK',source_voucher=c['voucher'],reference='GENERIC',narration='Synthetic generic withholding',rule_key='CWIP-TDS',party_id=p['party_id'],base='1000',gross='1000',mode='Automatic',resident_reviewed=True,confirmed=True,counter_account='2000')))

    def test_existing_b09_expense_source_links_without_registering_tax_twice(self):
        p=self.rule();r=tds.post(self.s,'Expense',dict(date='2026-04-01',branch='VK',reference='B09-SOURCE',narration='Synthetic already recorded contractor cost',rule_key='CWIP-TDS',party_id=p['party_id'],base='1000',gross='1000',account=self.p['ledgers']['Civil'],counter_account='2000',resident_reviewed=True,confirmed=True))
        c=self.cost(mode='Link existing',source=r['voucher_id'],capital='1000',retention='0')
        self.unchanged(lambda:self.deduct(c,p,kind='Register existing deduction',reference='COPY',manual_tax='100',mode='Reviewed manual'))
        cwip.link_tds(self.s,c['id'],r['id'],'Already recorded by B09; link only',True)
        self.assertEqual(len(self.s['vouchers']),1);self.assertEqual(len(self.s['tds_records']),1);self.assertEqual(self.report()['costs'][0][11],10000)
        self.unchanged(lambda:cwip.link_tds(self.s,c['id'],r['id'],'Duplicate',True))
    def test_voucher_history_links_cost_commission_and_original_journal(self):
        c=self.cost();e=self.commission(c)
        traced=history.trace_voucher(self.s,e['voucher'])['cwip_context'];self.assertEqual(traced['costs'][0]['id'],c['id']);self.assertEqual(traced['events'][-1]['asset'],e['asset']);self.assertTrue(history.integrity_report(self.s)['passed'])
    def test_non_depreciable_land_can_be_commissioned_with_reviewed_residual(self):
        c=self.cost(capital='1000',retention='0');e=self.commission(c,'1000',category='Land — non-depreciable',months=0,residual='1000');a=fa.asset(self.s,e['asset']);self.assertEqual(fa.target(a,'2030-01-01'),0)
