import copy
import json
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
from pypdf import PdfReader
from riju import core,foundation,fixed_assets as fa,history
from riju.storage import Vault


def pair(dr,cr,n,branch='JP'):return [dict(account=dr,branch=branch,debit=n,credit=0),dict(account=cr,branch=branch,debit=0,credit=n)]

class FixedAssetTests(unittest.TestCase):
    def setUp(self):self.s=core.empty_state()
    def create(self,**extra):
        v=dict(reference='EQ-1',name='Synthetic equipment',branch='JP',category='Manual / specific Schedule II category',mode='Post purchase',acquired='2026-04-01',put_to_use='2026-04-01',baseline='2026-04-01',cost='1200',opening_accum='0',residual='0',months=12,ppe='1030',accum='1050',expense='4570',counter='2000',basis='Synthetic reviewed single-shift SLM useful life / source',reviewed=True)
        v.update(extra);return fa.register(self.s,v)
    def result(self,end='2026-04-30',branch='ALL'):return fa.report(self.s,'2026-04-01',end,branch)
    def test_full_month_slm_posting_and_proof_are_exact(self):
        a=self.create();self.assertEqual(fa.target(a,'2026-04-30'),10000);self.assertFalse(self.result()['ready'])
        e=fa.post_depreciation(self.s,'2026-04-30');self.assertEqual(e['items'][0]['depreciation'],10000);self.assertTrue(self.result()['ready']);self.assertFalse(any(row[-1] for row in self.result()['proof']));core.verify_audit(self.s)
    def test_put_to_use_partial_month_and_residual_cap(self):
        a=self.create(put_to_use='2026-04-16');self.assertEqual(fa.target(a,'2026-04-15'),0);self.assertEqual(fa.target(a,'2026-04-30'),5000);self.assertEqual(fa.target(a,'2026-05-31'),15000)
        self.assertEqual(fa.target(a,'2027-04-15'),120000);self.assertEqual(fa.target(a,'2030-01-01'),120000)
    def test_residual_and_rounding_do_not_drift_over_months(self):
        a=self.create(cost='1000.01',residual='50',months=3)
        amounts=[]
        for end in ('2026-04-30','2026-05-31','2026-06-30'):amounts.append(fa.post_depreciation(self.s,end)['items'][0]['depreciation'])
        self.assertEqual(sum(amounts),95001);self.assertEqual(fa.accumulated(self.s,a,'2026-06-30'),95001);self.assertEqual(self.result('2026-06-30')['register'][0][6],5000)
    def test_catchup_only_posts_delta_and_duplicate_run_is_atomic(self):
        self.create();fa.post_depreciation(self.s,'2026-04-30');e=fa.post_depreciation(self.s,'2026-06-30','2026-07-02');self.assertEqual(e['items'][0]['depreciation'],20000)
        before=copy.deepcopy(self.s)
        with self.assertRaisesRegex(ValueError,'No unposted'):fa.post_depreciation(self.s,'2026-06-30','2026-07-02')
        self.assertEqual(before,self.s)
    def test_opening_asset_links_existing_journal_without_duplicate_cost(self):
        lines=pair('1030','3000',120000);lines[1]['credit']=60000;lines.append(dict(account='1050',branch='JP',debit=0,credit=60000));v=core.post(self.s,'2026-04-01','Opening',lines,opening=True)
        a=self.create(mode='Link existing',acquired='2020-01-01',put_to_use='2020-01-01',cost_voucher=v['id'],opening_accum='600',months=6,override_basis='Reviewed opening carrying value / remaining life')
        self.assertEqual(len(self.s['vouchers']),1);fa.post_depreciation(self.s,'2026-04-30');self.assertEqual(fa.accumulated(self.s,a,'2026-04-30'),70000);self.assertTrue(self.result()['ready'])
    def test_source_allocation_capacity_and_component_additions(self):
        v=core.post(self.s,'2026-04-01','Grouped acquisition',pair('1030','2000',120000))
        self.create(reference='A',mode='Link existing',cost_voucher=v['id'],cost='600');self.create(reference='B',mode='Link existing',cost_voucher=v['id'],cost='600')
        before=copy.deepcopy(self.s)
        with self.assertRaises(ValueError):self.create(reference='C',mode='Link existing',cost_voucher=v['id'],cost='1')
        self.assertEqual(before,self.s);fa.post_depreciation(self.s,'2026-04-30');self.assertTrue(self.result()['ready'])
    def test_review_overrides_and_invalid_dates_are_atomic(self):
        before=copy.deepcopy(self.s)
        for values in [dict(reviewed=False),dict(basis=''),dict(put_to_use='2026-03-31'),dict(cost='0'),dict(residual='1201'),dict(months=0),dict(category='Furniture and fittings — reference'),dict(residual='100'),dict(counter='4500')]:
            with self.assertRaises(ValueError):self.create(**values)
        self.assertEqual(before,self.s)
        a=self.create(category='Furniture and fittings — reference',months=120,residual='60');self.assertEqual(a['estimates'][0]['months'],120)
    def test_land_is_nondepreciable(self):
        a=self.create(category='Land — non-depreciable',months=0,residual='1200');self.assertEqual(fa.target(a,'2030-12-31'),0);self.assertTrue(self.result()['ready'])
        with self.assertRaises(ValueError):fa.post_depreciation(self.s,'2026-04-30')
    def test_monthly_branch_batch_is_balanced_by_unit(self):
        self.create();self.create(reference='KM-1',branch='KM');e=fa.post_depreciation(self.s,'2026-04-30');v=self.s['vouchers'][-1]
        self.assertEqual(len(e['items']),2)
        for unit in ('JP','KM'):self.assertEqual(sum(l['debit']-l['credit'] for l in v['lines'] if l['branch']==unit),0)
        self.assertTrue(self.result()['ready'])
    def test_manual_or_unregistered_ledger_movements_block_duplicate_depreciation(self):
        self.create();core.post(self.s,'2026-04-30','Manual depreciation',pair('4570','1050',10000));before=copy.deepcopy(self.s)
        with self.assertRaisesRegex(ValueError,'Reconcile'):fa.post_depreciation(self.s,'2026-04-30')
        self.assertEqual(before,self.s);self.assertFalse(self.result()['ready'])
    def test_prospective_estimate_preserves_earlier_schedule(self):
        a=self.create();fa.post_depreciation(self.s,'2026-04-30');old=fa.target(a,'2026-04-30')
        fa.revise_estimate(self.s,a['id'],'2026-05-01',22,'0','Technical review revised remaining life',True)
        current=fa.asset(self.s,a['id']);self.assertEqual(fa.target(current,'2026-04-30'),old);self.assertEqual(fa.target(current,'2026-05-31'),15000)
        fa.post_depreciation(self.s,'2026-05-31');self.assertTrue(self.result('2026-05-31')['ready'])
    def test_estimate_requires_prior_depreciation_and_dependent_withdrawal_guard(self):
        a=self.create()
        with self.assertRaisesRegex(ValueError,'previous|before'):fa.revise_estimate(self.s,a['id'],'2026-05-01',10,'0','Review',True)
        first=fa.post_depreciation(self.s,'2026-04-30');fa.revise_estimate(self.s,a['id'],'2026-05-01',10,'0','Review',True)
        with self.assertRaises(ValueError):core.reverse(self.s,first['voucher'],'2026-05-02')
        may=fa.post_depreciation(self.s,'2026-05-31')
        with self.assertRaises(ValueError):fa.withdraw_estimate(self.s,a['id'],'Change')
        core.reverse(self.s,may['voucher'],'2026-06-01');fa.withdraw_estimate(self.s,a['id'],'Correction');self.assertEqual(fa.target(fa.asset(self.s,a['id']),'2026-05-31'),20000)
    def test_transfer_moves_gross_and_accumulated_and_allocates_expense_by_unit(self):
        a=self.create();fa.post_depreciation(self.s,'2026-04-30');fa.transfer(self.s,a['id'],'2026-05-01','KM','Ownership transfer approved',True);fa.post_depreciation(self.s,'2026-05-31')
        self.assertEqual(fa.branch_at(self.s,a,'2026-05-31'),'KM');self.assertTrue(self.result('2026-05-31')['ready']);self.assertEqual(core.balances(self.s,'2026-05-31','JP')['1030'],0);self.assertEqual(core.balances(self.s,'2026-05-31','KM')['1050'],-20000)
        jp=self.result('2026-05-31','JP');km=self.result('2026-05-31','KM');self.assertEqual(sum(r[4] for r in jp['schedule']),10000);self.assertEqual(sum(r[4] for r in km['schedule']),10000);self.assertEqual(sum(r[5] for r in jp['schedule']),10000)
    def test_midmonth_transfer_charges_daily_portions_to_each_owner(self):
        a=self.create();fa.transfer(self.s,a['id'],'2026-04-16','KM','Midmonth ownership change',True);fa.post_depreciation(self.s,'2026-04-30')
        self.assertEqual(core.balances(self.s,'2026-04-30','JP')['4570'],5000);self.assertEqual(core.balances(self.s,'2026-04-30','KM')['4570'],5000);self.assertTrue(self.result()['ready'])
    def test_disposal_catches_up_removes_cost_accum_and_records_gain(self):
        a=self.create();e=fa.dispose(self.s,a['id'],'2026-04-30','1150','1010','','','Reviewed proceeds, separate taxes',True)
        self.assertEqual(e['gain_loss'],5000);self.assertEqual(e['items'][0]['depreciation'],10000);b=core.balances(self.s,'2026-04-30');self.assertEqual(b['1030'],0);self.assertEqual(b['1050'],0);self.assertTrue(self.result()['ready'])
        with self.assertRaises(ValueError):fa.post_depreciation(self.s,'2026-05-31')
    def test_scrapping_records_loss_and_auto_ledgers_do_not_reuse_occupied_codes(self):
        foundation.save_account(self.s,'4400','Unrelated revenue','Revenue','Other revenue');a=self.create();e=fa.dispose(self.s,a['id'],'2026-04-30','0','1010','','','Approved writeoff, reviewed tax',True)
        self.assertEqual(e['gain_loss'],-110000);self.assertNotEqual(self.s['asset_ledgers']['gain'],'4400');self.assertTrue(self.result()['ready'])
    def test_reversal_dependencies_and_dated_reporting_are_preserved(self):
        a=self.create();apr=fa.post_depreciation(self.s,'2026-04-30');may=fa.post_depreciation(self.s,'2026-05-31')
        with self.assertRaises(ValueError):core.reverse(self.s,apr['voucher'],'2026-06-01')
        with self.assertRaises(ValueError):core.reverse(self.s,a['cost_voucher'],'2026-06-01')
        core.reverse(self.s,may['voucher'],'2026-06-01');self.assertEqual(fa.accumulated(self.s,a,'2026-05-31'),20000);self.assertEqual(fa.accumulated(self.s,a,'2026-06-01'),10000)
        june=self.result('2026-06-30');self.assertTrue(any(r[5]==-10000 for r in june['schedule'] if r[2]=='2026-06-01'))
        fa.post_depreciation(self.s,'2026-06-30');self.assertTrue(self.result('2026-06-30')['ready'])
    def test_disposal_and_transfer_reverse_before_earlier_charges(self):
        a=self.create();apr=fa.post_depreciation(self.s,'2026-04-30');moved=fa.transfer(self.s,a['id'],'2026-05-01','KM','Move',True);sale=fa.dispose(self.s,a['id'],'2026-05-31','500','1010','','','Sale',True)
        with self.assertRaises(ValueError):core.reverse(self.s,moved['voucher'],'2026-06-01')
        core.reverse(self.s,sale['voucher'],'2026-06-01');core.reverse(self.s,moved['voucher'],'2026-06-01');self.assertEqual(fa.branch_at(self.s,a,'2026-06-01'),'JP');core.reverse(self.s,apr['voucher'],'2026-06-01');core.reverse(self.s,a['cost_voucher'],'2026-06-01');self.assertFalse(fa.present(self.s,a,'2026-06-01'))
    def test_locked_dates_and_failed_disposal_are_atomic(self):
        a=self.create();self.s['locked_through']='2026-04-30';before=copy.deepcopy(self.s)
        with self.assertRaises(ValueError):fa.post_depreciation(self.s,'2026-04-30')
        with self.assertRaises(ValueError):fa.dispose(self.s,a['id'],'2026-04-30','10','1010','','','Sale',True)
        self.assertEqual(before,self.s)
    def test_later_coverage_and_future_history_prevent_backdated_rework(self):
        a=self.create();fa.post_depreciation(self.s,'2026-05-31');before=copy.deepcopy(self.s)
        with self.assertRaises(ValueError):fa.post_depreciation(self.s,'2026-04-30','2026-06-01')
        with self.assertRaises(ValueError):fa.transfer(self.s,a['id'],'2026-05-01','KM','Backdated',True)
        self.assertEqual(before,self.s)
    def test_legacy_calculator_rows_are_retained_until_explicit_source_review(self):
        self.s['assets'].append(dict(name='Legacy',branch='JP',date='2026-04-01',cost=120000,residual=0,life=12,tax_rate=15));self.assertFalse(self.result()['ready'])
        self.create(legacy_index=0);fa.post_depreciation(self.s,'2026-04-30');self.assertTrue(self.result()['ready']);self.assertEqual(len(self.s['assets']),1)
    def test_excel_pdf_json_export_are_readonly_and_numeric(self):
        self.create();fa.post_depreciation(self.s,'2026-04-30');before=copy.deepcopy(self.s)
        with tempfile.TemporaryDirectory() as d:
            root=Path(d)
            for kind in ('xlsx','pdf','json'):fa.export(self.s,'2026-04-01','2026-04-30','ALL',root/('assets.'+kind),kind)
            w=load_workbook(root/'assets.xlsx');self.assertIn('PPE ledger reconciliation',w.sheetnames);self.assertTrue(any(c.data_type=='n' for row in w['Fixed asset register'] for c in row));w.close();text=' '.join(p.extract_text() for p in PdfReader(root/'assets.pdf').pages);self.assertIn('RECONCILED',text);self.assertIn('Monthly book depreciation',text)
        self.assertEqual(before,self.s)
    def test_schema_twelve_upgrade_preserves_original_backup_and_asset_events(self):
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);ext=root/'external';ext.mkdir();old=core.empty_state();old['version']=12;old['schema_revision']=12
            for key in ('book_assets','fixed_asset_events','asset_ledgers'):old.pop(key)
            history.checkpoint(old);path=root/'old.json';path.write_text(json.dumps(old),encoding='utf-8');raw=path.read_bytes();v=Vault(path,demo=True);self.s=v.load();self.s['settings']['backup_dir']=str(ext);self.create();fa.post_depreciation(self.s,'2026-04-30');v.save(self.s);self.assertEqual(v.last_upgrade_backup.read_bytes(),raw);self.assertEqual(v.load(),self.s);self.assertEqual(self.s['version'],15)
    def test_extra_shift_usage_is_reviewed_dated_and_capped(self):
        a=self.create();fa.save_shift_usage(self.s,a['id'],'2026-04-01','2026-04-15','1.5','Reviewed actual double-shift use; confirmed not NESD',True)
        self.assertEqual(fa.target(fa.asset(self.s,a['id']),'2026-04-30'),12500);fa.post_depreciation(self.s,'2026-04-30')
        with self.assertRaises(ValueError):fa.withdraw_shift_usage(self.s,a['id'],'Change')
        fa.save_shift_usage(self.s,a['id'],'2026-05-01','2027-03-31','2','Reviewed actual triple shift, eligible',True);fa.post_depreciation(self.s,'2027-03-31');self.assertEqual(fa.accumulated(self.s,a,'2027-03-31'),120000)
    def test_shift_eligibility_overlap_and_withdrawal_are_controlled(self):
        a=self.create();before=copy.deepcopy(self.s)
        with self.assertRaises(ValueError):fa.save_shift_usage(self.s,a['id'],'2026-04-01','2026-04-30','1.5','Not verified NESD',False)
        self.assertEqual(before,self.s);fa.save_shift_usage(self.s,a['id'],'2026-04-01','2026-04-30','1.5','Reviewed eligible',True)
        with self.assertRaises(ValueError):fa.save_shift_usage(self.s,a['id'],'2026-04-15','2026-05-15','2','Overlap',True)
        fa.withdraw_shift_usage(self.s,a['id'],'Usage correction');self.assertEqual(fa.target(fa.asset(self.s,a['id']),'2026-04-30'),10000)
    def test_leap_day_terminal_charge_reaches_residual_exactly(self):
        a=self.create(acquired='2024-02-29',put_to_use='2024-02-29',baseline='2024-02-29',cost='1000',residual='50',months=12)
        self.assertEqual(fa.target(a,'2025-02-27'),95000);self.assertEqual(fa.target(a,'2025-02-28'),95000)
    def test_future_put_to_use_does_not_charge_before_commissioning(self):
        a=self.create(put_to_use='2026-06-15');self.assertEqual(fa.target(a,'2026-05-31'),0);self.assertTrue(self.result()['ready'])
        with self.assertRaises(ValueError):fa.post_depreciation(self.s,'2026-05-31')
    def test_reversed_acquisition_keeps_history_and_can_be_restored(self):
        a=self.create();reverse=core.reverse(self.s,a['cost_voucher'],'2026-05-01');self.assertTrue(fa.present(self.s,a,'2026-04-30'));self.assertFalse(fa.present(self.s,a,'2026-05-31'));self.assertTrue(self.result('2026-05-31')['ready'])
        core.reverse(self.s,reverse['id'],'2026-05-02');self.assertTrue(fa.present(self.s,a,'2026-05-31'));fa.post_depreciation(self.s,'2026-05-31');self.assertTrue(self.result('2026-05-31')['ready'])
    def test_asset_master_group_and_history_are_protected(self):
        a=self.create();before=copy.deepcopy(self.s)
        with self.assertRaises(ValueError):foundation.save_account(self.s,'1030','Medical equipment','Asset','Other assets')
        self.assertEqual(before,self.s)
    def test_prior_fiscal_year_bookkeeping_and_year_locks(self):
        a=self.create(acquired='2020-04-01',put_to_use='2020-04-01',baseline='2020-04-01');fa.post_depreciation(self.s,'2021-03-31');self.assertTrue(fa.report(self.s,'2020-04-01','2021-03-31')['ready']);foundation.close_year(self.s,2020)
        before=copy.deepcopy(self.s)
        with self.assertRaises(ValueError):core.reverse(self.s,self.s['fixed_asset_events'][0]['voucher'],'2021-03-31')
        self.assertEqual(before,self.s)
    def test_prior_posted_depreciation_links_without_duplicate_expense(self):
        a=self.create();v=core.post(self.s,'2026-04-30','Existing manual depreciation',pair('4570','1050',10000));count=len(self.s['vouchers']);fa.link_prior_depreciation(self.s,a['id'],v['id'],'100','2026-04-30','Reviewed earlier manual journal',True)
        self.assertEqual(len(self.s['vouchers']),count);self.assertTrue(self.result()['ready']);before=copy.deepcopy(self.s)
        with self.assertRaises(ValueError):fa.link_prior_depreciation(self.s,a['id'],v['id'],'100','2026-04-30','Duplicate',True)
        self.assertEqual(before,self.s);fa.post_depreciation(self.s,'2026-05-31');self.assertEqual(fa.accumulated(self.s,a,'2026-05-31'),20000)
    def test_shared_prior_journal_capacity_is_allocated_per_asset(self):
        a=self.create();b=self.create(reference='EQ-2');v=core.post(self.s,'2026-04-30','Combined prior depreciation',pair('4570','1050',20000))
        fa.link_prior_depreciation(self.s,a['id'],v['id'],'100','2026-04-30','Reviewed allocation A',True);fa.link_prior_depreciation(self.s,b['id'],v['id'],'100','2026-04-30','Reviewed allocation B',True);self.assertTrue(self.result()['ready'])
        core.reverse(self.s,v['id'],'2026-05-01');self.assertEqual(fa.accumulated(self.s,a,'2026-05-01'),0);self.assertEqual(fa.accumulated(self.s,b,'2026-05-01'),0)
    def test_historical_source_link_after_later_reversals_preserves_asof_differences(self):
        a=self.create();june=fa.post_depreciation(self.s,'2026-06-30');core.reverse(self.s,june['voucher'],'2026-07-01')
        manual=core.post(self.s,'2026-05-31','Backfilled existing manual depreciation',pair('4570','1050',20000));fa.link_prior_depreciation(self.s,a['id'],manual['id'],'200','2026-05-31','Reviewed historical source; later duplicate reversed',True)
        self.assertFalse(self.result('2026-06-30')['ready']);fa.post_depreciation(self.s,'2026-07-31');self.assertTrue(self.result('2026-07-31')['ready'])
