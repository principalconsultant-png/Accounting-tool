import copy
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from pypdf import PdfReader
from riju import core,accounting_checks as c,financial_reports as f,history
from riju.storage import Vault

def pair(dr,cr,n,b='JP'):return [dict(account=dr,branch=b,debit=n,credit=0),dict(account=cr,branch=b,debit=0,credit=n)]

class ChecksTests(unittest.TestCase):
    def setUp(self):
        self.s=core.empty_state();core.post(self.s,'2025-04-01','Opening',pair('1010','3000',100000),opening=True)
        core.post(self.s,'2025-05-01','Prior revenue',pair('1010','4000',10000));core.post(self.s,'2026-05-01','Current revenue',pair('1000','4000',20000))
        self.rent=core.post(self.s,'2026-05-02','Rent',pair('4530','1000',30000),metadata={'payee':'Synthetic landlord','reference':'RENT-1'})
        self.pol=dict(effective='2026-04-01',cash_payment='200',cumulative_tds_base='200',unusual_posting='250',clearing_days='3',basis='Synthetic review thresholds, not law',confirmed=True);c.save_policy(self.s,self.pol)
    def runchecks(self,start='2026-05-01',end='2026-05-31',branch='ALL'):return c.run(self.s,start,end,branch)
    def answer(self,q,branch='ALL'):return c.answer(self.s,q,'2026-05-01','2026-05-31',branch)
    def test_threshold_cash_negative_cumulative_readonly_and_units(self):
        before=copy.deepcopy(self.s)
        with patch('socket.socket',side_effect=AssertionError('Offline')):r=self.runchecks()
        rules={a['rule'] for a in r['alerts']};self.assertTrue({'Cash-payment review','Negative cash','Potential cumulative TDS source','Large posting'}<=rules);self.assertEqual(before,self.s)
        cash=next(a for a in r['alerts'] if a['rule']=='Cash-payment review');self.assertEqual(cash['sources'],[self.rent['id']]);self.assertEqual(cash['amount'],30000)
        self.assertFalse(self.runchecks(branch='KM')['alerts'])
        c.save_policy(self.s,dict(self.pol,effective='2026-06-01',cash_payment='9999'));self.assertEqual(c.policy(self.s,'2026-05-31')['thresholds']['cash_payment'],20000)
        old=core.empty_state();self.assertEqual(c.policy(old,'2026-05-31')['id'],'UNREVIEWED');self.assertFalse(c.run(old,'2026-05-01','2026-05-31')['alerts'])
    def test_duplicate_across_months_reversals_and_review_staleness(self):
        later=core.post(self.s,'2026-06-01','Second rent',pair('4530','1010',30000),metadata={'payee':'Synthetic landlord','reference':'rent-1'})
        r=self.runchecks('2026-06-01','2026-06-30');a=next(a for a in r['alerts'] if a['rule']=='Potential duplicate source');self.assertEqual(set(a['sources']),{self.rent['id'],later['id']})
        c.review(self.s,r,a['id'],'Explained','Two genuine independently evidenced stages');self.assertEqual(next(a for a in self.runchecks('2026-06-01','2026-06-30')['alerts'] if a['rule']=='Potential duplicate source')['status'],'Explained')
        core.reverse(self.s,later['id'],'2026-06-02');self.assertFalse(any(a['rule']=='Potential duplicate source' for a in self.runchecks('2026-06-01','2026-06-30')['alerts']))
        with self.assertRaises(ValueError):c.review(self.s,r,a['id'],'Explained','stale')
        core.verify_audit(self.s)
    def test_tds_backdated_cumulative_source_shortfall(self):
        from test_tds import TdsTests
        helper=TdsTests();helper.setUp();helper.expense();helper.expense(date='2026-04-01',reference='LATE')
        r=c.run(helper.s,'2026-04-01','2026-04-30');a=next(a for a in r['alerts'] if a['rule']=='Cumulative TDS shortfall');self.assertEqual(a['amount'],12000);self.assertEqual(len(a['sources']),2);self.assertEqual(a['branch'],'Entity')
    def test_cwip_manual_difference_and_budget_sources(self):
        from test_cwip import CwipTests
        helper=CwipTests();helper.setUp();helper.cost();core.post(helper.s,'2026-04-04','Unclassified capital',pair('1020','2000',10000,'VK'))
        r=c.run(helper.s,'2026-04-01','2026-04-30','VK');self.assertTrue(any(a['rule']=='CWIP classification/reconciliation' and a['amount']==10000 for a in r['alerts']))
        ans=c.answer(helper.s,'CWIP balance','2026-04-01','2026-04-30','VK');self.assertTrue(any(p['mode']=='CWIP ledger reconciliation' for p in ans['packs']))
    def test_answers_arithmetic_vouchers_comparatives_and_refusals(self):
        before=copy.deepcopy(self.s)
        with patch('socket.socket',side_effect=AssertionError('Offline')):
            for q,amount in [('cash balance',-10000),('current cash balance',-10000),('current total expenses',30000),('bank balance',110000),('profit',-10000),('total expenses',30000),('rent expense',30000),('4000 balance',30000),('revenue',20000)]:
                r=self.answer(q);self.assertEqual(r['amount'],amount,q);self.assertEqual(sum(row[-1] for row in r['packs'][0]['rows']),amount)
        r=self.answer('compare profit');self.assertEqual(len(r['packs']),3);self.assertTrue(any(row[1]==self.rent['id'] for row in r['packs'][1]['rows']));self.assertTrue(r['packs'][2]['rows'])
        self.assertEqual(self.answer('cash balance at Kammanahalli')['amount'],0);self.assertEqual(before,self.s)
        for q in ['delete voucher','how much rent was paid','why cash decreased','cash receipts','top revenue','travel expenses','cash balance 2026','cash balance yesterday','cash balance 2026-05-03','balances at JP Nagar and Vaikom','is this legally compliant','forecast profit']:
            with self.subTest(q=q),self.assertRaises(ValueError):self.answer(q)
    def test_exports_literal_sources_review_history_and_persistence(self):
        r=self.runchecks();c.review(self.s,r,r['alerts'][0]['id'],'Investigating','=Literal evidence <&>')
        packs=[c.snapshots(self.s,'2026-05-01','2026-05-31'),*self.answer('rent expense')['packs']];before=copy.deepcopy(self.s)
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);f.export_pack_xlsx(packs,root/'checks.xlsx');f.export_pdf(packs,root/'checks.pdf');wb=load_workbook(root/'checks.xlsx');self.assertEqual(len(wb.sheetnames),3);wb.close();text=' '.join(p.extract_text() for p in PdfReader(root/'checks.pdf').pages);self.assertIn(self.rent['id'],text);self.assertIn('Literal evidence',text)
            history.checkpoint(self.s);vault=Vault(root/'checks.json',demo=True);vault.save(self.s);self.assertEqual(vault.load(),self.s)
        self.assertEqual(before['vouchers'],self.s['vouchers'])
    def test_policy_unknown_blank_zero_and_invalid_atomic(self):
        before=copy.deepcopy(self.s)
        for v in [dict(confirmed=False),dict(cash_payment='-1'),dict(clearing_days='1.5'),dict(effective='bad')]:
            with self.assertRaises(ValueError):c.save_policy(self.s,dict(self.pol,**v))
        self.assertEqual(before,self.s);c.save_policy(self.s,dict(self.pol,cash_payment='',unusual_posting='0'));p=c.policy(self.s,'2026-05-31');self.assertIsNone(p['thresholds']['cash_payment']);self.assertEqual(p['thresholds']['unusual_posting'],0)

    def test_cash_split_units_and_historical_clearing(self):
        core.post(self.s,'2026-05-02','Cash landlord other unit',pair('4530','1000',10000,'KM'),metadata={'payee':'Synthetic landlord'})
        a=next(a for a in self.runchecks()['alerts'] if a['rule']=='Cash-payment review');self.assertEqual(a['amount'],40000)
        from riju import reconciliation as r
        receipt=core.post(self.s,'2026-05-01','UPI service',pair('1110','4000',20000))
        r.import_statement(self.s,b'date,narration,reference,debit,credit,balance\n2026-06-02,UPI settlement,UP1,0,200,1300\n','upi.csv','1010','ALL','2026-06-01','2026-06-30','1100','1300')
        row=self.s['bank_lines'][0];r.settle(self.s,row['id'],'1110',[receipt['id']+':0'])
        self.assertTrue(any(a['rule']=='Aged collection clearing' for a in self.runchecks()['alerts']))
        self.assertFalse(any(a['rule']=='Aged collection clearing' for a in self.runchecks('2026-06-01','2026-06-30')['alerts']))
