import copy
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader
from riju import core,foundation,statements as s,financial_reports as f,history
from riju.storage import Vault

def pair(dr,cr,n,b='JP'):return [dict(account=dr,branch=b,debit=n,credit=0),dict(account=cr,branch=b,debit=0,credit=n)]

class StatementTests(unittest.TestCase):
    def setUp(self):
        self.s=core.empty_state();foundation.save_account(self.s,'2500','Current tax payable','Liability','Other liabilities')
        core.post(self.s,'2025-04-01','Opening',pair('1010','3000',100000),opening=True)
        core.post(self.s,'2025-05-01','Prior service',pair('1010','4000',20000));core.post(self.s,'2026-05-01','Service',pair('1010','4000',40000));core.post(self.s,'2026-05-02','Rent',pair('4530','1010',10000))
        self.save()
    def save(self,**kwargs):
        mappings=s.proposal(self.s);mappings['2500']['section']='Current liabilities: current tax'
        values=dict(start='2026-04-01',end='2027-03-31',framework=s.FRAMEWORKS[0],basis='Synthetic accountant review only',mappings=mappings,disclosures={},confirmed=True);values.update(kwargs);return s.save_policy(self.s,values)
    def tax(self,**kwargs):
        v=dict(start='2026-04-01',end='2027-03-31',branch='JP',expense_account='4610',payable_account='2500',rate='25',surcharge='0',cess='4',additions='0',deductions='0',losses='0',basis='Synthetic effective tax assumptions',confirmed=True);v.update(kwargs);return s.tax_estimate(self.s,v)
    def test_comparative_profit_balance_equity_and_readonly(self):
        before=copy.deepcopy(self.s);packs=s.report_pack(self.s,'2026-04-01','2027-03-31');bs,pl=packs[:2]
        self.assertFalse(any(bs['controls'].values()));result=next(r for r in pl['rows'] if r[2]=='Profit after tax');self.assertEqual(result[3:5],[30000,20000]);self.assertTrue(all(x['status']=='Review required' for x in packs));self.assertEqual(before,self.s)
        self.assertTrue(all(len(r)==len(p['columns']) for p in packs for r in p['rows']))
    def test_tax_delta_repeated_review_reversal_and_stale_preview(self):
        p=self.tax();self.assertEqual((p['pbt'],p['target'],p['delta']),(30000,7800,7800));v=s.post_tax(self.s,p)
        with self.assertRaises(ValueError):s.post_tax(self.s,p)
        zero=self.tax();self.assertEqual((zero['pbt'],zero['target'],zero['delta']),(30000,7800,0))
        with self.assertRaises(ValueError):s.post_tax(self.s,zero)
        core.reverse(self.s,v['id'],'2027-03-31');self.assertEqual(self.tax()['delta'],7800)
        with self.assertRaises(ValueError):self.tax(rate='NaN')
        with self.assertRaises(ValueError):self.tax(losses='999')
        with self.assertRaises(ValueError):self.tax(branch='ALL')
    def test_manual_booked_tax_avoids_duplicate_and_reduction_works(self):
        core.post(self.s,'2026-06-01','Manual provision',pair('4610','2500',10000))
        p=self.tax();self.assertEqual(p['delta'],-2200);s.post_tax(self.s,p);self.assertEqual(self.tax()['delta'],0)
        pl=s.report_pack(self.s,'2026-04-01','2027-03-31')[1];self.assertEqual(next(r for r in pl['rows'] if r[2]=='Profit after tax')[3],22200)
    def test_branch_clearing_eliminates_only_net_zero(self):
        core.post(self.s,'2026-05-03','Inter-unit',[dict(account='1130',branch='JP',debit=1000,credit=0),dict(account='1010',branch='JP',debit=0,credit=1000),dict(account='1010',branch='KM',debit=1000,credit=0),dict(account='1130',branch='KM',debit=0,credit=1000)])
        packs=s.report_pack(self.s,'2026-04-01','2027-03-31');clear=next(p for p in packs if p['mode']=='Inter-branch reconciliation');self.assertEqual(clear['rows'][-1][3:5],[1000,1000]);self.assertEqual(clear['controls']['Unreconciled inter-branch clearing'],0)
        branch=s.report_pack(self.s,'2026-04-01','2027-03-31','JP');self.assertFalse(any(branch[0]['controls'].values()))
        core.post(self.s,'2026-06-01','Uncleared',pair('1130','1010',500));clear=next(p for p in s.report_pack(self.s,'2026-04-01','2027-03-31') if p['mode']=='Inter-branch reconciliation');self.assertEqual(clear['rows'][-1][3:],[0,0,500]);self.assertEqual(clear['controls']['Unreconciled inter-branch clearing'],500)
    def test_mapping_disclosures_and_period_gates(self):
        with self.assertRaises(ValueError):self.save(confirmed=False)
        with self.assertRaises(ValueError):self.save(framework='Division III')
        with self.assertRaises(ValueError):s.report_pack(self.s,'2025-04-01','2026-03-31')
        before=copy.deepcopy(self.s);m=s.proposal(self.s);m.pop('1000')
        with self.assertRaises(ValueError):self.save(mappings=m)
        self.assertEqual(before,self.s)
        foundation.save_account(self.s,'4599','New expense','Expense','Other expenses')
        with self.assertRaises(ValueError):s.report_pack(self.s,'2026-04-01','2027-03-31')
        self.save(disclosures={t:dict(status='Not applicable',text='Synthetic evidenced review for '+t) for t in s.DISCLOSURES});packs=s.report_pack(self.s,'2026-04-01','2027-03-31');self.assertTrue(all(p['status']=='Prepared for accountant review' for p in packs))
    def test_excel_pdf_values_notes_and_persistence(self):
        self.save(framework=s.FRAMEWORKS[1],disclosures={s.DISCLOSURES[0]:dict(status='Reviewed',text='=Literal entity evidence <&>')})
        packs=s.report_pack(self.s,'2026-04-01','2027-03-31');before=copy.deepcopy(self.s)
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);f.export_pack_xlsx(packs,root/'s.xlsx');f.export_pdf(packs,root/'s.pdf');f.export_xlsx(s.tax_pack(self.s,self.tax()),root/'tax.xlsx');f.export_pdf(s.tax_pack(self.s,self.tax()),root/'tax.pdf')
            wb=load_workbook(root/'s.xlsx');self.assertEqual(len(wb.sheetnames),len(packs)+1);self.assertIn('Disclosure review',wb.sheetnames);wb.close();text=' '.join(p.extract_text() for p in PdfReader(root/'s.pdf').pages);self.assertIn('Inter-branch reconciliation',text);self.assertIn('Literal entity evidence',text)
            history.checkpoint(self.s);vault=Vault(root/'s.json',demo=True);vault.save(self.s);self.assertEqual(vault.load(),self.s)
        self.assertEqual(before['vouchers'],self.s['vouchers'])
