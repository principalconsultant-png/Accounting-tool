import copy
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
from pypdf import PdfReader
from riju import core,foundation,financial_reports as f


def pair(dr,cr,n,branch='JP'):return [dict(account=dr,branch=branch,debit=n,credit=0),dict(account=cr,branch=branch,debit=0,credit=n)]


class FinancialReportTests(unittest.TestCase):
    def setUp(self):
        self.s=core.empty_state();core.post(self.s,'2024-04-01','Opening',pair('1010','3000',100000),opening=True)
        core.post(self.s,'2024-05-01','Previous service',pair('1010','4000',20000));core.post(self.s,'2025-05-01','Current service',pair('1010','4000',30000));core.post(self.s,'2025-05-02','Rent',pair('4530','1010',10000))
    def report(self,mode,branch='ALL'):return f.snapshot(self.s,mode,'2025-04-01','2026-03-31',branch)
    def test_all_modes_readonly_and_proofs(self):
        before=copy.deepcopy(self.s)
        for mode in f.MODES:
            p=self.report(mode);self.assertTrue(p['columns']);self.assertTrue(p['rows'])
            self.assertTrue(all(len(row)==len(p['columns']) for row in p['rows']))
            if mode not in ('Profit and loss','Comparative P&L'):self.assertFalse(any(p['controls'].values()),mode)
        self.assertEqual(before,self.s)
    def test_pl_comparative_and_balance_sheet_accumulated_profit(self):
        p=self.report('Comparative P&L');profit=next(r for r in p['rows'] if r[2]=='Profit / (loss)');self.assertEqual(profit[3:],[20000,20000,0])
        bs=self.report('Balance sheet');result=next(r for r in bs['rows'] if r[2]=='Unclosed accumulated result');self.assertEqual(result[3],40000)
        self.assertEqual(bs['controls']['Assets minus liabilities and equity'],0)
    def test_openings_count_once_and_branch_scope(self):
        tb=f.snapshot(self.s,'Trial balance','2024-04-01','2025-03-31');bank=next(r for r in tb['rows'] if r[0]=='1010')
        self.assertEqual(bank[2:],[100000,0,20000,0,120000,0])
        self.assertEqual(self.report('Profit and loss','KM')['rows'][-1][3],0)
    def test_year_end_close_does_not_duplicate_profit_in_balance_sheet(self):
        foundation.close_year(self.s,2024)
        bs=self.report('Balance sheet');self.assertEqual(next(r[3] for r in bs['rows'] if r[2]=='Unclosed accumulated result'),20000)
        self.assertEqual(bs['controls']['Assets minus liabilities and equity'],0)
    def test_leap_comparative_and_invalid_dates(self):
        self.assertEqual(f.previous_year('2024-02-29'),'2023-02-28')
        with self.assertRaises(ValueError):f.snapshot(self.s,'Trial balance','2026-01-01','2025-01-01')
        with self.assertRaises(ValueError):self.report('Unknown')
    def test_matrix_all_units_despite_branch_selector(self):
        matrix=self.report('Branch matrix','KM');self.assertEqual(matrix['branch'],'ALL');self.assertIn('JP Nagar',matrix['columns']);self.assertEqual(matrix['controls']['Branch-to-entity difference'],0)
    def test_excel_numbers_literal_text_and_pdf_for_every_mode(self):
        self.s['settings']['entity']='=Hospital <Test>'
        before=copy.deepcopy(self.s)
        with tempfile.TemporaryDirectory() as d:
            root=Path(d)
            for index,mode in enumerate(f.MODES):
                pack=self.report(mode);xlsx=root/f'{index}.xlsx';pdf=root/f'{index}.pdf';f.export_xlsx(pack,xlsx);f.export_pdf(pack,pdf)
                wb=load_workbook(xlsx,data_only=False);ws=wb.active;self.assertEqual(ws['A1'].value,'=Hospital <Test>');self.assertEqual(ws['A1'].data_type,'s')
                header=next(r[0].row for r in ws.iter_rows() if r[0].value==pack['columns'][0] and len(r)>1 and r[1].value==pack['columns'][1])
                for i,row in enumerate(pack['rows'],header+1):
                    for col in pack['money_columns']:self.assertAlmostEqual(ws.cell(i,col+1).value,row[col]/100)
                self.assertTrue(ws.freeze_panes);self.assertTrue(ws.print_title_rows);wb.close()
                text='\n'.join(p.extract_text() for p in PdfReader(pdf).pages);self.assertIn(mode,text);self.assertIn('2025-04-01',text);self.assertIn('Page 1',text)
        self.assertEqual(before,self.s)
    def test_long_pdf_paginates_and_wide_matrix_keeps_all_branches(self):
        for i in range(65):core.post(self.s,'2025-05-03','Long '+('test narrative <&> '*18),pair('1010','4000',100))
        for i in range(8):foundation.save_branch(self.s,'X'+str(i),dict(name='Additional unit '+str(i),beds=0))
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'gl.pdf';f.export_pdf(self.report('General ledger'),p);self.assertGreater(len(PdfReader(p).pages),1)
            f.export_pdf(self.report('Branch matrix'),p);text=' '.join(page.extract_text() for page in PdfReader(p).pages);self.assertIn('Additional unit 7',text);self.assertIn('Entity total',text)
    def test_csv_retains_formula_safety(self):
        self.s['settings']['entity']='=Formula'
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'report.csv';f.export_csv(self.report('Profit and loss'),p);self.assertIn("'=Formula",p.read_text(encoding='utf-8-sig'))

    def test_customer_pack_workbook_tabs_and_combined_pdf_are_readonly(self):
        packs=f.report_pack(self.s,'2025-04-01','2026-03-31','KM');before=copy.deepcopy(self.s)
        self.assertEqual(len({p['prepared'] for p in packs}),1)
        with tempfile.TemporaryDirectory() as d:
            root=Path(d);f.export_pack_xlsx(packs,root/'pack.xlsx');f.export_pdf(packs,root/'pack.pdf')
            wb=load_workbook(root/'pack.xlsx');self.assertEqual(wb.sheetnames,['Contents',*f.MODES]);self.assertTrue(wb['Contents']['A7'].hyperlink);self.assertEqual(wb['Branch matrix']['A3'].value,'All operating units');wb.close()
            pages=PdfReader(root/'pack.pdf').pages;text='\n'.join(page.extract_text() for page in pages)
            for mode in f.MODES:self.assertIn(mode,text)
            self.assertIn('Mixed scopes; see section headings',text);self.assertGreater(len(pages),7)
        self.assertEqual(before,self.s)
